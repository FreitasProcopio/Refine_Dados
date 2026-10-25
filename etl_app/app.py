"""
Refina — Refinaria de dados em camadas Bronze, Prata e Ouro (desktop, PySide6)
=======================================================
Os dados ficam num banco DuckDB local, em disco, e não na RAM:
  • arquivos de rede são copiados em blocos grandes para um cache local (e reaproveitados);
  • planilhas são lidas em blocos de linhas e gravadas aos poucos no banco;
  • as transformações rodam em SQL com limite de memória (o excedente vai para o disco);
  • a tela só busca as linhas que estão visíveis.

Uso:
    python app.py                         # abre a interface
    python app.py --run receita.json      # executa uma receita sem interface
    Refina.exe --run receita.json         # idem, depois de compilado
"""
from __future__ import annotations

import os
import sys

# ---------------------------------------------------------------------------
# Zoom da interface (precisa ser definido antes de o Qt iniciar)
# ---------------------------------------------------------------------------
UI_SCALES = (0.75, 0.8, 0.85, 0.9, 1.0, 1.1, 1.25)
DEFAULT_UI_SCALE = 0.9


def saved_ui_scale() -> float:
    try:
        from PySide6.QtCore import QSettings as _S
        v = float(_S("Refina", "Refina").value("ui_scale", DEFAULT_UI_SCALE))
    except Exception:
        v = DEFAULT_UI_SCALE
    return min(UI_SCALES, key=lambda x: abs(x - v))


def apply_ui_scale():
    os.environ["QT_SCALE_FACTOR"] = f"{saved_ui_scale():g}"


# ---------------------------------------------------------------------------
# Tela de abertura: aparece logo no início, antes de carregar pandas/DuckDB/etc.
# ---------------------------------------------------------------------------
_SPLASH = None
_EARLY_APP = None


def _early_splash():
    global _SPLASH, _EARLY_APP
    if __name__ != "__main__" or "--run" in sys.argv:
        return
    try:
        from PySide6.QtCore import QPointF, QRectF, Qt as _Qt
        from PySide6.QtGui import QColor as _C, QFont as _F, QPainter as _P, QPixmap as _Pix, QPolygonF
        from PySide6.QtWidgets import QApplication as _App, QSplashScreen as _Splash
    except Exception:
        return

    class Splash(_Splash):
        W, H = 460, 250

        def __init__(self):
            ratio = 2
            pm = _Pix(self.W * ratio, self.H * ratio)
            pm.setDevicePixelRatio(ratio)
            pm.fill(_C("#1E1A28"))
            p = _P(pm)
            p.setRenderHint(_P.Antialiasing)
            p.setPen(_Qt.NoPen)
            # funil em três faixas: o dado bruto entra largo (bronze) e sai refinado (ouro)
            for pts, c in (([(30, 62), (82, 62), (74, 78), (38, 78)], "#B4703A"),
                           ([(40, 82), (72, 82), (64, 98), (48, 98)], "#7E8A9C"),
                           ([(50, 102), (62, 102), (59, 112), (53, 112)], "#C49A1B")):
                p.setBrush(_C(c))
                p.drawPolygon(QPolygonF([QPointF(x, y) for x, y in pts]))
            p.setPen(_C("#FFFFFF"))
            f = _F("Segoe UI Variable Display")
            f.setFamilies(["Segoe UI Variable Display", "Segoe UI", "Inter", "Noto Sans"])
            f.setPixelSize(30)
            f.setWeight(_F.DemiBold)
            p.setFont(f)
            p.drawText(QRectF(100, 58, 340, 40), _Qt.AlignLeft | _Qt.AlignVCenter, "Refina")
            f.setPixelSize(14)
            f.setWeight(_F.Normal)
            p.setFont(f)
            p.setPen(_C("#8F89A6"))
            p.drawText(QRectF(101, 92, 340, 24), _Qt.AlignLeft | _Qt.AlignVCenter, "Refinaria de dados")
            p.end()
            super().__init__(pm)
            self.setWindowFlag(_Qt.WindowStaysOnTopHint, True)
            self._msg, self._pct = "Abrindo…", 5

        def step(self, msg, pct):
            self._msg, self._pct = msg, pct
            self.repaint()
            _App.processEvents()

        def drawContents(self, p):
            p.setRenderHint(_P.Antialiasing)
            f = _F()
            f.setFamilies(["Segoe UI Variable Text", "Segoe UI", "Inter", "Noto Sans"])
            f.setPixelSize(13)
            p.setFont(f)
            p.setPen(_C("#BAB4CC"))
            p.drawText(QRectF(32, 170, 400, 22), _Qt.AlignLeft | _Qt.AlignVCenter, self._msg)
            p.setPen(_Qt.NoPen)
            p.setBrush(_C("#332D44"))
            p.drawRoundedRect(QRectF(32, 202, 396, 6), 3, 3)
            p.setBrush(_C("#5B4FD6"))
            p.drawRoundedRect(QRectF(32, 202, max(8, 396 * self._pct / 100), 6), 3, 3)

    if _App.instance() is None:
        apply_ui_scale()
    _EARLY_APP = _App.instance() or _App(sys.argv)
    _SPLASH = Splash()
    _SPLASH.show()
    _SPLASH.step("Carregando as bibliotecas de dados…", 10)


def splash_step(msg, pct):
    if _SPLASH is not None:
        _SPLASH.step(msg, pct)


_early_splash()

import contextlib
import ctypes
import datetime as dt
import decimal
import hashlib
import itertools
import json
import re
import shutil
import sqlite3
import tempfile
import unicodedata
import threading
import time
import traceback
import xml.etree.ElementTree as ET
import ssl
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import duckdb
import numpy as np
import pandas as pd

try:
    from python_calamine import CalamineWorkbook
    HAS_CALAMINE = True
except ImportError:
    HAS_CALAMINE = False

from PySide6.QtCore import (QAbstractTableModel, QByteArray, QEvent, QModelIndex, QObject, QProcess, QRunnable,
                            QSettings,
                            QSize, Qt, QThreadPool, QTimer, QUrl, Signal)
from PySide6.QtGui import QColor, QDesktopServices, QIcon, QKeySequence, QPainter, QPalette, QPixmap, QShortcut
from PySide6.QtSvg import QSvgRenderer
splash_step("Carregando a interface…", 45)
from PySide6.QtWidgets import (QAbstractItemView, QApplication, QCheckBox, QComboBox, QDialog,
                               QDialogButtonBox, QFileDialog, QFormLayout, QFrame, QGroupBox,
                               QHBoxLayout, QHeaderView, QLabel, QLineEdit, QListWidget,
                               QListWidgetItem, QMainWindow, QMessageBox, QPlainTextEdit, QProgressDialog,
                               QPushButton, QScrollArea, QSizePolicy, QSpinBox, QSplitter, QStackedWidget,
                               QTableView, QTableWidget, QTableWidgetItem, QTabWidget, QToolButton, QTreeWidget, QTreeWidgetItem, QVBoxLayout, QWidget)

# ============================================================================
# CONFIGURAÇÃO
# ============================================================================
APP_NAME = "Refina"
EXCEL_STREAM_EXT = {".xlsx", ".xlsm"}             # lidos em fluxo (openpyxl) ou rápido (calamine)
EXCEL_CALAMINE_EXT = {".xls", ".xlsb", ".ods"}    # só calamine
CSV_EXT = {".csv", ".txt"}
DB_EXT = {".db", ".sqlite", ".sqlite3", ".duckdb"}
PARQUET_EXT = {".parquet"}
JSON_EXT = {".json", ".jsonl", ".ndjson"}
SUPPORTED_EXTENSIONS = EXCEL_STREAM_EXT | EXCEL_CALAMINE_EXT | CSV_EXT | DB_EXT | PARQUET_EXT | JSON_EXT
FILE_FILTER = ("Dados (*.xlsx *.xlsm *.xlsb *.xls *.ods *.csv *.txt *.parquet *.json *.db *.sqlite *.sqlite3 *.duckdb)"
               ";;Todos os arquivos (*.*)")

JOIN_TYPES = {"left": "LEFT JOIN", "inner": "INNER JOIN", "full": "FULL OUTER JOIN", "cross": "CROSS JOIN"}
AGG_FUNCTIONS = ["sum", "mean", "count", "min", "max", "nunique", "first", "last"]
FILTER_OPERATORS = ["contém", "começa com", "termina com", "igual a", "diferente de",
                    "está na lista", "não está na lista", "maior que", "menor que",
                    "maior ou igual", "menor ou igual", "não é nulo", "é nulo"]
NUMERIC_OPERATORS = ("maior que", "menor que", "maior ou igual", "menor ou igual")
NULL_OPERATORS = ("não é nulo", "é nulo")
LIST_OPERATORS = ("está na lista", "não está na lista")
TEXT_CASES = {"Manter como está": "", "MAIÚSCULAS": "upper", "minúsculas": "lower"}
CAST_TYPES = ["int", "float", "str", "datetime", "bool"]
EXCEL_MAX_ROWS = 1_048_575
UNDO_DEPTH = 30
DEFAULT_CACHE_GB = 10

LAYERS = ("bronze", "silver", "gold")
LAYER_LABEL = {"bronze": "Bronze", "silver": "Prata", "gold": "Ouro"}
LAYER_COLOR = {"bronze": "#B4703A", "silver": "#7E8A9C", "gold": "#C49A1B"}
SILVER_SOURCES = ("bronze", "silver")
GOLD_SOURCES = ("silver", "gold")


def fmt_int(n) -> str:
    return f"{int(n):,}".replace(",", ".")


def fmt_bytes(n) -> str:
    n = float(n or 0)
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if n < 1024 or unit == "TB":
            return f"{n:.0f} {unit}" if unit in ("B", "KB") else f"{n:.1f} {unit}".replace(".", ",")
        n /= 1024


def qi(name) -> str:
    """Identificador SQL entre aspas."""
    return '"' + str(name).replace('"', '""') + '"'


def ql(value) -> str:
    """Texto literal SQL."""
    return "'" + str(value).replace("'", "''") + "'"


# ============================================================================
# SISTEMA: memória, pastas de trabalho, detecção de caminho de rede
# ============================================================================
def total_ram_mb() -> int:
    try:
        if os.name == "nt":
            class MEMSTAT(ctypes.Structure):
                _fields_ = [("dwLength", ctypes.c_ulong), ("dwMemoryLoad", ctypes.c_ulong),
                            ("ullTotalPhys", ctypes.c_ulonglong), ("ullAvailPhys", ctypes.c_ulonglong),
                            ("ullTotalPageFile", ctypes.c_ulonglong), ("ullAvailPageFile", ctypes.c_ulonglong),
                            ("ullTotalVirtual", ctypes.c_ulonglong), ("ullAvailVirtual", ctypes.c_ulonglong),
                            ("ullAvailExtendedVirtual", ctypes.c_ulonglong)]
            st = MEMSTAT()
            st.dwLength = ctypes.sizeof(MEMSTAT)
            ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(st))
            return int(st.ullTotalPhys // 2**20)
        return int(os.sysconf("SC_PAGE_SIZE") * os.sysconf("SC_PHYS_PAGES") // 2**20)
    except Exception:
        return 4096


def auto_mem_mb() -> int:
    """Padrão: 40% da RAM, entre 512 MB e 8 GB."""
    return int(min(max(total_ram_mb() * 0.4, 512), 8192))


def app_root() -> Path:
    base = os.environ.get("LOCALAPPDATA") or str(Path.home() / ".local" / "share")
    return Path(base) / "Refina"


def work_root() -> Path:
    return app_root() / "work"


def cache_root() -> Path:
    return app_root() / "cache"


def is_remote(path) -> bool:
    s = str(path)
    if s.startswith("\\\\") or s.startswith("//"):
        return True
    if os.name == "nt":
        drive = os.path.splitdrive(os.path.abspath(s))[0]
        if drive:
            try:
                return ctypes.windll.kernel32.GetDriveTypeW(drive + "\\") == 4   # DRIVE_REMOTE
            except Exception:
                return False
    return False


def clean_stale_sessions():
    """Apaga sessões antigas que ficaram para trás (ex.: o app foi fechado à força)."""
    root = work_root()
    if not root.exists():
        return
    limit = time.time() - 6 * 3600
    for d in root.glob("sessao_*"):
        try:
            if d.stat().st_mtime < limit:
                shutil.rmtree(d)
        except Exception:
            pass


# ============================================================================
# CONTEXTO DE EXECUÇÃO (progresso + cancelamento)
# ============================================================================
class Cancelled(Exception):
    pass


class Ctx:
    def __init__(self, on_progress=None):
        self._on = on_progress
        self._cancel = threading.Event()
        self._cons = set()
        self._last = 0.0
        self._lock = threading.Lock()

    @property
    def cancelled(self):
        return self._cancel.is_set()

    def progress(self, msg, pct=None, force=False):
        now = time.monotonic()
        if not force and now - self._last < 0.15:
            return
        self._last = now
        if self._on:
            self._on(msg, -1 if pct is None else int(max(0, min(100, pct))))

    def check(self):
        if self._cancel.is_set():
            raise Cancelled()

    def cancel(self):
        self._cancel.set()
        with self._lock:
            for con in list(self._cons):
                with contextlib.suppress(Exception):
                    con.interrupt()

    @contextlib.contextmanager
    def watch(self, con):
        with self._lock:
            self._cons.add(con)
        try:
            yield con
        finally:
            with self._lock:
                self._cons.discard(con)

    def query_progress(self) -> float:
        with self._lock:
            cons = list(self._cons)
        for con in cons:
            with contextlib.suppress(Exception):
                p = con.query_progress()
                if p is not None and p >= 0:
                    return float(p)
        return -1.0


# ============================================================================
# MOTOR: banco DuckDB da sessão (os dados ficam em disco, não na RAM)
# ============================================================================
class Engine:
    def __init__(self, mem_mb=0, cache_gb=DEFAULT_CACHE_GB):
        self.mem_mb = int(mem_mb) or auto_mem_mb()
        self.cache_gb = cache_gb
        root = work_root()
        root.mkdir(parents=True, exist_ok=True)
        self.dir = Path(tempfile.mkdtemp(prefix="sessao_", dir=str(root)))
        (self.dir / "spill").mkdir()
        self.con = duckdb.connect(str(self.dir / "dados.duckdb"), config={
            "memory_limit": f"{self.mem_mb}MB",
            "temp_directory": str(self.dir / "spill"),
        })
        self.ui = self.con.cursor()
        self._ids = itertools.count(1)
        self._id_lock = threading.Lock()
        self._meta = {}
        self._stats = {}
        self._profile = {}
        self.lineage = {}      # tabela → de onde veio (arquivo ou SQL), usado para gerar o código Python
        self.attached = None   # banco DuckDB externo conectado na tela "Banco de dados" (apelido "bd")
        self.last_sql_result = None
        self.last_report = None

    @property
    def budget(self) -> int:
        return self.mem_mb * 2**20

    def set_memory(self, mem_mb):
        self.mem_mb = int(mem_mb) or auto_mem_mb()
        self.ui.execute(f"SET memory_limit='{self.mem_mb}MB'")

    def new_tid(self) -> str:
        with self._id_lock:
            return f"t_{next(self._ids):06d}"

    def temp_file(self, suffix) -> Path:
        with self._id_lock:
            n = next(self._ids)
        return self.dir / f"tmp_{n}{suffix}"

    def cursor(self):
        return self.con.cursor()

    def meta(self, tid, con=None) -> dict:
        m = self._meta.get(tid)
        if m is None:
            c = con or self.ui
            desc = c.execute(f"DESCRIBE {tid}").fetchall()
            rows = c.execute(f"SELECT count(*) FROM {tid}").fetchone()[0]
            m = {"cols": [d[0] for d in desc], "types": [d[1] for d in desc], "rows": int(rows)}
            self._meta[tid] = m
        return m

    def columns(self, tid, con=None):
        return self.meta(tid, con)["cols"]

    def types(self, tid, con=None) -> dict:
        m = self.meta(tid, con)
        return dict(zip(m["cols"], m["types"]))

    def rows(self, tid):
        return self.meta(tid)["rows"]

    def null_counts(self, tid) -> list:
        if tid not in self._stats:
            cols = self.columns(tid)
            if not cols:
                self._stats[tid] = []
            else:
                expr = ", ".join(f"count(*) - count({qi(c)})" for c in cols)
                self._stats[tid] = [int(v) for v in self.ui.execute(f"SELECT {expr} FROM {tid}").fetchone()]
        return self._stats[tid]

    def forget(self, prefix):
        """Esquece o que estava em cache sobre tabelas que mudaram (ex.: do banco externo)."""
        for cache in (self._meta, self._stats, self._profile):
            for k in [k for k in cache if k.startswith(prefix)]:
                cache.pop(k, None)

    def tables(self):
        return [r[0] for r in self.ui.execute(
            "SELECT table_name FROM duckdb_tables() WHERE database_name = current_database()").fetchall()]

    def drop(self, name):
        with contextlib.suppress(Exception):
            self.ui.execute(f"DROP TABLE IF EXISTS {name}")
        self._meta.pop(name, None)
        self._stats.pop(name, None)
        self._profile.pop(name, None)

    def close(self):
        with contextlib.suppress(Exception):
            self.ui.close()
        with contextlib.suppress(Exception):
            self.con.close()
        shutil.rmtree(self.dir, ignore_errors=True)


# ============================================================================
# CÓPIA RÁPIDA DE ARQUIVOS DE REDE (cache local)
# ============================================================================
COPY_BLOCK = 8 * 2**20


def _cache_name(path: Path, st) -> str:
    h = hashlib.sha1(os.path.normcase(str(path)).encode("utf-8")).hexdigest()[:16]
    return f"{h}_{st.st_mtime_ns}_{st.st_size}{path.suffix.lower()}"


def cache_usage():
    root = cache_root()
    if not root.exists():
        return 0, 0
    files = [f for f in root.rglob("*") if f.is_file() and f.name != "meta.json"]
    return sum(f.stat().st_size for f in files), len(files)


def clear_cache():
    shutil.rmtree(cache_root(), ignore_errors=True)


def prune_cache(cap_bytes, keep=()):
    root = cache_root()
    if not root.exists():
        return
    keep = {os.path.normcase(str(k)) for k in keep}
    files = sorted((f for f in root.rglob("*") if f.is_file() and f.name != "meta.json"),
                   key=lambda f: f.stat().st_atime)
    total = sum(f.stat().st_size for f in files)
    for f in files:
        if total <= cap_bytes:
            break
        if os.path.normcase(str(f)) in keep:
            continue
        with contextlib.suppress(Exception):
            size = f.stat().st_size
            f.unlink()
            total -= size


def stage_files(files, ctx: Ctx, cache_gb=DEFAULT_CACHE_GB) -> dict:
    """Copia arquivos de rede para o cache local (em blocos de 8 MB, até 4 em paralelo).
    Se o arquivo não mudou desde a última cópia, reaproveita na hora."""
    out = {f: f for f in files}
    remote = [f for f in files if is_remote(f)]
    if not remote:
        return out
    root = cache_root()
    root.mkdir(parents=True, exist_ok=True)
    plan, total = [], 0
    for f in remote:
        st = os.stat(f)
        dest = root / _cache_name(f, st)
        out[f] = dest
        if dest.exists() and dest.stat().st_size == st.st_size:
            with contextlib.suppress(Exception):
                os.utime(dest)
            continue
        plan.append((f, dest))
        total += st.st_size
    if plan:
        done = [0]
        lock = threading.Lock()
        t0 = time.monotonic()

        def copy(item):
            src, dest = item
            tmp = dest.with_name(dest.name + f".{threading.get_ident()}.part")
            try:
                with open(src, "rb", buffering=0) as fi, open(tmp, "wb") as fo:
                    while True:
                        ctx.check()
                        block = fi.read(COPY_BLOCK)
                        if not block:
                            break
                        fo.write(block)
                        with lock:
                            done[0] += len(block)
                            speed = done[0] / max(time.monotonic() - t0, 1e-6)
                        ctx.progress(f"Copiando do servidor: {fmt_bytes(done[0])} de {fmt_bytes(total)} "
                                     f"({fmt_bytes(speed)}/s)", 100 * done[0] / max(total, 1))
                os.replace(tmp, dest)
            finally:
                with contextlib.suppress(Exception):
                    tmp.unlink()
            # remove versões antigas do mesmo arquivo
            prefix = dest.name.split("_", 1)[0] + "_"
            for old in root.glob(prefix + "*"):
                if old != dest and not old.name.endswith(".part"):
                    with contextlib.suppress(Exception):
                        old.unlink()

        with ThreadPoolExecutor(max_workers=min(4, len(plan))) as ex:
            list(ex.map(copy, plan))
    prune_cache(cache_gb * 2**30, keep=out.values())
    return out


# ============================================================================
# INSPEÇÃO: lista abas/tabelas sem carregar os dados
# ============================================================================
# ============================================================================
# LINKS DA INTERNET (http/https)
# ============================================================================
_URL_RE = re.compile(r"^https?://", re.I)
_WEB_MEMO = {}          # url → (arquivo, quando) — evita baixar duas vezes no mesmo carregamento
_WEB_MEMO_SECONDS = 600


def is_url(text) -> bool:
    return bool(_URL_RE.match(str(text).strip()))


_INVISIBLE = re.compile("[\u200b-\u200f\u202a-\u202e\u2060-\u2064\ufeff\u00ad]")


def clean_input(text) -> str:
    """Limpa o que foi colado: aspas, espaços e caracteres invisíveis que o Windows/Edge/Teams colocam
    no início do texto copiado (eles faziam um link ser tratado como pasta)."""
    t = _INVISIBLE.sub("", str(text)).strip().strip('\'"<>').strip()
    m = re.search(r"https?://", t, re.I)
    if m and m.start() > 0:                 # lixo antes do link
        t = t[m.start():]
    if t.lower().startswith("file:///"):     # file:///C:/pasta/arquivo.xlsx
        t = urllib.parse.unquote(t[8:])
    if not is_url(t) and "%" in t and not Path(t).exists() and Path(urllib.parse.unquote(t)).exists():
        t = urllib.parse.unquote(t)        # caminho colado com %20 no lugar dos espaços
    return t


def onedrive_local(url: str):
    """Link do OneDrive da empresa (…/personal/<usuário>/Documents/…) → arquivo sincronizado neste
    computador, quando existir. Assim não precisa de login nem de download."""
    p = urllib.parse.urlparse(url)
    if "sharepoint.com" not in p.netloc.lower():
        return None
    m = re.search(r"/personal/[^/]+/Documents/(.+)$", urllib.parse.unquote(p.path), re.I)
    if not m:
        return None
    rel = Path(*[part for part in m.group(1).split("/") if part])
    roots = [os.environ.get("OneDriveCommercial"), os.environ.get("OneDrive")]
    with contextlib.suppress(Exception):
        roots += [str(d) for d in Path.home().glob("OneDrive - *")]
    for root in dict.fromkeys(r for r in roots if r):
        cand = Path(root) / rel
        if cand.exists():
            return cand
    return None


def normalize_url(url: str) -> str:
    """Troca links de compartilhamento pelos links de download direto."""
    u = url.strip()
    p = urllib.parse.urlparse(u)
    host = p.netloc.lower()
    # Google Planilhas: .../spreadsheets/d/<id>/edit#gid=0 → exporta em .xlsx (todas as abas)
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", p.path)
    if "docs.google.com" in host and m and "/export" not in p.path and "output=" not in p.query:
        return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx"
    # Google Drive: .../file/d/<id>/view → download direto
    m = re.search(r"/file/d/([A-Za-z0-9_-]+)", p.path)
    if "drive.google.com" in host and m:
        return f"https://drive.google.com/uc?export=download&id={m.group(1)}"
    # Dropbox: ?dl=0 → ?dl=1
    if "dropbox.com" in host:
        q = dict(urllib.parse.parse_qsl(p.query))
        q["dl"] = "1"
        return urllib.parse.urlunparse(p._replace(query=urllib.parse.urlencode(q)))
    # OneDrive / SharePoint: link de compartilhamento → download
    if ("sharepoint.com" in host or "1drv.ms" in host or "onedrive.live.com" in host) and "download=1" not in p.query:
        return u + ("&" if p.query else "?") + "download=1"
    # GitHub: página do arquivo → conteúdo bruto
    m = re.match(r"^/([^/]+)/([^/]+)/blob/(.+)$", p.path)
    if host in ("github.com", "www.github.com") and m:
        return f"https://raw.githubusercontent.com/{m.group(1)}/{m.group(2)}/{m.group(3)}"
    return u


def _sniff_ext(head: bytes, ctype: str):
    lead = head.lstrip()[:15].lower()
    if "html" in ctype or lead.startswith((b"<!doctype html", b"<html")):
        return "html"
    if head.startswith(b"PK"):
        return ".xlsx"
    if head.startswith(b"PAR1"):
        return ".parquet"
    if head.startswith(b"SQLite format 3"):
        return ".db"
    if head[8:12] == b"DUCK":
        return ".duckdb"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return ".xls"
    if lead[:1] in (b"[", b"{") or "json" in ctype:
        return ".json"
    return ".csv"


def download_url(url: str, ctx: Ctx, timeout=60) -> Path:
    """Baixa o link para o cache local (em blocos, com progresso). Se o site informar que o arquivo
    não mudou desde a última vez (ETag/Last-Modified), reaproveita a cópia sem baixar de novo."""
    url = url.strip()
    memo = _WEB_MEMO.get(url)
    if memo and memo[0].exists() and time.time() - memo[1] < _WEB_MEMO_SECONDS:
        return memo[0]
    real = normalize_url(url)
    folder = cache_root() / "web" / hashlib.sha1(url.encode("utf-8")).hexdigest()[:16]
    folder.mkdir(parents=True, exist_ok=True)
    meta_path = folder / "meta.json"
    meta = {}
    with contextlib.suppress(Exception):
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
    cached = folder / meta["file"] if meta.get("file") else None
    host = urllib.parse.urlparse(real).netloc

    req = urllib.request.Request(real, headers={"User-Agent": f"{APP_NAME}/1.0 (Windows)", "Accept": "*/*"})
    if cached and cached.exists():
        if meta.get("etag"):
            req.add_header("If-None-Match", meta["etag"])
        if meta.get("last_modified"):
            req.add_header("If-Modified-Since", meta["last_modified"])

    ctx.progress(f"Conectando a {host}…", None, force=True)
    tls = ssl.create_default_context()       # usa os certificados do Windows
    resp = None
    for attempt in range(3):
        ctx.check()
        try:
            resp = urllib.request.urlopen(req, timeout=timeout, context=tls)
            break
        except urllib.error.HTTPError as exc:
            if exc.code == 304 and cached and cached.exists():
                _WEB_MEMO[url] = (cached, time.time())
                return cached
            if exc.code in (401, 403) and "sharepoint.com" in host.lower():
                raise ValueError(f"Não consegui baixar {url}: links do SharePoint/OneDrive da empresa exigem "
                                 "login. Se o arquivo estiver sincronizado no seu computador (pasta OneDrive ou "
                                 "da empresa no Explorador de Arquivos), carregue pelo caminho local: botão 'Arquivo…'.")
            msgs = {401: "o link pede login", 403: "acesso negado (o link pode ser privado ou exigir login)",
                    404: "arquivo não encontrado (confira o link)", 429: "o site limitou o número de acessos; "
                    "tente de novo em alguns minutos"}
            raise ValueError(f"Não consegui baixar {url}: {msgs.get(exc.code, f'o site respondeu erro {exc.code}')}.")
        except urllib.error.URLError as exc:
            reason = exc.reason
            if isinstance(reason, ssl.SSLCertVerificationError):
                raise ValueError(f"Não consegui baixar {url}: o certificado de segurança do site não foi aceito. "
                                 "Em redes corporativas isso acontece quando o proxy inspeciona o tráfego; "
                                 "peça à TI para instalar o certificado do proxy no Windows.")
            if attempt == 2:
                raise ValueError(f"Não consegui acessar {host}: {reason}. Confira o link e a conexão "
                                 "(em rede corporativa, o proxy pode estar bloqueando).")
        except TimeoutError:
            if attempt == 2:
                raise ValueError(f"O site {host} demorou demais para responder. Tente de novo mais tarde.")
        time.sleep(1.5 * (attempt + 1))

    with resp:
        ctype = (resp.headers.get("Content-Type") or "").lower()
        name = resp.headers.get_filename() or Path(urllib.parse.unquote(urllib.parse.urlparse(resp.geturl()).path)).name
        name = _safe_filename(name or "dados")
        total = int(resp.headers.get("Content-Length") or 0)
        tmp = folder / "download.part"
        done, t0, head = 0, time.monotonic(), b""
        with open(tmp, "wb") as fo:
            while True:
                ctx.check()
                block = resp.read(1 << 20)
                if not block:
                    break
                if len(head) < 64:
                    head += block[:64]
                fo.write(block)
                done += len(block)
                speed = done / max(time.monotonic() - t0, 1e-6)
                total_txt = f" de {fmt_bytes(total)}" if total else ""
                ctx.progress(f"Baixando de {host}: {fmt_bytes(done)}{total_txt} ({fmt_bytes(speed)}/s)",
                             100 * done / total if total else None)
        etag, last_mod = resp.headers.get("ETag"), resp.headers.get("Last-Modified")

    kind = _sniff_ext(head, ctype)
    if kind == "html":
        tmp.unlink(missing_ok=True)
        if "sharepoint.com" in host.lower() or "microsoftonline" in resp.geturl().lower():
            raise ValueError(f"O link {url} abriu a tela de login da Microsoft. Links do SharePoint/OneDrive da "
                             "empresa exigem login; se o arquivo estiver sincronizado no seu computador, "
                             "carregue pelo caminho local (botão 'Arquivo…').")
        raise ValueError(f"O link {url} abriu uma página da web, e não um arquivo de dados. Use o link de "
                         "download direto (em sites que pedem login, baixe o arquivo e carregue pela pasta).")
    if Path(name).suffix.lower() not in SUPPORTED_EXTENSIONS:
        name = Path(name).stem + kind
    final = folder / name
    os.replace(tmp, final)
    for old in folder.iterdir():
        if old.name not in (final.name, "meta.json"):
            with contextlib.suppress(Exception):
                old.unlink()
    meta = {"url": url, "file": final.name, "etag": etag, "last_modified": last_mod,
            "downloaded": dt.datetime.now().isoformat(timespec="seconds")}
    meta_path.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    _WEB_MEMO[url] = (final, time.time())
    return final


def resolve_inputs(paths, ctx: Ctx) -> list:
    """Transforma caminhos, pastas e links em [(arquivo_local, chave_de_origem, é_link)].
    A chave é o que fica gravado na receita (o link, no caso da internet)."""
    out = []
    for raw in paths:
        text = clean_input(raw)
        if not text:
            continue
        if is_url(text):
            local = onedrive_local(text)
            if local is not None and local.is_dir():
                out.extend((f, str(f), False) for f in expand_paths([str(local)]))
            elif local is not None:
                out.append((local, text, True))
            else:
                out.append((download_url(text, ctx), text, True))
        else:
            out.extend((f, str(f), False) for f in expand_paths([text]))
    if not out:
        raise ValueError("Nenhum arquivo suportado encontrado.")
    return out


def expand_paths(paths) -> list:
    out = []
    for raw in paths:
        p = Path(str(raw).strip().strip('\'"'))
        if not p.exists():
            raise ValueError(f"O caminho não existe: {raw}")
        if p.is_file():
            out.append(p)
        else:
            out.extend(sorted(f for f in p.iterdir()
                              if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
                              and not f.name.startswith("~$")))
    if not out:
        raise ValueError("Nenhum arquivo suportado encontrado (.xlsx, .xls, .csv, .parquet, .db...).")
    return out


def file_kind(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in EXCEL_STREAM_EXT:
        return "excel"
    if ext in EXCEL_CALAMINE_EXT:
        return "excel_calamine"
    if ext in CSV_EXT:
        return "csv"
    if ext in PARQUET_EXT:
        return "parquet"
    if ext in JSON_EXT:
        return "json"
    if ext in DB_EXT:
        with open(path, "rb") as f:
            head = f.read(16)
        if head.startswith(b"SQLite format 3"):
            return "sqlite"
        if head[8:12] == b"DUCK":
            return "duckdb"
        raise ValueError(f"{path.name}: não é um banco SQLite nem DuckDB.")
    raise ValueError(f"Extensão não suportada: {ext}")


_DIM_RE = re.compile(rb'<(?:\w+:)?dimension\s+ref="\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?"')


def _col_num(letters) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def xlsx_sheets(path) -> list:
    """Lê só o índice do .xlsx (poucos KB): nomes das abas, visibilidade e dimensões."""
    with zipfile.ZipFile(path) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        targets = {r.get("Id"): r.get("Target") for r in rels.iter() if r.get("Id")}
        names = set(z.namelist())
        out = []
        for sh in wb.iter():
            if not sh.tag.endswith("}sheet"):
                continue
            rid = next((v for k, v in sh.attrib.items() if k.endswith("}id")), None)
            target = targets.get(rid, "")
            if "worksheets/" not in target:
                continue                                # gráfico, macro etc.
            member = target.lstrip("/") if target.startswith("/") else "xl/" + target
            member = member.replace("xl/./", "xl/")
            info = {"name": sh.get("name"), "hidden": sh.get("state", "visible") != "visible",
                    "rows": None, "cols": None, "xml_size": 0, "member": member}
            if member in names:
                zi = z.getinfo(member)
                info["xml_size"] = zi.file_size
                with z.open(member) as fh:
                    m = _DIM_RE.search(fh.read(131072))
                if m and m.group(3):
                    g = [x.decode("ascii") for x in m.groups()]
                    c1, r1, c2, r2 = _col_num(g[0]), int(g[1]), _col_num(g[2]), int(g[3])
                    info["rows"] = max(r2 - r1, 0)          # sem a linha do cabeçalho
                    info["cols"] = c2 - c1 + 1
                elif m and zi.file_size < 4096:
                    info["rows"], info["cols"] = 0, 1
            out.append(info)
    return out


# --- Leitor de .xlsx em fluxo (bem mais rápido que o openpyxl, memória constante) ---
_BUILTIN_DATE_FMTS = set(range(14, 23)) | set(range(27, 37)) | set(range(45, 48)) | \
    set(range(50, 59)) | set(range(71, 82))
_FMT_CLEAN = re.compile(r'"[^"]*"|\[[^\]]*\]|\\.|_.|\*.')


def _is_date_format(code: str) -> bool:
    code = _FMT_CLEAN.sub("", code or "").lower()
    return "general" not in code and any(ch in code for ch in "dmyhs")


def _rel_targets(z, rels_member, base):
    out = {}
    if rels_member not in z.namelist():
        return out
    for r in ET.fromstring(z.read(rels_member)).iter():
        typ, target = r.get("Type") or "", r.get("Target") or ""
        if target:
            member = target.lstrip("/") if target.startswith("/") else base + target
            out[typ.rsplit("/", 1)[-1]] = member
    return out


def _xlsx_context(z):
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    date1904 = any(el.tag.endswith("}workbookPr") and el.get("date1904") in ("1", "true") for el in wb.iter())
    rels = _rel_targets(z, "xl/_rels/workbook.xml.rels", "xl/")
    strings = []
    ss = rels.get("sharedStrings", "xl/sharedStrings.xml")
    if ss in z.namelist():
        with z.open(ss) as fh:
            for _, el in ET.iterparse(fh, events=("end",)):
                if el.tag.endswith("}si"):
                    parts = []
                    for child in el:
                        if child.tag.endswith("}t"):
                            parts.append(child.text or "")
                        elif child.tag.endswith("}r"):
                            parts.extend(t.text or "" for t in child if t.tag.endswith("}t"))
                    strings.append("".join(parts))
                    el.clear()
    date_styles = set()
    st = rels.get("styles", "xl/styles.xml")
    if st in z.namelist():
        root = ET.fromstring(z.read(st))
        custom = {}
        for el in root.iter():
            if el.tag.endswith("}numFmt"):
                with contextlib.suppress(Exception):
                    custom[int(el.get("numFmtId"))] = el.get("formatCode") or ""
        for el in root:
            if el.tag.endswith("}cellXfs"):
                for i, xf in enumerate(x for x in el if x.tag.endswith("}xf")):
                    fid = int(xf.get("numFmtId") or 0)
                    if fid in _BUILTIN_DATE_FMTS or (fid in custom and _is_date_format(custom[fid])):
                        date_styles.add(str(i))
    epoch = dt.datetime(1904, 1, 1) if date1904 else dt.datetime(1899, 12, 30)
    return strings, date_styles, epoch


def _col_index(ref: str) -> int:
    n = 0
    for ch in ref:
        o = ord(ch)
        if o < 65:
            break
        n = n * 26 + (o & 31)
    return n - 1


def iter_xlsx_rows(path, sheet, min_row=1):
    """Gera as linhas de uma aba lendo o XML em fluxo."""
    with zipfile.ZipFile(path) as z:
        strings, date_styles, epoch = _xlsx_context(z)
        member = sheet.get("member") or next(
            (s["member"] for s in xlsx_sheets(path) if s["name"] == sheet["name"]), None)
        if not member:
            raise ValueError(f"Aba não encontrada: {sheet['name']}")
        with z.open(member) as fh:
            sheet_data, rown = None, 0
            for ev, el in ET.iterparse(fh, events=("start", "end")):
                tag = el.tag
                if ev == "start":
                    if tag.endswith("}sheetData"):
                        sheet_data = el
                    continue
                if not tag.endswith("}row"):
                    continue
                r = el.get("r")
                rown = int(r) if r else rown + 1
                if rown >= min_row:
                    row = []
                    for c in el:
                        if not c.tag.endswith("}c"):
                            continue
                        ref = c.get("r")
                        col = _col_index(ref) if ref else len(row)
                        t = c.get("t")
                        value = None
                        if t == "inlineStr":
                            value = "".join(x.text or "" for x in c.iter() if x.tag.endswith("}t"))
                        else:
                            vtext = None
                            for x in c:
                                if x.tag.endswith("}v"):
                                    vtext = x.text
                                    break
                            if vtext is not None:
                                if t == "s":
                                    value = strings[int(vtext)]
                                elif t in ("str", "e"):
                                    value = vtext
                                elif t == "b":
                                    value = vtext == "1"
                                elif t == "d":
                                    with contextlib.suppress(ValueError):
                                        value = dt.datetime.fromisoformat(vtext.rstrip("Z"))
                                elif c.get("s") in date_styles:
                                    f = float(vtext)
                                    moment = epoch + dt.timedelta(microseconds=round(f * 86_400_000) * 1000)
                                    value = moment.time() if 0 <= f < 1 else moment
                                elif "." in vtext or "E" in vtext or "e" in vtext:
                                    value = float(vtext)
                                else:
                                    value = int(vtext)
                        if col == len(row):
                            row.append(value)
                        elif col > len(row):
                            row.extend([None] * (col - len(row)))
                            row.append(value)
                        else:
                            row[col] = value
                    yield row
                el.clear()
                if sheet_data is not None:
                    sheet_data.clear()


def calamine_sheets(path) -> list:
    if not HAS_CALAMINE:
        raise ValueError(f"{Path(path).name}: para ler {Path(path).suffix} instale o pacote python-calamine.")
    wb = CalamineWorkbook.from_path(str(path))
    try:
        out = []
        for m in wb.sheets_metadata:
            if "Work" not in str(m.typ):
                continue
            out.append({"name": m.name, "hidden": "Visible" not in str(m.visible), "rows": None, "cols": None})
        return out
    finally:
        with contextlib.suppress(Exception):
            wb.close()


def sqlite_tables(path) -> list:
    conn = sqlite3.connect(f"file:{Path(path).as_posix()}?mode=ro", uri=True) \
        if not str(path).startswith("\\\\") else sqlite3.connect(str(path))
    try:
        out = []
        for (t,) in conn.execute("SELECT name FROM sqlite_master WHERE type IN ('table','view') "
                                 "AND name NOT LIKE 'sqlite_%' ORDER BY name"):
            cols = len(conn.execute(f'PRAGMA table_info("{t}")').fetchall())
            rows = None
            with contextlib.suppress(Exception):
                rows = conn.execute(f'SELECT max(rowid) FROM "{t}"').fetchone()[0]
            out.append({"name": t, "hidden": False, "rows": rows, "cols": cols})
        return out
    finally:
        conn.close()


def duckdb_tables(path) -> list:
    con = duckdb.connect(str(path), read_only=True)
    try:
        rows = con.execute("SELECT schema_name, table_name, estimated_size, column_count FROM duckdb_tables() "
                           "WHERE NOT internal ORDER BY schema_name, table_name").fetchall()
        return [{"name": t if s == "main" else f"{s}.{t}", "hidden": False, "rows": n, "cols": c}
                for s, t, n, c in rows]
    finally:
        con.close()


def inspect_file(path: Path, source=None, web=False) -> dict:
    info = {"path": source or str(path), "file": path.name, "stem": path.stem, "kind": None, "items": [],
            "size": 0, "remote": False if web else is_remote(path), "web": web, "error": None}
    try:
        info["size"] = path.stat().st_size
        kind = info["kind"] = file_kind(path)
        if kind == "excel":
            info["items"] = xlsx_sheets(path)
        elif kind == "excel_calamine":
            info["items"] = calamine_sheets(path)
        elif kind == "sqlite":
            info["items"] = sqlite_tables(path)
        elif kind == "duckdb":
            info["items"] = duckdb_tables(path)
        elif kind == "parquet":
            import pyarrow.parquet as pq
            md = pq.read_metadata(str(path))
            info["items"] = [{"name": path.stem, "hidden": False, "rows": md.num_rows, "cols": md.num_columns}]
        else:
            info["items"] = [{"name": path.stem, "hidden": False, "rows": None, "cols": None}]
        if not info["items"]:
            info["error"] = "Nenhuma aba/tabela com dados."
    except Exception as exc:
        info["error"] = str(exc)
    return info


def inspect_paths(paths, ctx: Ctx) -> list:
    files = resolve_inputs(paths, ctx)
    out = []
    for i, (f, key, web) in enumerate(files, 1):
        ctx.check()
        ctx.progress(f"Analisando {f.name} ({i}/{len(files)})…", 100 * i / len(files), force=True)
        out.append(inspect_file(f, key, web))
    return out


def dataset_name(stem, item, n_items) -> str:
    return stem if n_items == 1 else f"{stem}__{item}"


# ============================================================================
# INGESTÃO EM BLOCOS (a RAM usada é limitada ao tamanho do bloco)
# ============================================================================
_INT_TYPES = {"TINYINT", "SMALLINT", "INTEGER", "BIGINT", "HUGEINT", "UTINYINT", "USMALLINT",
              "UINTEGER", "UBIGINT", "UHUGEINT"}


def is_int_type(t):
    return t in _INT_TYPES


def is_numeric_type(t):
    return t in _INT_TYPES or t in ("FLOAT", "DOUBLE", "REAL") or t.startswith("DECIMAL")


def is_temporal_type(t):
    return t == "DATE" or t.startswith("TIMESTAMP")


def unify_types(a, b) -> str:
    if a == b:
        return a
    if is_int_type(a) and is_int_type(b):
        return "HUGEINT" if "HUGEINT" in (a, b) else "BIGINT"
    if is_numeric_type(a) and is_numeric_type(b):
        return "DOUBLE"
    if is_temporal_type(a) and is_temporal_type(b):
        return "TIMESTAMP"
    return "VARCHAR"


def chunk_rows(ncols, budget) -> int:
    return int(min(max(budget * 0.08 / max(ncols, 1) / 80, 2000), 100_000))


def text_value(v):
    """Regra única para virar texto (a mesma usada no código Python gerado):
    1.0 → '1', datas → 'AAAA-MM-DD HH:MM:SS', vazio → nulo."""
    if v is None or (isinstance(v, float) and v != v) or v is pd.NaT:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d 00:00:00")
    return str(v)


def prepare_chunk(df: pd.DataFrame) -> pd.DataFrame:
    """Normaliza tipos de um bloco para o DuckDB aceitar sem surpresas."""
    for c in df.columns:
        s = df[c]
        if s.dtype.kind == "M":
            if getattr(s.dt, "tz", None) is not None:
                s = s.dt.tz_localize(None)
            df[c] = s.astype("datetime64[us]")
            continue
        if s.dtype.kind == "f":
            nn = s.dropna()
            if len(nn) and (nn % 1 == 0).all() and nn.abs().max() < 2**53:
                df[c] = s.astype("Int64")
            continue
        if s.dtype != object:
            continue
        kind = pd.api.types.infer_dtype(s, skipna=True)
        if kind in ("empty", "string"):
            continue
        if kind in ("integer", "floating", "mixed-integer-float", "decimal"):
            num = pd.to_numeric(s, errors="coerce")
            nn = num.dropna()
            if num.dtype.kind == "f" and len(nn) and (nn % 1 == 0).all() and nn.abs().max() < 2**53:
                num = num.astype("Int64")
            df[c] = num
        elif kind == "boolean":
            df[c] = s.astype("boolean")
        elif kind in ("datetime", "datetime64", "date"):
            df[c] = pd.to_datetime(s, errors="coerce").astype("datetime64[us]")
        else:   # tipos misturados (ex.: números e textos na mesma coluna) → texto
            df[c] = s.map(text_value)
    return df


class TableWriter:
    """Acrescenta blocos de DataFrame numa tabela DuckDB, ajustando os tipos quando um bloco
    novo traz um tipo diferente (ex.: coluna que era número e passou a ter texto)."""

    def __init__(self, con, tid):
        self.con, self.tid = con, tid
        self.created = False
        self.types = {}
        self.unknown = set()      # colunas só com nulos até agora
        self.rows = 0

    def append(self, df: pd.DataFrame):
        if df.empty:
            return
        view = f"_bloco_{self.tid}"
        self.con.register(view, df)
        try:
            chunk_types = {r[0]: r[1] for r in self.con.execute(f"DESCRIBE {view}").fetchall()}
            has_data = {c: bool(df[c].notna().any()) for c in df.columns}
            if not self.created:
                self.con.execute(f"CREATE TABLE {self.tid} AS SELECT * FROM {view}")
                self.created = True
                self.types = dict(chunk_types)
                self.unknown = {c for c, ok in has_data.items() if not ok}
            else:
                for c, t in chunk_types.items():
                    if c not in self.types:
                        self.con.execute(f"ALTER TABLE {self.tid} ADD COLUMN {qi(c)} {t}")
                        self.types[c] = t
                        if not has_data[c]:
                            self.unknown.add(c)
                        continue
                    if not has_data[c]:
                        continue
                    if c in self.unknown:
                        new = t
                        self.unknown.discard(c)
                    else:
                        new = unify_types(self.types[c], t)
                    if new != self.types[c]:
                        self.con.execute(f"ALTER TABLE {self.tid} ALTER COLUMN {qi(c)} TYPE {new}")
                        self.types[c] = new
                self.con.execute(f"INSERT INTO {self.tid} BY NAME SELECT * FROM {view}")
            self.rows += len(df)
        finally:
            self.con.unregister(view)

    def finish(self, names, generated=()):
        if not self.created:
            cols = ", ".join(f"{qi(n)} VARCHAR" for n in (names or ["coluna_1"]))
            self.con.execute(f"CREATE TABLE {self.tid} ({cols})")
            return
        # colunas sem cabeçalho e sem nenhum valor (formatação sobrando no Excel) saem
        for c in generated:
            if c in self.types and len(self.types) > 1:
                if self.con.execute(f"SELECT count({qi(c)}) FROM {self.tid}").fetchone()[0] == 0:
                    self.con.execute(f"ALTER TABLE {self.tid} DROP COLUMN {qi(c)}")
                    self.types.pop(c)
                    self.unknown.discard(c)
        for c in list(self.unknown):
            if c in self.types and self.types[c] != "VARCHAR":
                self.con.execute(f"ALTER TABLE {self.tid} ALTER COLUMN {qi(c)} TYPE VARCHAR")


def _header_names(header):
    names, generated, seen = [], [], set()
    for i, v in enumerate(header):
        if v is None or str(v).strip() == "":
            name = f"coluna_{i + 1}"
            generated.append(name)
        elif isinstance(v, float) and v.is_integer():
            name = str(int(v))
        elif isinstance(v, (dt.datetime, dt.date)):
            name = v.strftime("%d/%m/%Y")
        else:
            name = " ".join(str(v).split())
        base, k = name, 2
        while name.lower() in seen:          # DuckDB não diferencia maiúsculas em nomes de coluna
            name = f"{base}_{k}"
            k += 1
        if base in generated and name != base:
            generated[generated.index(base)] = name
        seen.add(name.lower())
        names.append(name)
    return names, generated


def ingest_rows(con, tid, rows, label, est_rows, budget, ctx, blank_is_none=False) -> int:
    """Lê linhas (o 1º registro não vazio é o cabeçalho) e grava em blocos."""
    header = None
    for row in rows:
        vals = [None if (blank_is_none and v == "") else v for v in row]
        if any(v is not None and str(v).strip() != "" for v in vals):
            header = vals
            break
    names, generated = _header_names(header or [])
    width = len(names)
    writer = TableWriter(con, tid)
    size = chunk_rows(width, budget)
    buf, n = [], 0

    def flush():
        nonlocal buf
        if not buf:
            return
        fixed = [r + [None] * (width - len(r)) if len(r) < width else r for r in buf]
        writer.append(prepare_chunk(pd.DataFrame(fixed, columns=names)))
        buf = []

    for row in rows:
        r = [None if v == "" else v for v in row] if blank_is_none else list(row)
        if len(r) > width:
            extra = r[width:]
            if any(v is not None for v in extra):
                last = max(i for i, v in enumerate(extra) if v is not None)
                for j in range(width, width + last + 1):
                    nm = f"coluna_{j + 1}"
                    while nm.lower() in {x.lower() for x in names}:
                        nm += "_"
                    names.append(nm)
                    generated.append(nm)
                width = len(names)
            r = r[:width]
        if all(v is None for v in r):
            continue
        buf.append(r)
        n += 1
        if len(buf) >= size:
            ctx.check()
            flush()
            pct = 100 * n / est_rows if est_rows else None
            ctx.progress(f"Lendo {label}: {fmt_int(n)} linhas", pct)
    ctx.check()
    flush()
    writer.finish(names, generated)
    return n


def _ingest_excel_sheet(con, tid, path, kind, sheet, header_row, budget, ctx, label):
    est = sheet.get("rows")
    # calamine é o mais rápido, mas carrega a aba inteira (~2,5× o XML); só usa se couber com folga
    use_calamine = HAS_CALAMINE and (
        kind == "excel_calamine" or sheet.get("xml_size", 0) * 2.5 < budget * 0.6)
    if kind == "excel_calamine" and not HAS_CALAMINE:
        raise ValueError("Para ler este formato instale o pacote python-calamine.")
    if use_calamine:       # rápido; a aba cabe com folga na memória
        wb = CalamineWorkbook.from_path(str(path))
        try:
            sh = wb.get_sheet_by_name(sheet["name"])
            est = est or sh.height
            rows = iter(sh.iter_rows())
            for _ in range(max(header_row - 1, 0)):
                next(rows, None)
            ingest_rows(con, tid, rows, label, est, budget, ctx, blank_is_none=True)
        finally:
            with contextlib.suppress(Exception):
                wb.close()
    else:                  # em fluxo: lê linha a linha, memória constante
        try:
            rows = iter_xlsx_rows(path, sheet, max(header_row, 1))
            ingest_rows(con, tid, rows, label, est, budget, ctx)
        except (Cancelled, duckdb.InterruptException):
            raise
        except Exception:      # arquivo fora do padrão: tenta pelo openpyxl (mais lento)
            import openpyxl
            con.execute(f"DROP TABLE IF EXISTS {tid}")
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            try:
                ws = wb[sheet["name"]]
                ws.reset_dimensions()
                rows = ws.iter_rows(min_row=max(header_row, 1), values_only=True)
                ingest_rows(con, tid, rows, label, est, budget, ctx)
            finally:
                wb.close()


def _ingest_csv(con, tid, path, header_row, ctx):
    skip = max(header_row - 1, 0)
    attempts = [("utf-8", False), ("latin-1", False), ("utf-8", True), ("latin-1", True)]
    last = None
    for enc, all_text in attempts:
        ctx.check()
        try:
            opts = f"header=true, skip={skip}, encoding={ql(enc)}, sample_size=100000"
            if all_text:
                opts += ", all_varchar=true"
            con.execute(f"CREATE OR REPLACE TABLE {tid} AS SELECT * FROM read_csv({ql(path)}, {opts})")
            return
        except duckdb.InterruptException:
            raise
        except Exception as exc:
            last = exc
    raise ValueError(f"Não foi possível ler o CSV: {last}")


def _ingest_sqlite(con, tid, path, table, est, budget, ctx, label):
    conn = sqlite3.connect(str(path))
    try:
        ncols = len(conn.execute(f'PRAGMA table_info("{table}")').fetchall())
        writer = TableWriter(con, tid)
        n = 0
        for chunk in pd.read_sql_query(f'SELECT * FROM "{table}"', conn, chunksize=chunk_rows(ncols, budget)):
            ctx.check()
            chunk.columns = _header_names(list(chunk.columns))[0]
            writer.append(prepare_chunk(chunk))
            n += len(chunk)
            ctx.progress(f"Lendo {label}: {fmt_int(n)} linhas", 100 * n / est if est else None)
        writer.finish([str(c[1]) for c in conn.execute(f'PRAGMA table_info("{table}")').fetchall()])
    finally:
        conn.close()


def _ingest_duckdb(con, tid, path, table, ctx):
    alias = f"origem_{tid}"
    con.execute(f"ATTACH {ql(path)} AS {alias} (READ_ONLY)")
    try:
        schema, _, name = table.rpartition(".")
        ref = f"{alias}.{qi(schema or 'main')}.{qi(name)}"
        con.execute(f"CREATE TABLE {tid} AS SELECT * FROM {ref}")
    finally:
        con.execute(f"DETACH {alias}")


def _ingest_json(con, tid, path):
    """JSON vira tabela. Se for o formato comum de API ({"data": [ {...}, ... ]}), abre a lista."""
    try:
        con.execute(f"CREATE TABLE {tid} AS SELECT * FROM read_json_auto({ql(path)}, maximum_object_size=268435456)")
    except duckdb.Error as exc:
        raise ValueError(f"Não foi possível ler o JSON: {exc}")
    rows = con.execute(f"SELECT count(*) FROM {tid}").fetchone()[0]
    lists = [(c, t) for c, t, *_ in con.execute(f"DESCRIBE {tid}").fetchall()
             if t.startswith("STRUCT(") and t.endswith("[]")]
    if rows == 1 and len(lists) == 1:
        col = lists[0][0]
        tmp = f"{tid}_lista"
        con.execute(f"CREATE TABLE {tmp} AS SELECT u.* FROM (SELECT unnest({qi(col)}) AS u FROM {tid})")
        con.execute(f"DROP TABLE {tid}")
        con.execute(f"ALTER TABLE {tmp} RENAME TO {tid}")


def load_files(engine: Engine, con, params, ctx) -> list:
    header_row = int(params.get("header_row") or 1)
    select = {os.path.normcase(k): v for k, v in (params.get("select") or {}).items()}
    headers = {os.path.normcase(k): v for k, v in (params.get("headers") or {}).items()}
    plan = []
    for f, key, _web in resolve_inputs(params["paths"], ctx):
        chosen = select.get(os.path.normcase(key))
        if chosen is not None and not chosen:
            continue
        plan.append((f, key, chosen))
    if not plan:
        raise ValueError("Nenhuma aba ou arquivo foi selecionado.")

    local = stage_files([f for f, _, _ in plan], ctx, engine.cache_gb)
    results = []
    for f, key, chosen in plan:
        ctx.check()
        src = local[f]
        ctx.progress(f"Abrindo {f.name}…", None, force=True)
        info = inspect_file(Path(src))
        if info["error"]:
            raise ValueError(f"{f.name}: {info['error']}")
        items = info["items"]
        by_name = {it["name"]: it for it in items}
        names = chosen if chosen is not None else [it["name"] for it in items]
        missing = [n for n in names if n not in by_name]
        if missing:
            raise ValueError(f"{f.name}: aba/tabela não encontrada: {', '.join(missing)}")
        file_headers = headers.get(os.path.normcase(key), {})
        for item_name in names:
            item = by_name[item_name]
            hdr = int(file_headers.get(item_name, header_row))
            tid = engine.new_tid()
            label = f.name if len(items) == 1 else f"{f.name} › {item_name}"
            kind = info["kind"]
            if kind in ("excel", "excel_calamine"):
                _ingest_excel_sheet(con, tid, src, kind, item, hdr, engine.budget, ctx, label)
            elif kind == "csv":
                ctx.progress(f"Lendo {label}…", None, force=True)
                _ingest_csv(con, tid, str(src), hdr, ctx)
            elif kind == "parquet":
                ctx.progress(f"Lendo {label}…", None, force=True)
                con.execute(f"CREATE TABLE {tid} AS SELECT * FROM read_parquet({ql(str(src))})")
            elif kind == "json":
                ctx.progress(f"Lendo {label}…", None, force=True)
                _ingest_json(con, tid, str(src))
            elif kind == "sqlite":
                _ingest_sqlite(con, tid, src, item_name, item.get("rows"), engine.budget, ctx, label)
            elif kind == "duckdb":
                ctx.progress(f"Lendo {label}…", None, force=True)
                _ingest_duckdb(con, tid, str(src), item_name, ctx)
            engine.lineage[tid] = {"kind": kind, "source": key, "item": item_name, "header": hdr,
                                   "n_items": len(items)}
            name = dataset_name(f.stem, item_name, len(items))
            base, k, used = name, 2, {n for _, n, _ in results}
            while name in used:          # dois arquivos com o mesmo nome na mesma carga
                name = f"{base}_{k}"
                k += 1
            results.append(("bronze", name, tid))
    return results


def load_sample(engine, con) -> list:
    clientes = pd.DataFrame({"id_cliente": [1, 2, 3, 4], "nome": ["Ana", "Bruno", "Carla", "Diego"],
                             "cidade": ["SP", "RJ", "BH", "SP"]})
    pedidos = pd.DataFrame({"id_pedido": [101, 102, 103, 104], "id_cliente": [1, 1, 2, 5],
                            "valor": [100.0, 250.0, 80.0, 300.0]})
    out = []
    for name, df in (("clientes", clientes), ("pedidos", pedidos)):
        tid = engine.new_tid()
        w = TableWriter(con, tid)
        w.append(df)
        w.finish(list(df.columns))
        engine.lineage[tid] = {"kind": "sample"}
        out.append(("bronze", name, tid))
    return out


# ============================================================================
# TRANSFORMAÇÕES EM SQL (Prata e Ouro) — rodam no DuckDB, sem carregar na RAM
# ============================================================================
def parse_number(value):
    s = str(value).strip()
    for cand in (s, s.replace(".", "").replace(",", ".")):
        try:
            return float(cand)
        except ValueError:
            pass
    return None


def parse_date(value):
    s = str(value).strip()
    for fmt, has_time in (("%d/%m/%Y %H:%M:%S", True), ("%d/%m/%Y %H:%M", True), ("%d/%m/%Y", False),
                          ("%Y-%m-%d %H:%M:%S", True), ("%Y-%m-%d", False)):
        try:
            return dt.datetime.strptime(s, fmt), has_time
        except ValueError:
            pass
    return None, False


def num_expr(col, typ) -> str:
    """Converte qualquer coluna em número; aceita '1.234,56' e '1,5'."""
    q = qi(col)
    if is_numeric_type(typ):
        return f"CAST({q} AS DOUBLE)"
    if typ == "BOOLEAN":
        return f"CAST(CAST({q} AS INTEGER) AS DOUBLE)"
    t = f"TRIM(CAST({q} AS VARCHAR))"
    return f"COALESCE(TRY_CAST({t} AS DOUBLE), TRY_CAST(REPLACE(REPLACE({t}, '.', ''), ',', '.') AS DOUBLE))"


_TRUE = ("true", "1", "sim", "s", "yes", "y", "verdadeiro", "x", "v")
_FALSE = ("false", "0", "não", "nao", "n", "no", "falso", "f")


def cast_expr(col, typ, target) -> str:
    q = qi(col)
    if target == "int":
        return f"TRY_CAST(ROUND({num_expr(col, typ)}) AS BIGINT)"
    if target == "float":
        return num_expr(col, typ)
    if target == "str":
        return f"CAST({q} AS VARCHAR)"
    if target == "datetime":
        if is_temporal_type(typ):
            return f"CAST({q} AS TIMESTAMP)"
        if is_numeric_type(typ):   # data serial do Excel (ex.: 45292 = 01/01/2024)
            return (f"CASE WHEN {q} BETWEEN 1 AND 2958465 THEN TIMESTAMP '1899-12-30' + "
                    f"to_microseconds(CAST(ROUND({q} * 86400000000) AS BIGINT)) END")
        t = f"TRIM(CAST({q} AS VARCHAR))"
        fmts = ["%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"]
        parts = [f"TRY_CAST({t} AS TIMESTAMP)"] + [f"try_strptime({t}, {ql(f)})" for f in fmts]
        return f"COALESCE({', '.join(parts)})"
    if target == "bool":
        if typ == "BOOLEAN":
            return q
        if is_numeric_type(typ):
            return f"CASE WHEN {q} IS NULL THEN NULL ELSE {q} <> 0 END"
        t = f"lower(TRIM(CAST({q} AS VARCHAR)))"
        yes = ", ".join(ql(v) for v in _TRUE)
        no = ", ".join(ql(v) for v in _FALSE)
        return f"CASE WHEN {t} IN ({yes}) THEN TRUE WHEN {t} IN ({no}) THEN FALSE END"
    raise ValueError(f"Tipo desconhecido: {target}")


def sql_dedup(engine, con, src, subset):
    cols = subset or engine.columns(src, con)
    keys = ", ".join(qi(c) for c in cols)
    return (f"SELECT * FROM {src} WHERE rowid IN (SELECT min(rowid) FROM {src} GROUP BY {keys}) "
            f"ORDER BY rowid")


def sql_nulls(engine, con, src, fill, dropna):
    types = engine.types(src, con)
    sel = "*"
    if fill not in (None, ""):
        num = parse_number(fill)
        parts = []
        for c, t in types.items():
            if is_numeric_type(t):
                lit = "NULL" if num is None else repr(num)
                parts.append(f"COALESCE({qi(c)}, TRY_CAST({lit} AS {t})) AS {qi(c)}")
            else:
                parts.append(f"COALESCE({qi(c)}, TRY_CAST({ql(fill)} AS {t})) AS {qi(c)}")
        sel = ", ".join(parts)
    sql = f"SELECT {sel} FROM {src}"
    if dropna:
        cond = " AND ".join(f"{qi(c)} IS NOT NULL" for c in types)
        sql = f"SELECT * FROM ({sql}) WHERE {cond}"
    return sql


def sql_cast(engine, con, src, column, target):
    typ = engine.types(src, con)[column]
    return f"SELECT * REPLACE ({cast_expr(column, typ, target)} AS {qi(column)}) FROM {src}"


def sql_merge(engine, con, a, b, left_on, right_on, how):
    how = how if how in JOIN_TYPES else "left"
    ca, cb = engine.columns(a, con), engine.columns(b, con)
    la, lb = {c.lower() for c in ca}, {c.lower() for c in cb}
    same_key = how != "cross" and left_on.lower() == right_on.lower()
    overlap = la & lb
    if same_key:
        overlap -= {left_on.lower()}
    sel, used = [], set()

    def add(expr, name):
        base, k = name, 2
        while name.lower() in used:
            name = f"{base}_{k}"
            k += 1
        used.add(name.lower())
        sel.append(f"{expr} AS {qi(name)}")

    for c in ca:
        if same_key and c.lower() == left_on.lower():
            expr = f"COALESCE(A.{qi(c)}, B.{qi(right_on)})" if how == "full" else f"A.{qi(c)}"
            add(expr, c)
        else:
            add(f"A.{qi(c)}", f"{c}_A" if c.lower() in overlap else c)
    for c in cb:
        if same_key and c.lower() == right_on.lower():
            continue
        add(f"B.{qi(c)}", f"{c}_B" if c.lower() in overlap else c)

    sql = f"SELECT {', '.join(sel)} FROM {a} AS A {JOIN_TYPES[how]} {b} AS B"
    if how != "cross":
        sql += f" ON {join_condition(engine, con, a, b, left_on, right_on)}"
    return sql + " ORDER BY A.rowid NULLS LAST, B.rowid"


def join_condition(engine, con, a, b, left_on, right_on) -> str:
    """Condição de junção A×B que tolera tipos diferentes (número × texto etc.)."""
    ta, tb = engine.types(a, con)[left_on], engine.types(b, con)[right_on]
    ka, kb = f"A.{qi(left_on)}", f"B.{qi(right_on)}"
    if ta == tb or (is_numeric_type(ta) and is_numeric_type(tb)) or \
            (is_temporal_type(ta) and is_temporal_type(tb)):
        return f"{ka} = {kb}"
    if is_numeric_type(ta):
        return f"CAST({ka} AS DOUBLE) = {num_expr(right_on, tb).replace(qi(right_on), kb)}"
    if is_numeric_type(tb):
        return f"{num_expr(left_on, ta).replace(qi(left_on), ka)} = CAST({kb} AS DOUBLE)"
    return f"TRIM(CAST({ka} AS VARCHAR)) = TRIM(CAST({kb} AS VARCHAR))"


def merge_check(engine, a, b, left_on, right_on, how, ctx) -> dict:
    """Diagnóstico antes do merge: quanto de A encontra par em B e se B multiplica linhas."""
    con = engine.cursor()
    try:
        with ctx.watch(con):
            ctx.progress("Verificando as chaves…", None, force=True)
            cond = join_condition(engine, con, a, b, left_on, right_on)
            ka, kb = qi(left_on), qi(right_on)
            rows_a, null_a = con.execute(f"SELECT count(*), count(*) - count({ka}) FROM {a}").fetchone()
            rows_b, keys_b = con.execute(f"SELECT count(*), count(DISTINCT {kb}) FROM {b}").fetchone()
            dup_b = con.execute(f"SELECT count(*) FROM (SELECT {kb} FROM {b} WHERE {kb} IS NOT NULL "
                                f"GROUP BY {kb} HAVING count(*) > 1)").fetchone()[0]
            matched = con.execute(f"SELECT count(*) FROM {a} AS A WHERE EXISTS "
                                  f"(SELECT 1 FROM {b} AS B WHERE {cond})").fetchone()[0]
            result = con.execute(f"SELECT count(*) FROM {a} AS A {JOIN_TYPES[how]} {b} AS B "
                                 f"ON {cond}").fetchone()[0]
    except duckdb.InterruptException:
        raise Cancelled()
    finally:
        con.close()
    return {"rows_a": rows_a, "null_a": null_a, "rows_b": rows_b, "keys_b": keys_b, "dup_b": dup_b,
            "matched": matched, "result": result, "how": how}


def sql_filter(engine, con, src, column, operator, value):
    typ = engine.types(src, con)[column]
    q = qi(column)
    if operator == "não é nulo":
        cond = f"{q} IS NOT NULL"
    elif operator == "é nulo":
        cond = f"{q} IS NULL"
    elif operator == "contém":
        cond = f"contains(lower(CAST({q} AS VARCHAR)), lower({ql(value)}))"
    elif operator == "começa com":
        cond = f"starts_with(lower(TRIM(CAST({q} AS VARCHAR))), lower({ql(str(value).strip())}))"
    elif operator == "termina com":
        cond = f"suffix(lower(TRIM(CAST({q} AS VARCHAR))), lower({ql(str(value).strip())}))"
    elif operator in LIST_OPERATORS:
        vals = [v.strip() for v in re.split(r"[;\n]", str(value)) if v.strip()]
        if not vals:
            raise ValueError("Informe os valores da lista separados por ponto e vírgula (;).")
        nums = [parse_number(v) for v in vals]
        if is_numeric_type(typ) and all(n is not None for n in nums):
            cond = f"{num_expr(column, typ)} IN ({', '.join(repr(n) for n in nums)})"
        else:
            cond = f"lower(TRIM(CAST({q} AS VARCHAR))) IN ({', '.join(ql(v.lower()) for v in vals)})"
        if operator == "não está na lista":
            cond = f"NOT COALESCE({cond}, FALSE)"
    else:
        sym = {"igual a": "=", "diferente de": "IS DISTINCT FROM", "maior que": ">", "menor que": "<",
               "maior ou igual": ">=", "menor ou igual": "<="}[operator]
        when, has_time = parse_date(value) if is_temporal_type(typ) else (None, False)
        if when is not None:
            if has_time:
                left, right = f"CAST({q} AS TIMESTAMP)", f"TIMESTAMP {ql(when.isoformat(sep=' '))}"
            else:
                left, right = f"CAST({q} AS DATE)", f"DATE {ql(when.date().isoformat())}"
        elif operator in NUMERIC_OPERATORS or is_numeric_type(typ):
            n = parse_number(value)
            if n is None:
                if operator in NUMERIC_OPERATORS:
                    raise ValueError(f"O operador '{operator}' precisa de um valor numérico ou de uma data.")
                left, right = f"CAST({q} AS VARCHAR)", ql(value)
            else:
                left, right = num_expr(column, typ), repr(n)
        else:
            left, right = f"CAST({q} AS VARCHAR)", ql(value)
        cond = f"{left} {sym} {right}"
    return f"SELECT * FROM {src} WHERE {cond}"


def sql_text_clean(engine, con, src, columns, trim=True, spaces=True, accents=False, case="", empty_null=True):
    """Padroniza textos: espaços (inclusive o espaço 'invisível' do Excel), acentos e maiúsculas."""
    types = engine.types(src, con)
    parts = []
    for c in columns:
        if c not in types:
            raise ValueError(f"A coluna '{c}' não existe em {src}.")
        e = f"CAST({qi(c)} AS VARCHAR)"
        if spaces:
            e = f"regexp_replace({e}, '[\\s\\x{{00A0}}]+', ' ', 'g')"
        if trim:
            e = f"regexp_replace({e}, '^[\\s\\x{{00A0}}]+|[\\s\\x{{00A0}}]+$', '', 'g')"
        if accents:
            e = f"strip_accents({e})"
        if case in ("upper", "lower"):
            e = f"{case}({e})"
        if empty_null:
            e = f"NULLIF({e}, '')"
        parts.append(f"{e} AS {qi(c)}")
    if not parts:
        raise ValueError("Escolha ao menos uma coluna de texto.")
    return f"SELECT * REPLACE ({', '.join(parts)}) FROM {src}"


def profile_table(engine, tid, ctx) -> pd.DataFrame:
    """Perfil das colunas (nulos, distintos, mínimo, máximo, média) via SUMMARIZE do DuckDB."""
    cached = engine._profile.get(tid)
    if cached is not None:
        return cached
    con = engine.cursor()
    try:
        with ctx.watch(con):
            ctx.progress("Calculando o perfil das colunas…", None, force=True)
            rows = con.execute(f"SUMMARIZE {tid}").fetchall()
            names = [d[0] for d in con.description]
    except duckdb.InterruptException:
        raise Cancelled()
    finally:
        con.close()

    def short(v, n=60):
        if v is None:
            return ""
        txt = str(v)
        return txt if len(txt) <= n else txt[:n - 1] + "…"

    def num(v):
        try:
            f = float(v)
        except (TypeError, ValueError):
            return short(v)
        return f"{f:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    out = []
    for r in rows:
        d = dict(zip(names, r))
        nulls = float(d.get("null_percentage") or 0)
        typ = d.get("column_type") or ""
        out.append({"Coluna": d["column_name"], "Tipo": typ,
                    "% vazio": f"{nulls:.1f}%".replace(".", ","),
                    "Distintos (aprox.)": fmt_int(d.get("approx_unique") or 0),
                    "Mínimo": short(d.get("min")), "Máximo": short(d.get("max")),
                    "Média": num(d.get("avg")) if is_numeric_type(typ) else ""})
    df = pd.DataFrame(out)
    engine._profile[tid] = df
    return df


def sql_aggregate(engine, con, src, group_by, columns, func):
    f = {"sum": "sum({c})", "mean": "avg({c})", "count": "count({c})", "min": "min({c})", "max": "max({c})",
         "nunique": "count(DISTINCT {c})",
         "first": "first({c} ORDER BY rowid) FILTER (WHERE {c} IS NOT NULL)",
         "last": "last({c} ORDER BY rowid) FILTER (WHERE {c} IS NOT NULL)"}[func]
    aggs = ", ".join(f"{f.format(c=qi(c))} AS {qi(c)}" for c in columns)
    if not group_by:
        return f"SELECT {aggs} FROM {src}"
    g = ", ".join(qi(c) for c in group_by)
    order = ", ".join(f"{qi(c)} NULLS LAST" for c in group_by)
    return f"SELECT {g}, {aggs} FROM {src} GROUP BY {g} ORDER BY {order}"


def sql_compute(engine, con, src, col_a, operation, col_b, new_name):
    types = engine.types(src, con)
    a, b = num_expr(col_a, types[col_a]), num_expr(col_b, types[col_b])
    expr = f"{a} / NULLIF({b}, 0)" if operation == "/" else f"{a} {operation} {b}"
    existing = next((c for c in types if c.lower() == new_name.lower()), None)
    if existing:
        return f"SELECT * REPLACE ({expr} AS {qi(existing)}) FROM {src}"
    return f"SELECT *, {expr} AS {qi(new_name)} FROM {src}"


# ============================================================================
# BANCO DUCKDB EXTERNO: criar tabelas, gravar datasets, rodar SQL
# ============================================================================
DB_TYPES = ["TEXT", "INTEGER", "BIGINT", "DOUBLE", "DECIMAL(18,2)", "BOOLEAN", "DATE", "DATETIME"]
DB_WRITE_MODES = {
    "append": "Acrescentar linhas (INSERT)",
    "truncate": "Limpar a tabela e inserir (DELETE + INSERT)",
    "replace": "Recriar a tabela com as colunas do dataset (CREATE OR REPLACE)",
}
_TYPE_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_ ]*(\(\s*\d+\s*(,\s*\d+\s*)?\))?(\[\])?$")


def db_type_for(duck_type: str) -> str:
    """Tipo do DuckDB → tipo no estilo do BD.py (TEXT, DATETIME…)."""
    t = duck_type.upper()
    if t == "VARCHAR":
        return "TEXT"
    if t.startswith("TIMESTAMP"):
        return "DATETIME"
    if t in ("TINYINT", "SMALLINT", "INTEGER", "BIGINT", "HUGEINT", "UTINYINT", "USMALLINT", "UINTEGER", "UBIGINT"):
        return "BIGINT"
    if t in ("FLOAT", "DOUBLE", "REAL"):
        return "DOUBLE"
    if t.startswith("DECIMAL") or t in ("DATE", "BOOLEAN", "TIME"):
        return t
    return "TEXT"


def db_norm(path) -> str:
    p = clean_input(path)
    if not p:
        raise ValueError("Informe o caminho do banco DuckDB (.duckdb).")
    return str(Path(p).expanduser().resolve())


def db_check_file(path: str):
    p = Path(path)
    if p.exists() and p.is_file() and p.stat().st_size > 0:
        with open(p, "rb") as fh:
            head = fh.read(16)
        if head.startswith(b"SQLite format 3"):
            raise ValueError(f"{p.name} é um banco SQLite, não DuckDB. Para ler, use a tela Upload.")
        if head[8:12] != b"DUCK":
            raise ValueError(f"{p.name} não é um banco DuckDB.")


def friendly_db_error(exc, path) -> ValueError:
    msg, name = str(exc), Path(path).name
    if "lock" in msg.lower() or "being used by another process" in msg.lower():
        return ValueError(f"O banco {name} está aberto em outro programa (DBeaver, Power BI ou um script Python). "
                          "Feche-o lá e tente de novo.")
    return ValueError(f"Não consegui abrir {name}: {msg}")


@contextlib.contextmanager
def external_db(engine, con, path):
    """Disponibiliza o banco externo com um apelido. Reaproveita a conexão aberta na tela, se for o mesmo."""
    p = db_norm(path)
    if engine.attached and os.path.normcase(engine.attached) == os.path.normcase(p):
        yield "bd"
        return
    db_check_file(p)
    Path(p).parent.mkdir(parents=True, exist_ok=True)
    alias = f"ext_{engine.new_tid()[2:]}"
    try:
        con.execute(f"ATTACH {ql(p)} AS {alias}")
    except duckdb.Error as exc:
        raise friendly_db_error(exc, p)
    try:
        yield alias
    finally:
        with contextlib.suppress(Exception):
            con.execute(f"DETACH {alias}")


def db_tables(con, alias) -> list:
    return [r[0] for r in con.execute(
        "SELECT table_name FROM duckdb_tables() WHERE database_name = ? AND schema_name = 'main' "
        "ORDER BY table_name", [alias]).fetchall()]


def db_ref(alias, table) -> str:
    return f"{alias}.main.{qi(table)}"


def db_create(engine, con, p):
    cols = p.get("columns") or []
    if not str(p.get("table") or "").strip():
        raise ValueError("Dê um nome à tabela.")
    if not cols:
        raise ValueError("A tabela precisa de pelo menos uma coluna.")
    parts = []
    for c in cols:
        name, typ = str(c["name"]).strip(), str(c["type"]).strip().upper()
        if not name:
            raise ValueError("Há uma coluna sem nome.")
        if not _TYPE_RE.match(typ):
            raise ValueError(f"Tipo inválido na coluna '{name}': {typ}")
        parts.append(f"{qi(name)} {typ}")
    with external_db(engine, con, p["db"]) as a:
        con.execute(f"CREATE TABLE IF NOT EXISTS {db_ref(a, p['table'].strip())} ({', '.join(parts)})")


def db_write(engine, con, p, src) -> dict:
    mode, table = p.get("mode", "append"), str(p["table"]).strip()
    if mode not in DB_WRITE_MODES:
        raise ValueError(f"Modo de gravação desconhecido: {mode}")
    scols = engine.columns(src, con)
    n = con.execute(f"SELECT count(*) FROM {src}").fetchone()[0]
    with external_db(engine, con, p["db"]) as a:
        ref = db_ref(a, table)
        exists = table.lower() in {t.lower() for t in db_tables(con, a)}
        if mode == "replace" or not exists:
            con.execute(f"CREATE OR REPLACE TABLE {ref} AS SELECT * FROM {src}")
            return {"table": table, "rows": n, "created": True, "common": scols, "ignored": [], "missing": []}
        tcols = [(r[0], r[1]) for r in con.execute(f"DESCRIBE {ref}").fetchall()]
        smap = {c.lower(): c for c in scols}
        tlow = {t.lower() for t, _ in tcols}
        common = [(t, ty, smap[t.lower()]) for t, ty in tcols if t.lower() in smap]
        if not common:
            raise ValueError(f"Nenhuma coluna do dataset '{p.get('src')}' existe na tabela {table}. "
                             f"Colunas da tabela: {', '.join(t for t, _ in tcols)}")
        ignored = [c for c in scols if c.lower() not in tlow]
        missing = [t for t, _ in tcols if t.lower() not in smap]
        dest = ", ".join(qi(t) for t, _, _ in common)
        sel = ", ".join((f"TRY_CAST({qi(s)} AS {ty})" if p.get("try_cast") else qi(s)) for _, ty, s in common)
        con.execute("BEGIN TRANSACTION")
        try:
            if mode == "truncate":
                con.execute(f"DELETE FROM {ref}")
            con.execute(f"INSERT INTO {ref} ({dest}) SELECT {sel} FROM {src}")
            con.execute("COMMIT")
        except duckdb.ConversionException as exc:
            con.execute("ROLLBACK")
            raise ValueError(f"Alguns valores não cabem no tipo das colunas da tabela {table}: {exc}\n\n"
                             "Nada foi gravado. Marque 'Valores incompatíveis viram vazio' para gravar mesmo assim, "
                             "ou ajuste o tipo na tela Prata.")
        except Exception:
            with contextlib.suppress(Exception):
                con.execute("ROLLBACK")
            raise
        return {"table": table, "rows": n, "created": False, "common": [t for t, _, _ in common],
                "ignored": ignored, "missing": missing}


def db_sql(engine, con, p):
    """Roda SQL no banco externo (tabelas sem prefixo = tabelas do banco). Devolve as linhas, se houver."""
    sql = str(p.get("sql") or "").strip()
    if not sql:
        raise ValueError("Escreva o SQL a executar.")
    with external_db(engine, con, p["db"]) as a:
        cur_db, cur_schema = con.execute("SELECT current_database(), current_schema()").fetchone()
        con.execute(f"USE {a}")
        try:
            con.execute(sql)
            if con.description:
                cols = [d[0] for d in con.description]
                rows = con.fetchmany(1000)
                return pd.DataFrame(rows, columns=cols)
            return None
        except duckdb.Error as exc:
            raise ValueError(f"Erro no SQL: {exc}")
        finally:
            with contextlib.suppress(Exception):
                con.execute(f"USE {qi(cur_db)}.{qi(cur_schema)}")


# ============================================================================
# OPERAÇÕES DO PIPELINE (interface E execução agendada)
# ============================================================================
def available(state: dict, layers) -> dict:
    out = {}
    for layer in layers:
        out.update(state[layer])
    return out


def _get(state, name, layers):
    ds = available(state, layers)
    if name not in ds:
        nomes = "/".join(LAYER_LABEL[l] for l in layers)
        raise KeyError(f"Dataset '{name}' não encontrado em {nomes}.")
    return ds[name]


def run_op(engine: Engine, state: dict, op: str, p: dict, ctx: Ctx) -> list:
    """Executa uma operação e devolve [(camada, nome, tabela), ...]."""
    con = engine.cursor()
    try:
        with ctx.watch(con):
            if op == "load_paths":
                return load_files(engine, con, p, ctx)
            if op == "load_sample":
                return load_sample(engine, con)
            if op == "db_create":
                db_create(engine, con, p)
                return []
            if op == "db_write":
                engine.last_report = db_write(engine, con, p, _get(state, p["src"], LAYERS))
                return []
            if op == "db_sql":
                engine.last_sql_result = db_sql(engine, con, p)
                return []
            if op == "dedup":
                layer, sql = "silver", sql_dedup(engine, con, _get(state, p["src"], SILVER_SOURCES), p.get("subset"))
            elif op == "nulls":
                layer, sql = "silver", sql_nulls(engine, con, _get(state, p["src"], SILVER_SOURCES),
                                                 p.get("fill"), p.get("dropna", False))
            elif op == "text_clean":
                layer, sql = "silver", sql_text_clean(engine, con, _get(state, p["src"], SILVER_SOURCES),
                                                      p["columns"], p.get("trim", True), p.get("spaces", True),
                                                      p.get("accents", False), p.get("case", ""),
                                                      p.get("empty_null", True))
            elif op == "cast":
                layer, sql = "silver", sql_cast(engine, con, _get(state, p["src"], SILVER_SOURCES),
                                                p["column"], p["type"])
            elif op == "merge":
                layer, sql = "silver", sql_merge(engine, con, _get(state, p["a"], SILVER_SOURCES),
                                                 _get(state, p["b"], SILVER_SOURCES),
                                                 p.get("left_on"), p.get("right_on"), p["how"])
            elif op == "filter":
                layer, sql = "silver", sql_filter(engine, con, _get(state, p["src"], SILVER_SOURCES),
                                                  p["column"], p["operator"], p.get("value", ""))
            elif op == "aggregate":
                layer, sql = "gold", sql_aggregate(engine, con, _get(state, p["src"], GOLD_SOURCES),
                                                   p.get("group_by") or [], p["columns"], p["func"])
            elif op == "select_cols":
                tid_src = _get(state, p["src"], LAYERS)
                cols = resolve_columns(engine, tid_src, p["columns"], p["src"], con)
                layer, sql = "gold", f"SELECT {projection_sql(cols)} FROM {tid_src}"
            elif op == "compute":
                layer, sql = "gold", sql_compute(engine, con, _get(state, p["src"], GOLD_SOURCES),
                                                 p["col_a"], p["operation"], p["col_b"], p["new_col"])
            else:
                raise ValueError(f"Operação desconhecida: {op}")
            tid = engine.new_tid()
            ctx.progress("Processando…", None, force=True)
            con.execute(f"CREATE TABLE {tid} AS {sql}")
            engine.lineage[tid] = {"kind": "sql", "sql": sql}
            return [(layer, p["name"], tid)]
    except duckdb.InterruptException:
        raise Cancelled()
    finally:
        con.close()


def apply_results(state: dict, step: dict, results: list) -> None:
    step["outputs"] = [[layer, name] for layer, name, _ in results]
    step["_prev"] = [(layer, name, state[layer].get(name)) for layer, name, _ in results]
    step["_new"] = [tid for _, _, tid in results]
    for layer, name, tid in results:
        state[layer][name] = tid


def clean_step(step: dict) -> dict:
    return {k: v for k, v in step.items() if not k.startswith("_") and k != "outputs"}


def replay(engine, steps: list, ctx: Ctx, log=None):
    """Reexecuta uma receita do zero, lendo os arquivos de origem novamente."""
    state = {layer: {} for layer in LAYERS}
    done = []
    for i, raw in enumerate(steps, 1):
        ctx.check()
        step = clean_step(raw)
        ctx.progress(f"Passo {i}/{len(steps)}: {step.get('label', step['op'])}", None, force=True)
        try:
            results = run_op(engine, state, step["op"], step["params"], ctx)
        except Cancelled:
            raise
        except Exception as exc:
            raise RuntimeError(f"Falha no passo {i} ({step.get('label', step['op'])}): {exc}") from exc
        apply_results(state, step, results)
        done.append(step)
        if log:
            log(f"  {i}. {step.get('label', step['op'])} → {', '.join(n for _, n, _ in results)}")
    return state, done


# ============================================================================
# GERAR CÓDIGO PYTHON (BD.py, etl_bronze.py, etl_prata.py, etl_ouro.py)
# um script por camada, que roda sem o app
# ============================================================================
LAYER_SUFFIX = {"bronze": "bronze", "silver": "prata", "gold": "ouro"}
_TID_RE = re.compile(r"\bt_\d{6}\b")
_SQL_BREAKS = ("FROM", "WHERE", "GROUP BY", "ORDER BY", "LEFT JOIN", "INNER JOIN", "FULL OUTER JOIN",
               "CROSS JOIN", "ON", "AND", "OR")


def sql_ident(name: str) -> str:
    """Nome amigável → identificador SQL simples (sem acentos, espaços ou símbolos)."""
    txt = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode("ascii")
    txt = re.sub(r"[^0-9A-Za-z]+", "_", txt).strip("_") or "tabela"
    return f"t_{txt}" if txt[0].isdigit() else txt


def default_table_name(name, layer) -> str:
    base, suffix = sql_ident(name), LAYER_SUFFIX[layer]
    return base if base.lower().endswith("_" + suffix) else f"{base}_{suffix}"


def ouro_name(prata_table: str) -> str:
    return prata_table[:-6] + "_ouro" if prata_table.lower().endswith("_prata") else prata_table + "_ouro"


def py_str(text) -> str:
    """Literal Python legível (r"..." para caminhos do Windows) quando possível."""
    t = str(text)
    if '"' not in t and not t.endswith("\\") and "\n" not in t:
        return f'r"{t}"'
    return repr(t)


def py_plain(text) -> str:
    """Texto entre aspas duplas, como no modelo ("Sheet1")."""
    t = str(text)
    return f'"{t}"' if '"' not in t and "\\" not in t and "\n" not in t else repr(t)


def format_sql(sql: str, indent="        ") -> str:
    """Quebra o SQL gerado em linhas (sem mexer em textos entre aspas)."""
    out, i, n, depth, quote = [], 0, len(sql), 0, None
    upper = sql.upper()
    stack = []                       # para cada "(": é a lista de um REPLACE (...)?
    while i < n:
        ch = sql[i]
        if quote:
            out.append(ch)
            if ch == quote:
                quote = None
            i += 1
            continue
        if ch in ("'", '"'):
            quote = ch
            out.append(ch)
            i += 1
            continue
        if ch == "(":
            depth += 1
            stack.append(upper[:i].rstrip().endswith("REPLACE"))
        elif ch == ")":
            depth -= 1
            if stack:
                stack.pop()
        if depth == 0 and ch == " ":
            hit = next((k for k in _SQL_BREAKS if upper.startswith(k + " ", i + 1)), None)
            if hit == "FROM" and upper[:i].rstrip().endswith("DISTINCT"):
                hit = None
            if hit:
                out.append("\n" + indent + ("    " if hit in ("AND", "OR", "ON") else ""))
                i += 1
                continue
        if ch == "," and (depth == 0 or (depth == 1 and stack and stack[-1])):
            # listas longas de colunas: uma por linha
            out.append(",\n" + indent + "    ")
            i += 1
            while i < n and sql[i] == " ":
                i += 1
            continue
        out.append(ch)
        i += 1
    return indent + "".join(out)


def _sql_type(t: str) -> str:
    if t == "VARCHAR":
        return "TEXT"
    if t.startswith("TIMESTAMP"):
        return "DATETIME"
    return t


def collect_project(engine, steps, export_cols) -> dict:
    """Junta o que o gerador precisa: tabelas bronze (com origem), consultas prata/ouro (com SQL)."""
    bronze, queries, problems = {}, [], []
    for step in steps:
        new = step.get("_new") or []
        outs = step.get("outputs") or []
        for tid, (layer, name) in zip(new, outs):
            lin = engine.lineage.get(tid)
            if lin is None:
                problems.append(f"{name}: origem desconhecida")
                continue
            if layer == "bronze":
                bronze[name] = {"tid": tid, **lin}
            elif lin.get("kind") == "sql":
                queries.append({"layer": layer, "name": name, "tid": tid, "sql": lin["sql"],
                                "label": step.get("label", step["op"]), "desc": step.get("desc", "")})
    return {"bronze": bronze, "queries": queries, "problems": problems, "export_cols": export_cols or {}}


def build_scripts(engine, project, names: dict, banco: str, pasta_parquet: str) -> dict:
    """Monta o conteúdo dos arquivos. names: {(camada, dataset): nome_da_tabela}."""
    gerado = f"Gerado pelo {APP_NAME} em {dt.datetime.now():%d/%m/%Y %H:%M}"
    tid_table = {}
    for name, b in project["bronze"].items():
        tid_table[b["tid"]] = names[("bronze", name)]
    for q in project["queries"]:
        tid_table[q["tid"]] = names[(q["layer"], q["name"])]

    def sql_for(q):
        return _TID_RE.sub(lambda m: tid_table.get(m.group(0), m.group(0)), q["sql"])

    # ---------------------------------------------------------------- BD.py
    bd = [f'"""Cria as tabelas da camada Bronze. {gerado}."""', "import duckdb", "",
          f"CAMINHO_BANCO = {py_str(banco)}", "", "conn = duckdb.connect(CAMINHO_BANCO)", "",
          "cursor = conn.cursor()", ""]
    drops = []
    for name, b in project["bronze"].items():
        if b.get("kind") == "sample":
            continue
        table = names[("bronze", name)]
        m = engine.meta(b["tid"])
        cols = ",\n".join(f"        {qi(c)} {_sql_type(t)}" for c, t in zip(m["cols"], m["types"]))
        bd += ['cursor.execute("""', f"    CREATE TABLE IF NOT EXISTS {table} (", cols, "    )", '""")', ""]
        drops.append(f'# cursor.execute("DROP TABLE IF EXISTS {table}")')
    bd += [""] + drops + ["", "conn.commit()", "", "conn.close()", ""]

    # ---------------------------------------------------------------- etl_bronze.py
    entries = []
    for name, b in project["bronze"].items():
        kind, table = b.get("kind"), names[("bronze", name)]
        if kind == "sample":
            continue
        if kind in ("excel", "excel_calamine"):
            e = [f'            "caminho_excel": {py_str(b["source"])},',
                 f'            "tabela_destino": "{table}",',
                 f'            "sheet": {py_plain(b["item"])}']
            if int(b.get("header") or 1) != 1:
                e[-1] += ","
                e.append(f'            "linha_cabecalho": {int(b["header"])}')
        else:
            e = [f'            "caminho_arquivo": {py_str(b["source"])},',
                 f'            "tabela_destino": "{table}",',
                 f'            "tipo": "{kind}"']
            if kind in ("sqlite", "duckdb"):
                e[-1] += ","
                e.append(f'            "tabela_origem": {py_plain(b["item"])}')
            if kind == "csv" and int(b.get("header") or 1) != 1:
                e[-1] += ","
                e.append(f'            "linha_cabecalho": {int(b["header"])}')
        entries.append("        {\n" + "\n".join(e) + "\n        }")
    bronze_py = BRONZE_TEMPLATE.replace("@GERADO@", gerado).replace("@BANCO@", py_str(banco)) \
        .replace("@ARQUIVOS@", ",\n".join(entries))

    # ---------------------------------------------------------------- etl_prata.py / etl_ouro.py
    prata_q = [q for q in project["queries"] if q["layer"] == "silver"]
    ouro_q = [q for q in project["queries"] if q["layer"] == "gold"]
    used_tables = set()
    for q in project["queries"]:
        used_tables |= {tid_table[t] for t in _TID_RE.findall(q["sql"]) if t in tid_table}

    def query_block(q):
        table = names[(q["layer"], q["name"])]
        var = f"query_{table.lower()}"
        body = format_sql(sql_for(q))
        quote = '"""' if '"""' not in body else "'''"
        comment = f"    # {q['label']}: {q['desc']}".rstrip()
        if len(comment) > 118:
            comment = comment[:115] + "…"
        return var, table, (f"{comment}\n    {var} = r{quote}\n        CREATE OR REPLACE TABLE {table} AS\n"
                            f"{body}\n    {quote}\n")

    blocks, execs, counts, prints = [], [], [], []
    for q in prata_q:
        var, table, block = query_block(q)
        blocks.append(block)
        execs.append(f'        print("Criando tabela: {table}...")\n        conn.execute({var})\n')
        counts.append(f'    total_{table.lower()} = conn.execute("SELECT COUNT(*) FROM {table}").fetchone()[0]')
        prints.append(f'    print(f"Linhas em {q["name"].upper()}: '
                      f'{{total_{table.lower()}}}")')
    prata_py = PRATA_TEMPLATE.replace("@GERADO@", gerado).replace("@BANCO@", py_str(banco)) \
        .replace("@QUERIES@", "\n".join(blocks) or "    # (nenhuma operação na camada Prata)\n") \
        .replace("@EXECS@", "\n".join(execs) or "        pass\n") \
        .replace("@COUNTS@", "\n".join(counts) or "    pass") \
        .replace("@PRINTS@", "\n".join(prints) or '    print("Nenhuma tabela Prata.")')

    # Ouro: consultas da camada Ouro + cópia das tabelas Prata finais (que nada mais usa)
    gold_tables, oblocks, oexecs = [], [], []
    for q in ouro_q:
        var, table, block = query_block(q)
        oblocks.append(block)
        oexecs.append(f"    conn.execute({var})")
        if table not in gold_tables:
            gold_tables.append(table)
    final_prata = []
    for q in prata_q:
        t = names[("silver", q["name"])]
        if t not in used_tables and t not in final_prata:
            final_prata.append(t)
    for t in final_prata:
        o = ouro_name(t)
        oexecs.append(f'    conn.execute("CREATE OR REPLACE TABLE {o} AS SELECT * FROM {t}")')
        if o not in gold_tables:
            gold_tables.append(o)
    ocounts = "\n".join(f'    qtd_{t.lower()} = conn.execute("SELECT COUNT(*) FROM {t}").fetchone()[0]'
                        for t in gold_tables) or "    pass"
    oprints = "\n".join(f'    print(f"Total de linhas {t}: {{qtd_{t.lower()}}}")' for t in gold_tables) \
        or '    print("Nenhuma tabela Ouro.")'
    parquet = ",\n".join(f'        "{t}": os.path.join(PASTA_PARQUET, "{t}.parquet")' for t in gold_tables)

    # colunas escolhidas/renomeadas na aba Exportar
    back = {}
    for q in prata_q + ouro_q:
        back[names[(q["layer"], q["name"])]] = q["name"]
    for t in final_prata:
        back[ouro_name(t)] = back.get(t)
    colmap = []
    for t in gold_tables:
        spec = project["export_cols"].get(back.get(t))
        if spec:
            pares = ", ".join(f"({c['col']!r}, {c.get('as') or c['col']!r})" for c in spec)
            colmap.append(f'        "{t}": [{pares}]')
    ouro_py = OURO_TEMPLATE.replace("@GERADO@", gerado).replace("@BANCO@", py_str(banco)) \
        .replace("@PARQUET@", py_str(pasta_parquet)) \
        .replace("@QUERIES@", "\n".join(oblocks) or "    # (nenhuma operação na camada Ouro)\n") \
        .replace("@EXECS@", "\n".join(oexecs) or "    pass") \
        .replace("@COUNTS@", ocounts).replace("@PRINTS@", oprints) \
        .replace("@ARQUIVOS@", parquet).replace("@COLUNAS@", ",\n".join(colmap))

    runner = RUNNER_TEMPLATE.replace("@GERADO@", gerado)
    return {"BD.py": "\n".join(bd), "etl_bronze.py": bronze_py, "etl_prata.py": prata_py,
            "etl_ouro.py": ouro_py, "executar_tudo.py": runner}


BRONZE_TEMPLATE = r'''"""Camada Bronze: lê os arquivos de origem e grava nas tabelas bronze. @GERADO@."""
import datetime
import json
import sqlite3

import duckdb
import pandas as pd

CAMINHO_BANCO = @BANCO@


def texto(v):
    """Mesma regra do Refina para gravar em colunas de texto: 1.0 → '1', datas → 'AAAA-MM-DD HH:MM:SS'."""
    if v is None or (isinstance(v, float) and v != v) or v is pd.NaT:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d 00:00:00")
    return str(v)


def ler_dados(item):
    """Lê a origem de um item da lista `arquivos` e devolve um DataFrame."""
    if "caminho_excel" in item:
        return pd.read_excel(item["caminho_excel"], sheet_name=item["sheet"],
                             header=item.get("linha_cabecalho", 1) - 1)
    caminho, tipo = item["caminho_arquivo"], item["tipo"]
    if tipo == "csv":
        pular = item.get("linha_cabecalho", 1) - 1
        try:
            return pd.read_csv(caminho, sep=None, engine="python", skiprows=pular, encoding="utf-8")
        except UnicodeDecodeError:
            return pd.read_csv(caminho, sep=None, engine="python", skiprows=pular, encoding="latin-1")
    if tipo == "parquet":
        return pd.read_parquet(caminho)
    if tipo == "json":
        if caminho.lower().startswith(("http://", "https://")):
            import urllib.request
            with urllib.request.urlopen(caminho) as resp:
                dados = json.load(resp)
        else:
            with open(caminho, encoding="utf-8") as fh:
                dados = json.load(fh)
        if isinstance(dados, dict):             # formato de API: {"data": [ {...}, ... ]}
            listas = [v for v in dados.values() if isinstance(v, list)]
            dados = listas[0] if len(listas) == 1 else [dados]
        return pd.json_normalize(dados)
    if tipo == "sqlite":
        with sqlite3.connect(caminho) as origem:
            return pd.read_sql_query(f'SELECT * FROM "{item["tabela_origem"]}"', origem)
    if tipo == "duckdb":
        origem = duckdb.connect(caminho, read_only=True)
        try:
            return origem.execute(f'SELECT * FROM "{item["tabela_origem"]}"').df()
        finally:
            origem.close()
    raise ValueError(f"Tipo de arquivo não suportado: {tipo}")


def camada_bronze():

    arquivos = [
@ARQUIVOS@
    ]

    for item in arquivos:
        caminho = item.get("caminho_excel") or item.get("caminho_arquivo")
        tabela = item["tabela_destino"]

        print("\n" + "="*50)
        print(f"Lendo o arquivo e preparando para a tabela: {tabela}...")

        # Conecta ao banco de dados (pode abrir uma vez para cada arquivo do loop)
        conn = duckdb.connect(CAMINHO_BANCO)

        # 1. Limpa a tabela bronze
        tabela_existe = conn.execute(
            f"SELECT COUNT(*) FROM information_schema.tables WHERE table_name = '{tabela}'"
        ).fetchone()[0]

        if tabela_existe > 0:
            conn.execute(f"DELETE FROM {tabela}")
            print(f"Tabela {tabela} limpa. Iniciando importação...")
        else:
            print(f"Tabela {tabela} não existia. A importação irá criá-la.")

        # 2. Lê os dados
        df = ler_dados(item)
        print(f"Arquivo processado: {caminho}")
        print(f"Linhas encontradas no arquivo: {len(df)}")

        # 3. Tratamento dos cabeçalhos (mesma regra do Refina: sem espaços sobrando)
        df.columns = [" ".join(str(c).split()) for c in df.columns]
        df = df.dropna(how="all")

        if tabela_existe == 0:
            conn.execute(f"CREATE TABLE {tabela} AS SELECT * FROM df")
            print("----------------------------------------")
            print(f"SUCESSO - Tabela {tabela} criada com {len(df)} linhas")
            conn.close()
            continue

        # 4. Obtém as colunas (e os tipos) da tabela de destino via DuckDB
        info = conn.execute(
            f"SELECT column_name, data_type FROM information_schema.columns WHERE table_name = '{tabela}'"
        ).fetchall()
        colunas_banco = [c for c, _ in info]
        tipos_banco = dict(info)

        # 5. Identifica interseção de colunas
        colunas_comuns = [col for col in df.columns if col in colunas_banco]

        if not colunas_comuns:
            print(f"ERRO: nenhuma coluna do arquivo foi encontrada na tabela {tabela}.")
            conn.close()
            continue

        # 6. Relatório de colunas
        print(f"Colunas em comum ({len(colunas_comuns)}): {colunas_comuns}")

        colunas_ignoradas = [col for col in df.columns if col not in colunas_banco]
        if colunas_ignoradas:
            print(f"Colunas do arquivo ignoradas: {colunas_ignoradas}")

        # 7. Filtra o DataFrame apenas com colunas válidas
        df = df[colunas_comuns].copy()

        # Colunas de texto no banco recebem texto (evita erro com colunas que misturam números e letras)
        for col in colunas_comuns:
            if tipos_banco.get(col) == "VARCHAR":
                df[col] = df[col].map(texto)

        if df.empty:
            print("Arquivo sem dados válidos para importar.")
            conn.close()
            continue

        # 8. Inserção direta via DuckDB
        colunas_str = ", ".join([f'"{col}"' for col in colunas_comuns])

        try:
            conn.execute(
                f"INSERT INTO {tabela} ({colunas_str}) SELECT {colunas_str} FROM df"
            )
            quantidade = len(df)
            print("----------------------------------------")
            print(f"SUCESSO - Linhas importadas na {tabela}: {quantidade}")
            print(f"Colunas importadas: {len(colunas_comuns)}")

        except Exception as e:
            print(f"ERRO ao inserir dados na tabela {tabela}: {e}")

        finally:
            # Garante que a conexão será fechada seja qual for o resultado
            conn.close()


def main():
    camada_bronze()

if __name__ == "__main__":
    main()
'''

PRATA_TEMPLATE = r'''"""Camada Prata: tratamentos e junções. @GERADO@."""
import duckdb

CAMINHO_BANCO = @BANCO@

def camada_prata():
    print("\nIniciando processamento da Camada Prata")

    # Conecta no banco
    conn = duckdb.connect(CAMINHO_BANCO)

    # Define as queries de transformação (na ordem em que foram aplicadas no Refina)
@QUERIES@
    # Executa as queries no DuckDB
    try:
@EXECS@
    except Exception as e:
        print(f"ERRO ao criar as tabelas prata:{e}")
        conn.close()
        return

    # Verifica quantas linhas foram geradas na camada Prata
@COUNTS@

    conn.close()

    print("========================================")
    print("CAMADA PRATA GERADA COM SUCESSO!")
    print("========================================")
@PRINTS@

if __name__ == "__main__":
    camada_prata()
'''

OURO_TEMPLATE = r'''"""Camada Ouro: consolidação e exportação para Parquet. @GERADO@."""
import os
import time
import duckdb

# Caminho do banco de dados DuckDB
CAMINHO_BANCO = @BANCO@

# Pasta de destino onde os arquivos .parquet serão salvos
PASTA_PARQUET = @PARQUET@

def camada_ouro():
    inicio = time.time()
    print("Iniciando a Camada Ouro e Exportação para Parquet...")

    # Garante que a pasta de destino exista no disco
    os.makedirs(PASTA_PARQUET, exist_ok=True)

    # Conecta ao banco DuckDB
    conn = duckdb.connect(CAMINHO_BANCO)

@QUERIES@
    # 1. Criação/Atualização das tabelas na Camada Ouro
    print("Criando tabelas na camada Ouro...")
@EXECS@

    # 2. Exibe no terminal a contagem de linhas
@COUNTS@

@PRINTS@

    # 3. Dicionário mapeando as tabelas e o caminho de cada arquivo Parquet
    arquivos_parquet = {
@ARQUIVOS@
    }

    # Colunas escolhidas (e renomeadas) na exportação — tabela ausente aqui = todas as colunas
    colunas_exportacao = {
@COLUNAS@
    }

    # 4. Exportação com o COPY do DuckDB
    print("\nExportando dados para Parquet...")
    for tabela, caminho_arquivo in arquivos_parquet.items():
        # Remove o arquivo antigo, se já existir
        if os.path.exists(caminho_arquivo):
            os.remove(caminho_arquivo)

        print(f"-> Gerando: {caminho_arquivo}")

        colunas = colunas_exportacao.get(tabela)
        selecao = ", ".join(f'"{c}" AS "{novo}"' for c, novo in colunas) if colunas else "*"
        caminho_sql = caminho_arquivo.replace("'", "''")

        # O comando COPY grava direto da tabela para o arquivo Parquet comprimido em SNAPPY
        conn.execute(f"""
            COPY (SELECT {selecao} FROM {tabela}) TO '{caminho_sql}' (FORMAT PARQUET, COMPRESSION 'SNAPPY');
        """)

    conn.close()

    fim = time.time()
    tempo_execucao = round(fim - inicio, 2)

    print("\n========================================")
    print("EXPORTAÇÃO PARQUET FINALIZADA COM SUCESSO!")
    print(f"Tempo total de execução: {tempo_execucao} segundos")
    print("========================================")

if __name__ == "__main__":
    camada_ouro()
'''

RUNNER_TEMPLATE = r'''"""Executa o pipeline completo: BD → Bronze → Prata → Ouro. @GERADO@."""
import os
import runpy
import time

PASTA = os.path.dirname(os.path.abspath(__file__))

if __name__ == "__main__":
    inicio = time.time()
    for script in ("BD.py", "etl_bronze.py", "etl_prata.py", "etl_ouro.py"):
        print(f"\n######## {script} ########")
        runpy.run_path(os.path.join(PASTA, script), run_name="__main__")
    print(f"\nPipeline completo em {time.time() - inicio:.1f} segundos.")
'''


# ============================================================================
# EXPORTAÇÃO EM FLUXO (grava em disco local e depois move para o destino)
# ============================================================================
def _safe_filename(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*]', "_", name).strip() or "dataset"


def _deliver(tmp: Path, dest, ctx: Ctx) -> Path:
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    ctx.progress(f"Salvando {dest.name}…", None, force=True)
    try:
        if dest.exists():
            dest.unlink()
        shutil.move(str(tmp), str(dest))
    except PermissionError:
        raise ValueError(f"Sem permissão para gravar {dest}. O arquivo está aberto em outro programa?")
    return dest


def resolve_columns(engine, tid, spec, dataset="dataset", con=None) -> list:
    """Converte a configuração [{"col", "as"}] em [(coluna, nome_final)], validando tudo."""
    if not spec:
        return []
    existing = {c.lower(): c for c in engine.columns(tid, con)}
    out, seen = [], set()
    for item in spec:
        col = existing.get(str(item["col"]).lower())
        if col is None:
            raise ValueError(f"{dataset}: a coluna '{item['col']}' não existe mais. "
                             f"Revise as colunas de exportação desse dataset.")
        alias = str(item.get("as") or col).strip() or col
        if alias.lower() in seen:
            raise ValueError(f"{dataset}: o nome '{alias}' aparece em mais de uma coluna.")
        seen.add(alias.lower())
        out.append((col, alias))
    if not out:
        raise ValueError(f"{dataset}: nenhuma coluna selecionada para exportar.")
    return out


def projection_sql(columns) -> str:
    if not columns:
        return "*"
    return ", ".join(qi(c) if c == a else f"{qi(c)} AS {qi(a)}" for c, a in columns)


def write_parquet(engine: Engine, tid, dest, columns=None, ctx: Ctx = None) -> Path:
    ctx = ctx or Ctx()
    tmp = engine.temp_file(".parquet")
    con = engine.cursor()
    try:
        with ctx.watch(con):
            ctx.progress(f"Gerando {Path(dest).name}…", None, force=True)
            con.execute(f"COPY (SELECT {projection_sql(columns)} FROM {tid}) TO {ql(str(tmp))} "
                        f"(FORMAT PARQUET, COMPRESSION ZSTD)")
    except duckdb.InterruptException:
        raise Cancelled()
    finally:
        con.close()
    return _deliver(tmp, dest, ctx)


_ILLEGAL = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


def _xl_value(v):
    if v is None or isinstance(v, (bool, int, decimal.Decimal)):
        return v
    if isinstance(v, float):
        return None if v != v else v
    if isinstance(v, str):
        if _ILLEGAL.search(v):
            v = _ILLEGAL.sub("", v)
        return v[:32767]
    if isinstance(v, dt.datetime):
        return v.replace(tzinfo=None) if v.tzinfo else v
    if isinstance(v, (dt.date, dt.time, dt.timedelta)):
        return v
    return str(v)


def _arrow_reader(result, rows):
    if hasattr(result, "to_arrow_reader"):
        return result.to_arrow_reader(rows)
    return result.fetch_record_batch(rows)


def write_xlsx(engine: Engine, tid, dest, columns=None, ctx: Ctx = None) -> Path:
    """Grava em fluxo (memória constante). Acima de 1.048.575 linhas, divide em várias abas."""
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    ctx = ctx or Ctx()
    meta = engine.meta(tid)
    total = meta["rows"]
    cols = [a for _, a in columns] if columns else meta["cols"]
    tmp = engine.temp_file(".xlsx")
    wb = openpyxl.Workbook(write_only=True)
    bold = Font(bold=True)
    ws, in_sheet, sheet_no, written = None, EXCEL_MAX_ROWS, 0, 0

    def new_sheet():
        nonlocal ws, in_sheet, sheet_no
        sheet_no += 1
        ws = wb.create_sheet("dados" if sheet_no == 1 else f"dados_{sheet_no}")
        header = []
        for c in cols:
            cell = WriteOnlyCell(ws, value=c)
            cell.font = bold
            header.append(cell)
        ws.append(header)
        in_sheet = 0

    con = engine.cursor()
    try:
        with ctx.watch(con):
            reader = _arrow_reader(con.execute(f"SELECT {projection_sql(columns)} FROM {tid}"), 20_000)
            for batch in reader:
                ctx.check()
                columns = [c.to_pylist() for c in batch.columns]
                for row in zip(*columns):
                    if in_sheet >= EXCEL_MAX_ROWS:
                        new_sheet()
                    ws.append([_xl_value(v) for v in row])
                    in_sheet += 1
                written += batch.num_rows
                ctx.progress(f"Gerando {Path(dest).name}: {fmt_int(written)} de {fmt_int(total)} linhas",
                             100 * written / max(total, 1))
        if ws is None:
            new_sheet()
        ctx.progress(f"Compactando {Path(dest).name}…", None, force=True)
        wb.save(str(tmp))
    except duckdb.InterruptException:
        raise Cancelled()
    finally:
        con.close()
    return _deliver(tmp, dest, ctx)


def write_duckdb(engine: Engine, items, dest, ctx: Ctx = None) -> Path:
    """Grava/atualiza tabelas num arquivo .duckdb (outras tabelas do arquivo são mantidas).
    items: [(nome_da_tabela, tabela_origem, colunas)]"""
    ctx = ctx or Ctx()
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    remote = is_remote(dest)
    work = engine.temp_file(".duckdb") if remote else dest
    if remote and dest.exists():
        ctx.progress(f"Copiando {dest.name} do servidor…", None, force=True)
        shutil.copyfile(dest, work)
    alias = f"destino_{engine.new_tid()}"
    con = engine.cursor()
    try:
        with ctx.watch(con):
            try:
                con.execute(f"ATTACH {ql(str(work))} AS {alias}")
            except duckdb.Error as exc:
                raise ValueError(f"Não consegui abrir {dest.name}. Se ele estiver aberto no DBeaver, Power BI "
                                 f"ou em outro programa, feche e tente de novo.\n\n{exc}")
            try:
                for name, tid, cols in items:
                    ctx.check()
                    ctx.progress(f"Gravando a tabela {name} em {dest.name}…", None, force=True)
                    con.execute(f"CREATE OR REPLACE TABLE {alias}.{qi(name)} AS "
                                f"SELECT {projection_sql(cols)} FROM {tid}")
            finally:
                con.execute(f"DETACH {alias}")
    except duckdb.InterruptException:
        raise Cancelled()
    finally:
        con.close()
    if remote:
        _deliver(work, dest, ctx)
    return dest


def export_state(engine, state, out_dir, layers, formats, ctx, export_columns=None, duckdb_file="dados.duckdb") -> list:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    export_columns = export_columns or {}
    written = []
    for layer in layers:
        for name, tid in state[layer].items():
            base = out / _safe_filename(name)
            cols = resolve_columns(engine, tid, export_columns.get(name), name)
            if "parquet" in formats:
                written.append(write_parquet(engine, tid, base.with_suffix(".parquet"), cols, ctx))
            if "xlsx" in formats:
                written.append(write_xlsx(engine, tid, base.with_suffix(".xlsx"), cols, ctx))
    if "duckdb" in formats:
        items = [(name, tid, resolve_columns(engine, tid, export_columns.get(name), name))
                 for layer in layers for name, tid in state[layer].items()]
        if items:
            written.append(write_duckdb(engine, items, out / (duckdb_file or "dados.duckdb"), ctx))
    return written


def run_recipe(recipe: dict, base_dir=None, ctx: Ctx = None, log=None) -> list:
    ctx = ctx or Ctx()
    out = Path(recipe.get("output_dir") or "saida")
    if not out.is_absolute() and base_dir:
        out = Path(base_dir) / out
    settings = recipe.get("settings") or {}
    engine = Engine(settings.get("mem_mb", 0), settings.get("cache_gb", DEFAULT_CACHE_GB))
    try:
        state, _ = replay(engine, recipe["steps"], ctx, log)
        return export_state(engine, state, out, recipe.get("layers", ["gold"]), recipe.get("formats", ["xlsx"]), ctx,
                            recipe.get("export_columns"), recipe.get("duckdb_file") or "dados.duckdb")
    finally:
        engine.close()


def cli_run(recipe_path: str) -> int:
    """Modo sem interface: python app.py --run receita.json (grava um .log ao lado da receita)."""
    rp = Path(recipe_path).resolve()
    log_path = rp.with_suffix(".log")

    def log(msg):
        line = f"[{dt.datetime.now():%Y-%m-%d %H:%M:%S}] {msg}"
        with contextlib.suppress(Exception):
            print(line)
        with open(log_path, "a", encoding="utf-8") as f:
            f.write(line + "\n")

    try:
        clean_stale_sessions()
        recipe = json.loads(rp.read_text(encoding="utf-8"))
        log(f"Iniciando receita {rp.name} ({len(recipe['steps'])} passos)")
        t0 = time.monotonic()
        files = run_recipe(recipe, base_dir=rp.parent, log=log)
        for f in files:
            log(f"  arquivo gerado: {f}")
        if not files:
            log("  aviso: nenhuma tabela nas camadas escolhidas para exportar.")
        log(f"Concluído com sucesso em {time.monotonic() - t0:.0f}s.")
        return 0
    except Exception:
        log("ERRO:\n" + traceback.format_exc())
        return 1


# ============================================================================
# ESTILO
# ============================================================================
INK = "#1E1A28"          # barra lateral
SURFACE = "#F4F3F7"      # fundo das telas
CARD = "#FFFFFF"
LINE = "#E5E2EC"
TEXT = "#211E2C"
MUTED = "#6F6A82"
PRIMARY = "#5B4FD6"      # ação principal (íris)
PRIMARY_DARK = "#4A3FC2"
PRIMARY_SOFT = "#ECEAFC"
WARN = "#B45309"
DANGER = "#B91C1C"
BLUE = PRIMARY           # compatibilidade

# Ícones de traço (24×24), desenhados para o app
ICON_PATHS = {
    "upload": '<path d="M12 15V4M7 9l5-5 5 5"/><path d="M4 15v4a1 1 0 0 0 1 1h14a1 1 0 0 0 1-1v-4"/>',
    "search": '<circle cx="11" cy="11" r="7"/><path d="M20 20l-4.2-4.2"/>',
    "layers": '<path d="M12 3l9 5-9 5-9-5 9-5z"/><path d="M3 13l9 5 9-5"/>',
    "award": '<circle cx="12" cy="9" r="6"/><path d="M8.6 13.8L7 21l5-3 5 3-1.6-7.2"/>',
    "clock": '<circle cx="12" cy="12" r="9"/><path d="M12 7v5l3 2"/>',
    "undo": '<path d="M9 14L4 9l5-5"/><path d="M4 9h11a5 5 0 0 1 0 10h-3"/>',
    "trash": '<path d="M4 7h16M10 11v6M14 11v6M6 7l1 13h10l1-13M9 7V4h6v3"/>',
    "download": '<path d="M12 4v11M7 10l5 5 5-5M4 20h16"/>',
    "database": '<ellipse cx="12" cy="6" rx="8" ry="3"/><path d="M4 6v12c0 1.7 3.6 3 8 3s8-1.3 8-3V6"/>'
                '<path d="M4 12c0 1.7 3.6 3 8 3s8-1.3 8-3"/>',
    "link": '<path d="M10 14a4 4 0 0 0 5.7 0l3-3a4 4 0 0 0-5.7-5.7l-1 1"/>'
            '<path d="M14 10a4 4 0 0 0-5.7 0l-3 3a4 4 0 0 0 5.7 5.7l1-1"/>',
    "filter": '<path d="M4 5h16l-6 8v6l-4-2v-4L4 5z"/>',
    "sparkles": '<path d="M11 3l1.7 4.3L17 9l-4.3 1.7L11 15l-1.7-4.3L5 9l4.3-1.7z"/>'
                '<path d="M18 14l.8 2.2L21 17l-2.2.8L18 20l-.8-2.2L15 17l2.2-.8z"/>',
    "eraser": '<path d="M7 20h13"/><path d="M4.5 15.5l9-9a2 2 0 0 1 2.8 0l2.2 2.2a2 2 0 0 1 0 2.8L11 19H8z"/>',
    "play": '<path d="M7 4.5l12 7.5-12 7.5z"/>',
    "folder": '<path d="M3 7a2 2 0 0 1 2-2h4l2 2h8a2 2 0 0 1 2 2v8a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2z"/>',
    "file": '<path d="M14 3H7a2 2 0 0 0-2 2v14a2 2 0 0 0 2 2h10a2 2 0 0 0 2-2V8z"/><path d="M14 3v5h5"/>',
    "refresh": '<path d="M20 11a8 8 0 0 0-14.3-4.9L4 8"/><path d="M4 4v4h4"/>'
               '<path d="M4 13a8 8 0 0 0 14.3 4.9L20 16"/><path d="M20 20v-4h-4"/>',
    "save": '<path d="M5 4h11l3 3v12a1 1 0 0 1-1 1H6a1 1 0 0 1-1-1z"/><path d="M8 4v5h7V4M8 20v-6h8v6"/>',
    "plus": '<path d="M12 5v14M5 12h14"/>',
    "eye": '<path d="M2 12s3.6-7 10-7 10 7 10 7-3.6 7-10 7S2 12 2 12z"/><circle cx="12" cy="12" r="3"/>',
    "chart": '<path d="M5 20V11M12 20V4M19 20v-7"/>',
    "calc": '<rect x="5" y="3" width="14" height="18" rx="2"/><path d="M8.5 7.5h7M8.5 12h.01M12 12h.01M15.5 12h.01'
            'M8.5 16h.01M12 16h.01M15.5 16h.01"/>',
    "package": '<path d="M3.5 7.5L12 3l8.5 4.5v9L12 21l-8.5-4.5z"/><path d="M3.5 7.5L12 12l8.5-4.5M12 12v9"/>',
    "dice": '<rect x="4" y="4" width="16" height="16" rx="3"/><path d="M9 9h.01M15 15h.01M15 9h.01M9 15h.01"/>',
    "up": '<path d="M12 19V5M6 11l6-6 6 6"/>',
    "down": '<path d="M12 5v14M6 13l6 6 6-6"/>',
    "list": '<path d="M9 6h11M9 12h11M9 18h11M4.5 6h.01M4.5 12h.01M4.5 18h.01"/>',
    "zap": '<path d="M13 3L5 13.5h6L10 21l8-10.5h-6z"/>',
    "network": '<rect x="9" y="3" width="6" height="5" rx="1"/><rect x="3" y="16" width="6" height="5" rx="1"/>'
               '<rect x="15" y="16" width="6" height="5" rx="1"/><path d="M12 8v4M6 16v-2h12v2"/>',
    "globe": '<circle cx="12" cy="12" r="9"/><path d="M3 12h18"/>'
             '<path d="M12 3c2.5 2.5 3.8 5.5 3.8 9s-1.3 6.5-3.8 9c-2.5-2.5-3.8-5.5-3.8-9S9.5 5.5 12 3z"/>',
    "x": '<path d="M7 7l10 10M17 7L7 17"/>',
    "chevron": '<path d="M6 9l6 6 6-6"/>',
    "chevron_up": '<path d="M6 15l6-6 6 6"/>',
    "check": '<path d="M5 12.5l4.5 4.5L19 7.5"/>',
    "dash": '<path d="M6 12h12"/>',
}

# emoji usado nos textos → ícone desenhado (o texto fica só com as palavras)
EMOJI_ICONS = {
    "📤": "upload", "🔎": "search", "🔍": "search", "⚙": "clock", "↩": "undo", "🗑": "trash", "⬇": "download",
    "🦆": "database", "🔗": "link", "🎯": "filter", "✨": "sparkles", "🧹": "eraser", "▶": "play", "📂": "folder",
    "📄": "file", "🔄": "refresh", "💾": "save", "➕": "plus", "👁": "eye", "📊": "chart", "🧮": "calc",
    "📦": "package", "🎲": "dice", "🧽": "eraser", "▲": "up", "▼": "down", "📜": "list", "⚡": "zap",
    "📡": "network", "📁": "folder", "🥉": "dot:bronze", "🥈": "dot:silver", "🥇": "dot:gold",
}
_EMOJI_PREFIX = re.compile(r"^\s*([^\w\s\"'(\[<.,:;!?+\-/\\]+)\s*")


def split_emoji(text):
    """'📤 Upload' → ('upload', 'Upload'). Sem emoji conhecido → (None, texto)."""
    m = _EMOJI_PREFIX.match(text or "")
    if not m:
        return None, text
    raw = m.group(1).replace("️", "").replace("‍", "")
    key = EMOJI_ICONS.get(raw) or EMOJI_ICONS.get(raw[:1])
    if key is None:
        return None, text
    return key, text[m.end():]


def icon_svg(name, color) -> str:
    if name.startswith("dot:"):
        c = LAYER_COLOR.get(name[4:], name[4:])
        return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24">'
                f'<circle cx="12" cy="12" r="6" fill="{c}"/></svg>')
    if name == "logo":   # funil em três faixas: entra bronze (largo), sai ouro (refinado)
        return ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 28 28" stroke-linejoin="round">'
                f'<path d="M1.5 3h25l-4 7.5h-17z" fill="{LAYER_COLOR["bronze"]}" '
                f'stroke="{LAYER_COLOR["bronze"]}" stroke-width="1.2"/>'
                f'<path d="M6.5 12.5h15l-3.5 6.5h-8z" fill="{LAYER_COLOR["silver"]}" '
                f'stroke="{LAYER_COLOR["silver"]}" stroke-width="1.2"/>'
                f'<path d="M11 21h6l-1.2 4.5h-3.6z" fill="{LAYER_COLOR["gold"]}" '
                f'stroke="{LAYER_COLOR["gold"]}" stroke-width="1.2"/></svg>')
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{color}" '
            f'stroke-width="1.9" stroke-linecap="round" stroke-linejoin="round">{ICON_PATHS[name]}</svg>')


def svg_pixmap(svg: str, size: int, ratio: int = 2) -> QPixmap:
    pm = QPixmap(size * ratio, size * ratio)
    pm.fill(Qt.transparent)
    painter = QPainter(pm)
    QSvgRenderer(QByteArray(svg.encode("utf-8"))).render(painter)
    painter.end()
    pm.setDevicePixelRatio(ratio)
    return pm


_ICON_CACHE = {}


def icon(name, color=MUTED, size=18, selected=None) -> QIcon:
    key = (name, color, size, selected)
    ic = _ICON_CACHE.get(key)
    if ic is None:
        ic = QIcon()
        ic.addPixmap(svg_pixmap(icon_svg(name, color), size), QIcon.Normal)
        if selected:
            ic.addPixmap(svg_pixmap(icon_svg(name, selected), size), QIcon.Selected)
        if not name.startswith("dot:") and name != "logo":
            ic.addPixmap(svg_pixmap(icon_svg(name, "#ABA6BA"), size), QIcon.Disabled)
        _ICON_CACHE[key] = ic
    return ic


def layer_icon(layer) -> QIcon:
    return icon(f"dot:{layer}", size=14)


def layer_html(layer) -> str:
    return f"<span style='color:{LAYER_COLOR[layer]}'>●</span> {LAYER_LABEL[layer]}"


QSS = r"""
* { font-family: 'Segoe UI Variable Text', 'Segoe UI', 'Inter', 'Noto Sans', sans-serif; font-size: 9.5pt; color: @TEXT@; }
QMainWindow, QWidget#page, QDialog { background: @SURFACE@; }
QScrollArea { background: transparent; border: none; }
QScrollArea > QWidget > QWidget { background: transparent; }
QWidget#page { background: @SURFACE@; }

/* ---------- barra lateral ---------- */
QFrame#sidebar { background: @INK@; border: none; }
QFrame#sidebar QLabel { color: #BAB4CC; background: transparent; }
QFrame#sidebar QLabel#appTitle { color: #FFFFFF; font-size: 14pt; font-weight: 600; }
QFrame#sidebar QLabel#appSubtitle { color: #8F89A6; font-size: 9pt; }
QFrame#sidebar QLabel#sectionLabel { color: #7D7795; font-size: 9pt; font-weight: 600; }
QFrame#sidebar QLabel#count { color: #FFFFFF; font-weight: 600; }
QListWidget#nav { background: transparent; border: none; outline: 0; font-size: 10pt; }
QListWidget#nav::item { color: #BAB4CC; padding: 7px 10px; margin: 1px 0; border-radius: 7px; }
QListWidget#nav::item:hover { background: #2B2638; color: #ECE9F4; }
QListWidget#nav::item:selected { background: @PRIMARY@; color: #FFFFFF; font-weight: 600; }
QFrame#statusCard { background: transparent; border-radius: 8px; }
QFrame#statusCard:hover { background: #2B2638; }
QPushButton#sideButton { background: transparent; color: #BAB4CC; border: 1px solid #3A3449; padding: 5px 8px; }
QPushButton#sideButton:hover { background: #2B2638; color: #FFFFFF; border-color: #4C4561; }

/* ---------- cabeçalhos e textos ---------- */
QLabel#pageTitle { font-size: 16pt; font-weight: 600; color: #1A1724; }
QLabel#pageCaption { color: @MUTED@; }
QLabel#hint { color: @MUTED@; font-size: 8.5pt; }
QLabel#metricLabel { color: @MUTED@; font-size: 8.5pt; }
QLabel#metricValue { color: #1A1724; font-size: 16pt; font-weight: 600; }
QFrame#metric { background: @CARD@; border: 1px solid @LINE@; border-radius: 10px; }
QLabel#report { background: #F7F6FA; border: 1px solid @LINE@; border-left: 3px solid @PRIMARY@;
                border-radius: 8px; padding: 10px 12px; }

/* ---------- cartões (grupos) com o título por dentro ---------- */
QGroupBox { background: @CARD@; border: 1px solid @LINE@; border-radius: 10px; margin-top: 0px;
            padding: 34px 12px 12px 12px; font-size: 10pt; font-weight: 600; }
QGroupBox::title { subcontrol-origin: padding; subcontrol-position: top left; left: 12px; top: 10px; color: #1A1724; }
QGroupBox QLabel, QGroupBox QCheckBox, QGroupBox QListWidget, QGroupBox QLineEdit, QGroupBox QComboBox,
QGroupBox QPushButton, QGroupBox QTableView, QGroupBox QSpinBox { font-size: 9.5pt; font-weight: 400; }
QGroupBox QPushButton[primary="true"] { font-weight: 600; }
QTabWidget QGroupBox { background: #FBFAFC; }

/* ---------- botões ---------- */
QPushButton { background: @CARD@; border: 1px solid #DCD9E6; border-radius: 7px; padding: 5px 12px;
              color: #2F2B3D; font-weight: 500; }
QPushButton:hover { background: #F6F5FA; border-color: #C3BED2; }
QPushButton:pressed { background: #ECEAF2; }
QPushButton:focus { border-color: @PRIMARY@; }
QPushButton:disabled { color: #ABA6BA; background: #F7F6FA; border-color: #E7E5EE; }
QPushButton[primary="true"] { background: @PRIMARY@; border: 1px solid @PRIMARY@; color: #FFFFFF; font-weight: 600; }
QPushButton[primary="true"]:hover { background: @PRIMARY_DARK@; border-color: @PRIMARY_DARK@; }
QPushButton[primary="true"]:pressed { background: #3A2F9A; }
QPushButton[primary="true"]:focus { border: 1px solid #3A2F9A; }

/* ---------- campos ---------- */
QLineEdit, QComboBox, QSpinBox { background: @CARD@; border: 1px solid #DCD9E6; border-radius: 8px;
    padding: 4px 9px; min-height: 18px; selection-background-color: #DEDAF8; selection-color: @TEXT@; }
QLineEdit:hover, QComboBox:hover, QSpinBox:hover { border-color: #C3BED2; }
QLineEdit:focus, QComboBox:focus, QSpinBox:focus { border: 1px solid @PRIMARY@; }
QLineEdit:disabled, QComboBox:disabled { background: #F7F6FA; color: #ABA6BA; }
QComboBox { padding-right: 30px; }
QComboBox::drop-down { border: none; width: 30px; subcontrol-origin: padding; subcontrol-position: center right; }
QComboBox::down-arrow { image: url(@A@/chevron.png); width: 14px; height: 14px; }
QComboBox QAbstractItemView { background: @CARD@; border: 1px solid @LINE@; padding: 4px; outline: 0;
    selection-background-color: @PRIMARY_SOFT@; selection-color: @TEXT@; }
QSpinBox { padding-right: 22px; }
QSpinBox::up-button, QSpinBox::down-button { border: none; background: transparent; width: 20px; }
QSpinBox::up-arrow { image: url(@A@/chevron_up.png); width: 10px; height: 10px; }
QSpinBox::down-arrow { image: url(@A@/chevron.png); width: 10px; height: 10px; }

/* ---------- caixas de marcar (inclusive dentro de listas e tabelas) ---------- */
QCheckBox { spacing: 8px; background: transparent; }
QCheckBox::indicator, QListView::indicator, QTreeView::indicator, QTableView::indicator {
    width: 16px; height: 16px; border: 1px solid #C3BED2; border-radius: 4px; background: #FFFFFF; }
QCheckBox::indicator:hover, QListView::indicator:hover, QTreeView::indicator:hover, QTableView::indicator:hover {
    border-color: @PRIMARY@; }
QCheckBox::indicator:checked, QListView::indicator:checked, QTreeView::indicator:checked, QTableView::indicator:checked {
    background: @PRIMARY@; border-color: @PRIMARY@; image: url(@A@/check.png); }
QCheckBox::indicator:indeterminate, QTreeView::indicator:indeterminate {
    background: @PRIMARY@; border-color: @PRIMARY@; image: url(@A@/dash.png); }

/* ---------- abas: sublinhado, conteúdo num cartão ---------- */
QTabWidget::pane { background: @CARD@; border: 1px solid @LINE@; border-radius: 10px; top: -1px; }
QTabBar { background: transparent; }
QTabBar::tab { background: transparent; color: @MUTED@; padding: 6px 11px 7px 11px; margin-right: 2px;
               border: none; border-bottom: 2px solid transparent; font-weight: 600; }
QTabBar::tab:hover { color: @TEXT@; }
QTabBar::tab:selected { color: @PRIMARY@; border-bottom: 2px solid @PRIMARY@; }

/* ---------- tabelas e listas ---------- */
QTableView, QTreeWidget, QListWidget { background: @CARD@; border: 1px solid @LINE@; border-radius: 8px;
    gridline-color: #EFEDF4; selection-background-color: @PRIMARY_SOFT@; selection-color: @TEXT@;
    alternate-background-color: #F9F8FB; outline: 0; }
QListWidget::item { padding: 3px 6px; border-radius: 6px; }
QListWidget::item:hover { background: #F4F3F8; }
QListWidget::item:selected { background: @PRIMARY_SOFT@; color: @TEXT@; }
QTreeWidget::item { padding: 3px 2px; }
QHeaderView { background: @CARD@; border: none; }
QHeaderView::section { background: #F7F6FA; color: #4F4A62; border: none; border-bottom: 1px solid @LINE@;
    border-right: 1px solid #EFEDF4; padding: 6px 8px; font-weight: 600; font-size: 9pt; }
QTableCornerButton::section { background: #F7F6FA; border: none; border-bottom: 1px solid @LINE@; }

/* ---------- rolagem, divisória, progresso, dicas ---------- */
QScrollBar:vertical { background: transparent; width: 10px; margin: 2px; }
QScrollBar:horizontal { background: transparent; height: 10px; margin: 2px; }
QScrollBar::handle:vertical { background: #CFCBDB; border-radius: 3px; min-height: 32px; }
QScrollBar::handle:horizontal { background: #CFCBDB; border-radius: 3px; min-width: 32px; }
QScrollBar::handle:hover { background: #B2ACC4; }
QScrollBar::add-line, QScrollBar::sub-line { width: 0px; height: 0px; }
QScrollBar::add-page, QScrollBar::sub-page { background: transparent; }
QSplitter::handle { background: transparent; width: 10px; }
QProgressBar { background: #EAE8F1; border: none; border-radius: 4px; max-height: 8px; text-align: center; color: transparent; }
QProgressBar::chunk { background: @PRIMARY@; border-radius: 4px; }
QProgressDialog, QMessageBox { background: @CARD@; }
QToolTip { background: @TEXT@; color: #FFFFFF; border: none; padding: 6px 8px; }
QStatusBar { background: @CARD@; border-top: 1px solid @LINE@; }
QStatusBar QLabel { color: @MUTED@; }
QPlainTextEdit { background: @CARD@; border: 1px solid #DCD9E6; border-radius: 8px; padding: 6px;
                 selection-background-color: #DEDAF8; selection-color: @TEXT@; }
QPlainTextEdit:focus { border: 1px solid @PRIMARY@; }
QPlainTextEdit#code { font-family: 'Cascadia Mono', 'Consolas', 'DejaVu Sans Mono', monospace; font-size: 9.5pt; }
QToolButton#sideZoom { background: transparent; color: #BAB4CC; border: 1px solid #3A3449; border-radius: 6px;
                       min-width: 22px; min-height: 20px; font-size: 11pt; font-weight: 600; }
QToolButton#sideZoom:hover { background: #2B2638; color: #FFFFFF; border-color: #4C4561; }
QListWidget#recent::item { padding-right: 34px; }
QToolButton#rowDelete { background: transparent; border: none; border-radius: 6px; padding: 3px; }
QToolButton#rowDelete:hover { background: #FDE8E8; }
QPushButton#linkButton { background: transparent; border: none; color: @PRIMARY@; padding: 0px 2px;
                         font-size: 9pt; font-weight: 600; }
QPushButton#linkButton:hover { color: @PRIMARY_DARK@; text-decoration: underline; }
QPushButton#linkButton:disabled { color: #ABA6BA; }
"""


def build_style() -> str:
    """Gera os pequenos ícones usados pelo estilo (setas, marcação) e devolve a folha de estilo."""
    folder = app_root() / "ui"
    folder.mkdir(parents=True, exist_ok=True)
    for name, color, size in (("chevron", MUTED, 14), ("chevron_up", MUTED, 14), ("check", "#FFFFFF", 16),
                              ("dash", "#FFFFFF", 16)):
        svg_pixmap(icon_svg(name, color), size, ratio=3).save(str(folder / f"{name}.png"))
    qss = QSS.replace("@A@", folder.as_posix())
    for k, v in (("SURFACE", SURFACE), ("CARD", CARD), ("LINE", LINE), ("TEXT", TEXT), ("MUTED", MUTED),
                 ("PRIMARY_DARK", PRIMARY_DARK), ("PRIMARY_SOFT", PRIMARY_SOFT), ("PRIMARY", PRIMARY), ("INK", INK)):
        qss = qss.replace(f"@{k}@", v)
    return qss


def light_palette() -> QPalette:
    pal = QPalette()
    for role, color in [(QPalette.Window, SURFACE), (QPalette.WindowText, TEXT),
                        (QPalette.Base, CARD), (QPalette.AlternateBase, "#F9F8FB"),
                        (QPalette.Text, TEXT), (QPalette.Button, CARD),
                        (QPalette.ButtonText, TEXT), (QPalette.Highlight, PRIMARY),
                        (QPalette.HighlightedText, "#FFFFFF"), (QPalette.ToolTipBase, TEXT),
                        (QPalette.ToolTipText, "#FFFFFF"), (QPalette.PlaceholderText, "#A09AB0")]:
        pal.setColor(role, QColor(color))
    return pal


def polish(root):
    """Troca os emojis de abas, títulos, campos e caixas de marcar por ícones desenhados."""
    for tw in root.findChildren(QTabWidget):
        tw.setIconSize(QSize(16, 16))
        for i in range(tw.count()):
            key, txt = split_emoji(tw.tabText(i))
            if key:
                tw.setTabText(i, txt.replace("&&", "&&"))
                tw.setTabIcon(i, icon(key, MUTED, 16))
    for gb in root.findChildren(QGroupBox):
        key, txt = split_emoji(gb.title())
        if key:
            gb.setTitle(txt)
    for le in root.findChildren(QLineEdit):
        key, txt = split_emoji(le.placeholderText())
        if key:
            le.setPlaceholderText(txt)
            if key == "search":
                le.addAction(icon("search", "#928CA6", 16), QLineEdit.LeadingPosition)
    for cb in root.findChildren(QCheckBox):
        key, txt = split_emoji(cb.text())
        if key:
            cb.setText(txt)
            cb.setIcon(icon(key, MUTED, 14))


# ============================================================================
# COMPONENTES DE UI
# ============================================================================
def make_button(text, slot=None, primary=False, icon_color=None) -> QPushButton:
    key, txt = split_emoji(text)
    btn = QPushButton(txt.strip())
    btn.setCursor(Qt.PointingHandCursor)
    if key:
        btn.setIcon(icon(key, icon_color or ("#FFFFFF" if primary else "#4F4A62"), 16))
        btn.setIconSize(QSize(16, 16))
    if primary:
        btn.setProperty("primary", True)
    if slot:
        btn.clicked.connect(slot)
    return btn


def left(widget) -> QWidget:
    box = QWidget()
    lay = QHBoxLayout(box)
    lay.setContentsMargins(0, 0, 0, 0)
    lay.addWidget(widget)
    lay.addStretch()
    return box


def hint(text) -> QLabel:
    lbl = QLabel(text)
    lbl.setObjectName("hint")
    lbl.setWordWrap(True)
    return lbl


def fill_combo(combo: QComboBox, items) -> None:
    current = combo.currentText()
    items = [str(i) for i in items]
    combo.blockSignals(True)
    combo.clear()
    combo.addItems(items)
    if current in items:
        combo.setCurrentText(current)
    combo.blockSignals(False)


def _fmt_cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v != v:
            return ""
        return str(int(v)) + ".0" if v.is_integer() and abs(v) < 1e15 else repr(v)
    if isinstance(v, dt.datetime):
        return v.strftime("%d/%m/%Y") if (v.hour, v.minute, v.second, v.microsecond) == (0, 0, 0, 0) \
            else v.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    return str(v)


class PandasModel(QAbstractTableModel):
    """Para tabelas pequenas (esquema)."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self._df = pd.DataFrame()

    def set_df(self, df):
        self.beginResetModel()
        self._df = df if df is not None else pd.DataFrame()
        self.endResetModel()

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self._df)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else self._df.shape[1]

    def data(self, index, role=Qt.DisplayRole):
        if index.isValid() and role in (Qt.DisplayRole, Qt.ToolTipRole):
            return _fmt_cell(self._df.iat[index.row(), index.column()])
        return None

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        if role != Qt.DisplayRole:
            return None
        return str(self._df.columns[section]) if orientation == Qt.Horizontal else str(section + 1)


class DuckModel(QAbstractTableModel):
    """Lê do DuckDB só os blocos de linhas que aparecem na tela (com cache pequeno)."""
    BLOCK = 500
    MAX_BLOCKS = 24
    sortRequested = Signal(int, object)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.engine, self.table = None, None
        self.rows, self.cols, self._numeric = 0, [], set()
        self.select = "*"
        self._cache = {}

    def set_source(self, engine, table, rows, cols, types, select="*"):
        self.beginResetModel()
        self.engine, self.table, self.rows, self.cols = engine, table, rows, list(cols)
        self.select = select
        self._numeric = {i for i, t in enumerate(types) if is_numeric_type(t)}
        self._cache = {}
        self.endResetModel()

    def clear(self):
        self.set_source(None, None, 0, [], [])

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else self.rows

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.cols)

    def _block(self, b):
        rows = self._cache.get(b)
        if rows is None:
            s = b * self.BLOCK
            try:
                rows = self.engine.ui.execute(
                    f"SELECT {self.select} FROM {self.table} WHERE rowid >= {s} AND rowid < {s + self.BLOCK} "
                    f"ORDER BY rowid").fetchall()
            except Exception:
                rows = []
            if len(self._cache) >= self.MAX_BLOCKS:
                self._cache.pop(next(iter(self._cache)))
            self._cache[b] = rows
        return rows

    def data(self, index, role=Qt.DisplayRole):
        if not index.isValid() or self.table is None:
            return None
        if role in (Qt.DisplayRole, Qt.ToolTipRole):
            rows = self._block(index.row() // self.BLOCK)
            i = index.row() % self.BLOCK
            if i < len(rows):
                return _fmt_cell(rows[i][index.column()])
            return ""
        if role == Qt.TextAlignmentRole and index.column() in self._numeric:
            return Qt.AlignRight | Qt.AlignVCenter
        return None

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        if role != Qt.DisplayRole:
            return None
        if orientation == Qt.Horizontal:
            return self.cols[section] if section < len(self.cols) else None
        return fmt_int(section + 1)

    def sort(self, column, order=Qt.AscendingOrder):
        self.sortRequested.emit(column, order)


class DataView(QWidget):
    """Tabela paginada sob demanda, com busca (Enter) e ordenação por clique no cabeçalho.
    Busca e ordenação geram uma tabela auxiliar no DuckDB (nada vai para a RAM)."""
    _seq = itertools.count(1)

    def __init__(self, min_height=320):
        super().__init__()
        self.engine, self.tid = None, None
        self._view = f"u_view_{next(self._seq)}"
        self._sort = None
        self._proj = None          # [(coluna, nome exibido)] ou None = todas
        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        top = QHBoxLayout()
        self.search = QLineEdit()
        self.search.setPlaceholderText("🔎 Buscar em todas as colunas (Enter)")
        self.search.returnPressed.connect(self._apply)
        self.search.textChanged.connect(lambda t: self._apply() if not t else None)
        self.info = QLabel()
        self.info.setObjectName("hint")
        top.addWidget(self.search, 1)
        top.addWidget(self.info)
        lay.addLayout(top)
        self.table = QTableView()
        self.model = DuckModel(self)
        self.model.sortRequested.connect(self._on_sort)
        self.table.setModel(self.model)
        self.table.setAlternatingRowColors(True)
        self.table.setWordWrap(False)
        vh = self.table.verticalHeader()
        vh.setSectionResizeMode(QHeaderView.Fixed)
        vh.setDefaultSectionSize(25)
        hh = self.table.horizontalHeader()
        hh.setMaximumSectionSize(420)
        hh.setResizeContentsPrecision(60)
        hh.setSortIndicator(-1, Qt.AscendingOrder)
        self.table.setSortingEnabled(True)
        self.table.setMinimumHeight(min_height)
        lay.addWidget(self.table, 1)

    def set_projection(self, proj):
        proj = list(proj) if proj else None
        if proj == self._proj:
            return
        self._proj = proj
        self._sort = None
        if self.tid is not None:
            self._apply()

    def set_table(self, engine, tid, proj=None):
        if (engine, tid) == (self.engine, self.tid):
            if proj is not None or self._proj is not None:
                self.set_projection(proj)
            return
        self.engine, self.tid, self._sort = engine, tid, None
        self._proj = list(proj) if proj else None
        self.search.blockSignals(True)
        self.search.clear()
        self.search.blockSignals(False)
        self._apply()

    def _on_sort(self, column, order):
        if self.tid is None:
            return
        self._sort = (column, order)
        self._apply()

    def _apply(self):
        if self.engine is None or self.tid is None:
            self.model.clear()
            self.info.setText("Sem dados para exibir.")
            return
        meta = self.engine.meta(self.tid)
        types = dict(zip(meta["cols"], meta["types"]))
        proj = [(c, a) for c, a in (self._proj or []) if c in types] or [(c, c) for c in meta["cols"]]
        shown_cols = [a for _, a in proj]
        shown_types = [types[c] for c, _ in proj]
        select = projection_sql(proj) if self._proj else "*"
        text = self.search.text().strip()
        source = self.tid
        if text or self._sort:
            where, order = "", "ORDER BY rowid"
            if text and proj:
                lit = ql(text.lower())
                where = "WHERE " + " OR ".join(
                    f"contains(lower(CAST({qi(c)} AS VARCHAR)), {lit})" for c, _ in proj)
            if self._sort and self._sort[0] < len(proj):
                col, ordr = self._sort
                d = "ASC" if ordr == Qt.AscendingOrder else "DESC"
                order = f"ORDER BY {qi(proj[col][0])} {d} NULLS LAST, rowid"
            QApplication.setOverrideCursor(Qt.WaitCursor)
            try:
                self.engine.ui.execute(f"CREATE OR REPLACE TABLE {self._view} AS "
                                       f"SELECT * FROM {self.tid} {where} {order}")
                n = self.engine.ui.execute(f"SELECT count(*) FROM {self._view}").fetchone()[0]
            finally:
                QApplication.restoreOverrideCursor()
            source = self._view
        else:
            n = meta["rows"]
            self.engine.drop(self._view)
            self.table.horizontalHeader().setSortIndicator(-1, Qt.AscendingOrder)
        self.model.set_source(self.engine, source, n, shown_cols, shown_types, select)
        vh = self.table.verticalHeader()
        vh.setFixedWidth(vh.fontMetrics().horizontalAdvance(fmt_int(max(n, 1))) + 24)
        self.table.resizeColumnsToContents()
        if text:
            self.info.setText(f"{fmt_int(n)} de {fmt_int(meta['rows'])} linhas")
        else:
            self.info.setText(f"{fmt_int(meta['rows'])} linhas × {len(shown_cols)} colunas")


class SmallTable(QTableView):
    def __init__(self, min_height=200):
        super().__init__()
        self.model_ = PandasModel(self)
        self.setModel(self.model_)
        self.setAlternatingRowColors(True)
        self.verticalHeader().setDefaultSectionSize(24)
        self.setMinimumHeight(min_height)

    def set_df(self, df):
        self.model_.set_df(df)
        self.resizeColumnsToContents()


class CheckList(QListWidget):
    def __init__(self, max_height=150):
        super().__init__()
        self.setMaximumHeight(max_height)

    def set_items(self, items):
        prev = set(self.checked())
        self.clear()
        for it in items:
            w = QListWidgetItem(str(it))
            w.setFlags(w.flags() | Qt.ItemIsUserCheckable)
            w.setCheckState(Qt.Checked if str(it) in prev else Qt.Unchecked)
            self.addItem(w)

    def checked(self):
        return [self.item(i).text() for i in range(self.count())
                if self.item(i).checkState() == Qt.Checked]


class StepsLog(QListWidget):
    def __init__(self, store, min_height=160):
        super().__init__()
        self.store = store
        self.setMinimumHeight(min_height)
        self.setSelectionMode(QListWidget.NoSelection)
        self.setWordWrap(True)
        store.changed.connect(self.refresh)
        self.refresh()

    def refresh(self):
        self.clear()
        if not self.store.steps:
            self.addItem("Nenhum passo aplicado ainda.")
            return
        n = len(self.store.steps)
        by_label = {v: k for k, v in LAYER_LABEL.items()}
        for i, step in enumerate(reversed(self.store.steps)):
            item = QListWidgetItem(f"{n - i}. {step['label']} — {step['desc']}")
            layer = by_label.get(step["layer"])
            if layer:
                item.setIcon(layer_icon(layer))
            elif step["layer"] == "Banco":
                item.setIcon(icon("database", PRIMARY, 14))
            self.addItem(item)


class MetricCard(QFrame):
    def __init__(self, label):
        super().__init__()
        self.setObjectName("metric")
        v = QVBoxLayout(self)
        v.setContentsMargins(14, 8, 14, 8)
        v.setSpacing(0)
        lbl = QLabel(label)
        lbl.setObjectName("metricLabel")
        self.value = QLabel("—")
        self.value.setObjectName("metricValue")
        v.addWidget(lbl)
        v.addWidget(self.value)


def layer_of(store, name, layers):
    """Camada onde o dataset vive (a camada mais à direita vence, como em available())."""
    for layer in reversed(layers):
        if name in store.data[layer]:
            return layer
    return None


class PreviewPanel(QGroupBox):
    """Painel lateral que mostra a origem escolhida e, depois de aplicar, o resultado."""

    def __init__(self, page):
        super().__init__("👁️ Pré-visualização")
        self.page = page
        self.current = None
        v = QVBoxLayout(self)
        self.head = QLabel()
        self.head.setWordWrap(True)
        self.head.setTextFormat(Qt.RichText)
        v.addWidget(self.head)
        self.view = DataView(min_height=280)
        v.addWidget(self.view, 1)
        row = QHBoxLayout()
        row.addStretch()
        self.open_btn = make_button("🔎 Abrir no Explorar", self._open)
        row.addWidget(self.open_btn)
        v.addLayout(row)
        self.show_dataset(None, None)

    def show_dataset(self, layer, name, note=""):
        store, engine = self.page.store, self.page.engine
        tid = store.data[layer].get(name) if layer else None
        if tid is None:
            self.current = None
            self.view.set_table(None, None)
            self.head.setText("<span style='color:#6E6884'>Escolha um dataset de origem para ver os dados aqui.</span>")
            self.open_btn.setEnabled(False)
            return
        self.current = (layer, name)
        m = engine.meta(tid)
        self.head.setText(f"{note}<b>{name}</b> &nbsp;&nbsp;{layer_html(layer)}&nbsp;&nbsp; "
                          f"{fmt_int(m['rows'])} linhas × {len(m['cols'])} colunas")
        self.view.set_table(engine, tid)
        self.open_btn.setEnabled(True)

    def _open(self):
        if self.current:
            self.page.win.open_in_explorer(*self.current)


class ClickCard(QFrame):
    def __init__(self, on_click):
        super().__init__()
        self._on_click = on_click
        self.setCursor(Qt.PointingHandCursor)

    def mousePressEvent(self, event):
        self._on_click()
        super().mousePressEvent(event)


class SelectDialog(QDialog):
    """Escolha de abas (Excel) e tabelas (bancos) antes de carregar."""

    def __init__(self, infos, parent=None):
        super().__init__(parent)
        self.infos = infos
        self.setWindowTitle("Escolher o que carregar")
        self.setWindowIcon(QIcon(svg_pixmap(icon_svg("logo", ""), 32)))
        self.resize(820, 520)
        v = QVBoxLayout(self)
        v.addWidget(hint("Marque as abas e tabelas que quer trazer para a Bronze. Só o que estiver marcado "
                         "é lido, e a escolha fica gravada na receita. O número de linhas é aproximado."))
        self.tree = QTreeWidget()
        self.tree.setHeaderLabels(["Arquivo / aba", "Linhas (aprox.)", "Colunas", "Cabeçalho na linha", "Tamanho"])
        self.tree.setAlternatingRowColors(True)
        self.tree.setUniformRowHeights(True)
        hdr = self.tree.header()
        hdr.setSectionResizeMode(0, QHeaderView.Stretch)
        for i in (1, 2, 3, 4):
            hdr.setSectionResizeMode(i, QHeaderView.ResizeToContents)
        self.tree.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.SelectedClicked
                                  | QAbstractItemView.EditKeyPressed)
        self.file_items = []
        for info in infos:
            top = QTreeWidgetItem(self.tree)
            top.setText(0, info["file"])
            top.setIcon(0, icon("globe" if info.get("web") else ("network" if info["remote"] else "file"), MUTED, 16))
            if info.get("web"):
                top.setToolTip(0, f"Baixado de {info['path']}")
            elif info["remote"]:
                top.setToolTip(0, "Arquivo de rede: será copiado para o cache local antes da leitura")
            top.setText(4, fmt_bytes(info["size"]))
            top.setData(0, Qt.UserRole, info["path"])
            if info["error"]:
                top.setText(1, "" + info["error"])
                top.setForeground(1, QColor("#B91C1C"))
                top.setFlags(Qt.ItemIsEnabled)
                self.file_items.append((top, info, []))
                continue
            top.setFlags(top.flags() | Qt.ItemIsUserCheckable | Qt.ItemIsAutoTristate)
            children = []
            for it in info["items"]:
                ch = QTreeWidgetItem(top)
                ch.setText(0, it["name"] + ("   (oculta)" if it.get("hidden") else ""))
                ch.setData(0, Qt.UserRole, it["name"])
                ch.setText(1, "?" if it.get("rows") is None else fmt_int(it["rows"]))
                ch.setText(2, "?" if it.get("cols") is None else str(it["cols"]))
                ch.setFlags(ch.flags() | Qt.ItemIsUserCheckable)
                if info["kind"] in ("excel", "excel_calamine", "csv"):
                    ch.setText(3, "1")
                    ch.setFlags(ch.flags() | Qt.ItemIsEditable)
                    ch.setToolTip(3, "Clique duas vezes para mudar (ex.: 3 se há título nas linhas 1 e 2)")
                ch.setCheckState(0, Qt.Unchecked if it.get("hidden") else Qt.Checked)
                children.append(ch)
            if len(info["items"]) == 1:
                top.setText(1, children[0].text(1))
                top.setText(2, children[0].text(2))
            top.setExpanded(len(info["items"]) > 1)
            self.file_items.append((top, info, children))
        v.addWidget(self.tree, 1)

        row = QHBoxLayout()
        row.addWidget(make_button("Marcar tudo", lambda: self._set_all(Qt.Checked)))
        row.addWidget(make_button("Desmarcar tudo", lambda: self._set_all(Qt.Unchecked)))
        row.addStretch()
        self.header_spin = QSpinBox()
        self.header_spin.setRange(1, 1000)
        self.header_spin.setValue(1)
        has_sheets = any(i["kind"] in ("excel", "excel_calamine", "csv") for i in infos if not i["error"])
        lbl = QLabel("Cabeçalho na linha")
        apply_btn = make_button("Aplicar às marcadas", self._apply_header)
        for w in (lbl, self.header_spin, apply_btn):
            w.setVisible(has_sheets)
            row.addWidget(w)
        self.header_spin.setToolTip("Use quando a planilha tem título ou linhas em branco antes dos nomes das colunas.")
        v.addLayout(row)
        self.summary = hint("")
        v.addWidget(self.summary)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("Carregar selecionados")
        bb.button(QDialogButtonBox.Ok).setProperty("primary", True)
        bb.button(QDialogButtonBox.Cancel).setText("Cancelar")
        bb.accepted.connect(self._accept)
        bb.rejected.connect(self.reject)
        v.addWidget(bb)
        self.tree.itemChanged.connect(lambda *_: self._update_summary())
        self._update_summary()
        polish(self)

    def _set_all(self, state):
        for top, info, children in self.file_items:
            for ch in children:
                ch.setCheckState(0, state)

    def _apply_header(self):
        for top, info, children in self.file_items:
            for ch in children:
                if ch.checkState(0) == Qt.Checked and ch.flags() & Qt.ItemIsEditable:
                    ch.setText(3, str(self.header_spin.value()))

    def headers(self) -> dict:
        out = {}
        for top, info, children in self.file_items:
            for ch in children:
                if not ch.flags() & Qt.ItemIsEditable:
                    continue
                try:
                    n = max(int(ch.text(3).strip()), 1)
                except ValueError:
                    n = 1
                if n != 1:
                    out.setdefault(info["path"], {})[ch.data(0, Qt.UserRole)] = n
        return out

    def selection(self) -> dict:
        out = {}
        for top, info, children in self.file_items:
            out[info["path"]] = [ch.data(0, Qt.UserRole) for ch in children if ch.checkState(0) == Qt.Checked]
        return out

    def _update_summary(self):
        n_items, n_rows, unknown, remote = 0, 0, False, 0
        for top, info, children in self.file_items:
            chosen = [ch for ch in children if ch.checkState(0) == Qt.Checked]
            if chosen and info["remote"]:
                remote += info["size"]
            by = {it["name"]: it for it in info["items"]}
            for ch in chosen:
                n_items += 1
                r = by[ch.data(0, Qt.UserRole)].get("rows")
                if r is None:
                    unknown = True
                else:
                    n_rows += r
        txt = f"{n_items} selecionada(s) · ≈ {fmt_int(n_rows)} linhas" + (" (+ abas sem estimativa)" if unknown else "")
        if remote:
            txt += f" · {fmt_bytes(remote)} a copiar do servidor (fica em cache para as próximas vezes)"
        self.summary.setText(txt)

    def _accept(self):
        if not any(self.selection().values()):
            QMessageBox.warning(self, APP_NAME, "Marque pelo menos uma aba ou tabela.")
            return
        self.accept()


class CodeDialog(QDialog):
    """Gera os scripts BD.py / etl_bronze.py / etl_prata.py / etl_ouro.py a partir do pipeline atual."""

    def __init__(self, win):
        super().__init__(win)
        self.win = win
        self.setWindowTitle("Gerar código Python")
        self.setWindowIcon(QIcon(svg_pixmap(icon_svg("logo", ""), 32)))
        self.resize(860, 620)
        st = win.settings
        self.project = collect_project(win.engine, win.store.steps, win.store.export_cols)
        v = QVBoxLayout(self)
        v.addWidget(hint("Gera os scripts do pipeline atual, um por camada "
                         "(BD.py, etl_bronze.py, etl_prata.py, etl_ouro.py e executar_tudo.py). Eles rodam só com "
                         "Python, sem o Refina, e dão o mesmo resultado do app."))
        form = QFormLayout()
        default_dir = str(Path.home() / "Documents" / "Refina_codigo")
        self.out_dir = QLineEdit(st.value("code_dir", default_dir))
        self.banco = QLineEdit(st.value("code_banco", str(Path(default_dir) / "pipeline.duckdb")))
        self.parquet = QLineEdit(st.value("code_parquet", str(Path(default_dir) / "Parquet")))
        for label, edit, pick in (("Salvar os scripts em", self.out_dir, "dir"),
                                  ("CAMINHO_BANCO (.duckdb)", self.banco, "file"),
                                  ("PASTA_PARQUET", self.parquet, "dir")):
            row = QHBoxLayout()
            row.addWidget(edit, 1)
            row.addWidget(make_button("Escolher…", lambda _=False, e=edit, k=pick: self._pick(e, k)))
            form.addRow(label, row)
        v.addLayout(form)

        v.addWidget(hint("Nome de cada tabela no banco (clique duas vezes para mudar). Use só letras, números e _."))
        self.table = QTableWidget(0, 3)
        self.table.setHorizontalHeaderLabels(["Dataset no app", "Camada", "Tabela no banco"])
        self.table.verticalHeader().setVisible(False)
        self.table.verticalHeader().setDefaultSectionSize(28)
        hh = self.table.horizontalHeader()
        hh.setSectionResizeMode(0, QHeaderView.Stretch)
        hh.setSectionResizeMode(1, QHeaderView.ResizeToContents)
        hh.setSectionResizeMode(2, QHeaderView.Stretch)
        self.keys = []
        saved = st.value("code_names", "{}")
        try:
            saved = json.loads(saved) if isinstance(saved, str) else {}
        except ValueError:
            saved = {}
        rows = [("bronze", n) for n, b in self.project["bronze"].items() if b.get("kind") != "sample"]
        seen = set()
        for q in self.project["queries"]:
            k = (q["layer"], q["name"])
            if k not in seen:
                seen.add(k)
                rows.append(k)
        for layer, name in rows:
            r = self.table.rowCount()
            self.table.insertRow(r)
            a = QTableWidgetItem(name)
            a.setFlags(Qt.ItemIsEnabled)
            b = QTableWidgetItem(layer_icon(layer), LAYER_LABEL[layer])
            b.setFlags(Qt.ItemIsEnabled)
            c = QTableWidgetItem(saved.get(f"{layer}|{name}") or default_table_name(name, layer))
            c.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable | Qt.ItemIsEditable)
            for col, it in enumerate((a, b, c)):
                self.table.setItem(r, col, it)
            self.keys.append((layer, name))
        v.addWidget(self.table, 1)
        skipped = [n for n, b in self.project["bronze"].items() if b.get("kind") == "sample"]
        notes = []
        if skipped:
            notes.append(f"Dados de exemplo não entram no código: {', '.join(skipped)}.")
        if self.project["problems"]:
            notes.append("Sem origem conhecida (reabra a receita para incluir): " + "; ".join(self.project["problems"]))
        if notes:
            v.addWidget(hint(" ".join(notes)))
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("Gerar arquivos")
        bb.button(QDialogButtonBox.Ok).setProperty("primary", True)
        bb.button(QDialogButtonBox.Cancel).setText("Cancelar")
        bb.accepted.connect(self._generate)
        bb.rejected.connect(self.reject)
        v.addWidget(bb)
        self.written = []

    def _pick(self, edit, kind):
        if kind == "dir":
            path = QFileDialog.getExistingDirectory(self, "Escolher pasta", edit.text())
        else:
            path, _ = QFileDialog.getSaveFileName(self, "Banco DuckDB", edit.text(), "DuckDB (*.duckdb)",
                                                  options=QFileDialog.Option.DontConfirmOverwrite)
        if path:
            edit.setText(str(Path(path)))

    def names(self):
        out = {}
        for r, key in enumerate(self.keys):
            out[key] = self.table.item(r, 2).text().strip()
        return out

    def _generate(self):
        names = self.names()
        bad = [n for n in names.values() if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", n)]
        if bad:
            return QMessageBox.warning(self, APP_NAME, "Nomes de tabela inválidos (use letras, números e _ e "
                                                       "comece com letra): " + ", ".join(bad))
        lower = [n.lower() for n in names.values()]
        dup = sorted({n for n in lower if lower.count(n) > 1})
        if dup:
            return QMessageBox.warning(self, APP_NAME, "Há tabelas com o mesmo nome: " + ", ".join(dup))
        out = Path(self.out_dir.text().strip() or ".")
        files = build_scripts(self.win.engine, self.project, names, self.banco.text().strip(),
                              self.parquet.text().strip())
        exist = [f for f in files if (out / f).exists()]
        if exist and QMessageBox.question(self, APP_NAME, f"Já existem em {out}:\n" + "\n".join(exist) +
                                          "\n\nSubstituir?") != QMessageBox.Yes:
            return
        out.mkdir(parents=True, exist_ok=True)
        for fname, content in files.items():
            (out / fname).write_text(content, encoding="utf-8")
        self.written = [out / f for f in files]
        st = self.win.settings
        st.setValue("code_dir", str(out))
        st.setValue("code_banco", self.banco.text().strip())
        st.setValue("code_parquet", self.parquet.text().strip())
        st.setValue("code_names", json.dumps({f"{l}|{n}": t for (l, n), t in names.items()}, ensure_ascii=False))
        self.accept()


# ============================================================================
# ESTADO GLOBAL E EXECUÇÃO EM SEGUNDO PLANO
# ============================================================================
class Store(QObject):
    changed = Signal()

    def __init__(self, engine: Engine):
        super().__init__()
        self.engine = engine
        self.data = {layer: {} for layer in LAYERS}
        self.steps = []
        self.export_cols = {}      # dataset → [{"col", "as"}] (só as colunas exportadas, na ordem)

    def available(self, layers) -> dict:
        return available(self.data, layers)

    def apply(self, step, results):
        apply_results(self.data, step, results)
        self.steps.append(step)
        for old in self.steps[:-UNDO_DEPTH]:     # passos antigos não podem mais ser desfeitos
            old["_prev"] = []
        self.changed.emit()

    def undo(self) -> bool:
        if not self.steps:
            return False
        step = self.steps.pop()
        for layer, name, old in reversed(step.get("_prev", [])):
            if old is None:
                self.data[layer].pop(name, None)
            else:
                self.data[layer][name] = old
        self.changed.emit()
        self.gc()
        return True

    def remove(self, layer, name):
        self.data[layer].pop(name, None)
        self.changed.emit()
        self.gc()

    def load(self, state, steps, export_cols=None):
        self.data, self.steps = state, steps
        self.export_cols = dict(export_cols or {})
        self.changed.emit()
        self.gc()

    def reset(self):
        self.data = {layer: {} for layer in LAYERS}
        self.steps = []
        self.export_cols = {}
        self.changed.emit()
        self.gc()

    def gc(self):
        """Apaga do disco as tabelas que nada mais usa."""
        keep = {tid for layer in LAYERS for tid in self.data[layer].values()}
        for s in self.steps:
            keep.update(tid for _, _, tid in s.get("_prev", []) if tid)
        for t in self.engine.tables():
            if t.startswith("t_") and t not in keep:
                self.engine.drop(t)

    def recipe_steps(self):
        return [clean_step(s) for s in self.steps]

    def recipe_export_cols(self):
        names = {n for layer in LAYERS for n in self.data[layer]}
        return {k: v for k, v in self.export_cols.items() if k in names}


class _WorkerSignals(QObject):
    done = Signal(object)
    failed = Signal(str, str)
    cancelled = Signal()
    progress = Signal(str, int)


class Worker(QRunnable):
    def __init__(self, fn, args):
        super().__init__()
        self.fn, self.args = fn, args
        self.signals = _WorkerSignals()
        self.ctx = Ctx(lambda msg, pct: self.signals.progress.emit(msg, pct))

    def run(self):
        try:
            result = self.fn(*self.args, self.ctx)
        except Cancelled:
            self.signals.cancelled.emit()
        except duckdb.InterruptException:
            self.signals.cancelled.emit()
        except Exception as exc:
            if self.ctx.cancelled:
                self.signals.cancelled.emit()
            else:
                self.signals.failed.emit(str(exc), traceback.format_exc())
        else:
            self.signals.done.emit(result)


# ============================================================================
# PÁGINAS
# ============================================================================
class Page(QWidget):
    def __init__(self, win, title, caption=""):
        super().__init__()
        self.setObjectName("page")
        self.win = win
        self.store = win.store
        self.engine = win.engine
        self._dirty = True
        self.body = QVBoxLayout(self)
        self.body.setContentsMargins(22, 16, 22, 16)
        self.body.setSpacing(8)
        head = QHBoxLayout()
        head.setSpacing(12)
        bar = QFrame()
        bar.setFixedSize(4, 22)
        bar.setStyleSheet(f"background: {getattr(self, 'ACCENT', PRIMARY)}; border-radius: 2px;")
        head.addWidget(bar)
        t = QLabel(split_emoji(title)[1])
        t.setObjectName("pageTitle")
        head.addWidget(t)
        head.addStretch()
        self.body.addLayout(head)
        if caption:
            c = QLabel(caption)
            c.setObjectName("pageCaption")
            c.setWordWrap(True)
            self.body.addWidget(c)
        self.store.changed.connect(self._on_change)

    def _on_change(self):
        self._dirty = True
        if self.isVisible():
            self._do_refresh()

    def showEvent(self, event):
        super().showEvent(event)
        if self._dirty:
            self._do_refresh()

    def _do_refresh(self):
        self._dirty = False
        self.refresh()

    def refresh(self):
        pass

    def warn(self, msg):
        QMessageBox.warning(self, APP_NAME, msg)

    def cols(self, layers, name, numeric=False):
        tid = self.store.available(layers).get(name)
        if not tid:
            return []
        if numeric:
            return [c for c, t in self.engine.types(tid).items() if is_numeric_type(t)]
        return self.engine.columns(tid)


class UploadPage(Page):
    ACCENT = LAYER_COLOR["bronze"]

    def __init__(self, win):
        super().__init__(win, "Ingestão de dados",
                         "Traga seus dados brutos para a camada Bronze. Você também pode arrastar "
                         "arquivos ou pastas direto para a janela (atalho: Ctrl+O para escolher arquivos).")
        cols = QHBoxLayout()
        cols.setSpacing(16)
        lc, rc = QVBoxLayout(), QVBoxLayout()
        cols.addLayout(lc, 1)
        cols.addLayout(rc, 1)
        self.body.addLayout(cols, 1)

        g = QGroupBox("Conectar a um arquivo, pasta ou link")
        v = QVBoxLayout(g)
        v.addWidget(hint("Aceita arquivo, pasta (local ou da rede) e link da internet (http/https), inclusive links "
                         "de Google Planilhas, Google Drive, OneDrive/SharePoint e Dropbox. Tudo fica em cache local e "
                         "só é baixado de novo quando mudar. O caminho ou link fica gravado na receita."))
        row = QHBoxLayout()
        self.path = QLineEdit()
        self.path.setPlaceholderText(r"C:\Users\...\Base.xlsx   ·   \\servidor\pasta   ·   https://site.com/dados.csv")
        self.path.returnPressed.connect(self._load_path)
        row.addWidget(self.path, 1)
        row.addWidget(make_button("Arquivo…", self._pick_file))
        row.addWidget(make_button("Pasta…", self._pick_dir))
        v.addLayout(row)
        row = QHBoxLayout()
        row.addWidget(make_button("🔄 Puxar dados do caminho", self._load_path, True))
        row.addWidget(make_button("📂 Escolher arquivos…", self._pick_files))
        row.addWidget(make_button("🎲 Dados de exemplo", self._sample))
        row.addStretch()
        v.addLayout(row)
        rec_row = QHBoxLayout()
        rec_lbl = QLabel("Recentes — clique duas vezes para carregar de novo")
        rec_lbl.setObjectName("hint")
        rec_row.addWidget(rec_lbl)
        rec_row.addStretch()
        self.clear_recent_btn = QPushButton("Limpar histórico")
        self.clear_recent_btn.setObjectName("linkButton")
        self.clear_recent_btn.setCursor(Qt.PointingHandCursor)
        self.clear_recent_btn.clicked.connect(self._clear_recent)
        rec_row.addWidget(self.clear_recent_btn)
        v.addLayout(rec_row)
        self.recent = QListWidget()
        self.recent.setObjectName("recent")
        self.recent.setMaximumHeight(150)
        self.recent.setToolTip("Clique para usar · clique duas vezes para carregar · Delete para remover do histórico")
        self.recent.installEventFilter(self)       # tecla Delete remove do histórico
        self.recent.setTextElideMode(Qt.ElideMiddle)
        self.recent.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.recent.itemDoubleClicked.connect(lambda it: self.win.load_paths([it.data(Qt.UserRole)]))
        self.recent.itemClicked.connect(lambda it: self.path.setText(it.data(Qt.UserRole)))
        v.addWidget(self.recent)
        lc.addWidget(g)

        g = QGroupBox("⚡ Desempenho e memória")
        form = QFormLayout(g)
        self.mem = QComboBox()
        ram = total_ram_mb()
        self.mem.addItem(f"Automático ({fmt_bytes(auto_mem_mb() * 2**20)} de {fmt_bytes(ram * 2**20)} de RAM)", 0)
        for mb in (512, 1024, 2048, 4096, 8192, 16384):
            if mb < ram:
                self.mem.addItem(fmt_bytes(mb * 2**20), mb)
        idx = self.mem.findData(win.settings.value("mem_mb", 0, int))
        self.mem.setCurrentIndex(max(idx, 0))
        self.mem.currentIndexChanged.connect(self._mem_changed)
        form.addRow("Memória máxima", self.mem)
        form.addRow("", hint("O app nunca passa desse limite: o que não cabe é processado em disco."))
        cache_row = QHBoxLayout()
        self.cache_lbl = QLabel()
        self.cache_lbl.setWordWrap(True)
        cache_row.addWidget(self.cache_lbl, 1)
        cache_row.addWidget(make_button("🧽 Limpar", self._clear_cache))
        form.addRow("Cache (rede e internet)", cache_row)
        lc.addWidget(g)
        lc.addStretch()

        g = QGroupBox("Datasets na Bronze")
        v = QVBoxLayout(g)
        self.list = QListWidget()
        self.list.setMinimumHeight(160)
        self.list.itemDoubleClicked.connect(lambda it: self.win.open_in_explorer("bronze", it.data(Qt.UserRole)))
        v.addWidget(self.list)
        row = QHBoxLayout()
        row.addWidget(make_button("🔎 Abrir no Explorar", self._open_selected))
        row.addWidget(make_button("Remover", self._remove))
        row.addStretch()
        v.addLayout(row)
        rc.addWidget(g)

        g = QGroupBox("📜 Histórico da operação")
        v = QVBoxLayout(g)
        v.addWidget(StepsLog(self.store))
        rc.addWidget(g, 1)
        self.refresh_recent()

    def refresh(self):
        self.list.clear()
        for name, tid in self.store.data["bronze"].items():
            m = self.engine.meta(tid)
            item = QListWidgetItem(f"{name}  —  {fmt_int(m['rows'])} linhas × {len(m['cols'])} colunas")
            item.setData(Qt.UserRole, name)
            self.list.addItem(item)
        self._refresh_cache()

    def refresh_recent(self):
        self.recent.clear()
        paths = self.win.recent_paths()
        for p in paths:
            kind = "globe" if is_url(p) else ("folder" if Path(p).suffix == "" else "file")
            it = QListWidgetItem(icon(kind, MUTED, 16), p)
            it.setData(Qt.UserRole, p)
            it.setToolTip(p)
            self.recent.addItem(it)
            # botão "x" à direita de cada linha
            holder = QWidget()
            holder.setStyleSheet("background: transparent;")
            hl = QHBoxLayout(holder)
            hl.setContentsMargins(0, 0, 4, 0)
            hl.addStretch()
            btn = QToolButton()
            btn.setObjectName("rowDelete")
            btn.setIcon(icon("x", "#928CA6", 14))
            btn.setIconSize(QSize(14, 14))
            btn.setAutoRaise(True)
            btn.setCursor(Qt.PointingHandCursor)
            btn.setToolTip("Remover do histórico")
            btn.clicked.connect(lambda _=False, p=p: self.win.remove_recent(p))
            hl.addWidget(btn)
            self.recent.setItemWidget(it, holder)
        if not paths:
            it = QListWidgetItem("Nenhum caminho usado ainda.")
            it.setFlags(Qt.NoItemFlags)
            self.recent.addItem(it)
        self.clear_recent_btn.setEnabled(bool(paths))
        self.recent.doItemsLayout()

    def eventFilter(self, obj, event):
        if obj is self.recent and event.type() == QEvent.KeyPress and event.key() in (Qt.Key_Delete, Qt.Key_Backspace):
            self._remove_selected_recent()
            return True
        return super().eventFilter(obj, event)

    def _remove_selected_recent(self):
        it = self.recent.currentItem()
        if it is not None and it.data(Qt.UserRole):
            row = self.recent.currentRow()
            self.win.remove_recent(it.data(Qt.UserRole))
            if self.recent.count() and self.recent.item(0).data(Qt.UserRole):
                self.recent.setCurrentRow(min(row, self.recent.count() - 1))

    def _clear_recent(self):
        if QMessageBox.question(self, APP_NAME, "Apagar todo o histórico de caminhos recentes?\n\n"
                                                "Os dados carregados e as receitas não são afetados.") == QMessageBox.Yes:
            self.win.clear_recent()

    def _refresh_cache(self):
        size, n = cache_usage()
        self.cache_lbl.setText(f"{fmt_bytes(size)} em {n} arquivo(s) · limite {DEFAULT_CACHE_GB} GB")

    def _mem_changed(self):
        mb = self.mem.currentData()
        self.win.settings.setValue("mem_mb", mb)
        self.engine.set_memory(mb)
        self.win.statusBar().showMessage(f"Memória máxima: {fmt_bytes(self.engine.mem_mb * 2**20)}", 5000)

    def _clear_cache(self):
        clear_cache()
        self._refresh_cache()

    def _pick_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Escolher arquivo", "", FILE_FILTER)
        if path:
            self.path.setText(str(Path(path)))

    def _pick_dir(self):
        path = QFileDialog.getExistingDirectory(self, "Escolher pasta")
        if path:
            self.path.setText(str(Path(path)))

    def _load_path(self):
        if not self.path.text().strip():
            return self.warn("Informe um caminho de arquivo ou pasta.")
        self.win.load_paths([self.path.text()])

    def pick_files(self):
        files, _ = QFileDialog.getOpenFileNames(self, "Escolher arquivos", "", FILE_FILTER)
        if files:
            self.win.load_paths(files)

    _pick_files = pick_files

    def _sample(self):
        self.win.run_step("Bronze", "load_sample", {}, "Ingestão exemplo", "Dados de exemplo (clientes + pedidos)")

    def _open_selected(self):
        item = self.list.currentItem()
        if item:
            self.win.open_in_explorer("bronze", item.data(Qt.UserRole))

    def _remove(self):
        item = self.list.currentItem()
        if item:
            self.store.remove("bronze", item.data(Qt.UserRole))


class ExplorerPage(Page):
    """Explora qualquer dataset das três camadas: dados, perfil das colunas, remover, exportar."""

    def __init__(self, win):
        super().__init__(win, "🔎 Explorar dados",
                         "Veja qualquer dataset da Bronze, Prata ou Ouro: os dados, o perfil de cada coluna "
                         "e atalhos para exportar ou remover.")
        row = QHBoxLayout()
        self.layer = QComboBox()
        self.layer.addItem("Todas as camadas", None)
        for layer in LAYERS:
            self.layer.addItem(layer_icon(layer), LAYER_LABEL[layer], layer)
        self.layer.currentIndexChanged.connect(lambda *_: self.refresh())
        self.combo = QComboBox()
        self.combo.setMinimumWidth(340)
        self.combo.currentIndexChanged.connect(self._show)
        row.addWidget(QLabel("Camada"))
        row.addWidget(self.layer)
        row.addSpacing(10)
        row.addWidget(QLabel("Dataset"))
        row.addWidget(self.combo, 1)
        row.addWidget(make_button("📦 Exportar…", self._export))
        row.addWidget(make_button("🗑️ Remover", self._remove))
        self.body.addLayout(row)
        self.empty = hint("Nenhum dado ainda. Vá para Upload e adicione arquivos.")
        self.body.addWidget(self.empty)
        row = QHBoxLayout()
        self.m_rows, self.m_cols, self.m_null, self.m_layer = (MetricCard("Linhas"), MetricCard("Colunas"),
                                                               MetricCard("Células vazias"), MetricCard("Camada"))
        for m in (self.m_rows, self.m_cols, self.m_null, self.m_layer):
            row.addWidget(m)
        self.body.addLayout(row)
        self.tabs = QTabWidget()
        self.view = DataView()
        prof = QWidget()
        pv = QVBoxLayout(prof)
        pv.addWidget(hint("Resumo de cada coluna: % de vazios, quantidade aproximada de valores distintos, "
                          "mínimo, máximo e média. Ajuda a achar colunas com problema antes de tratar."))
        self.profile = SmallTable(min_height=260)
        pv.addWidget(self.profile, 1)
        data_page = QWidget()
        dv = QVBoxLayout(data_page)
        dv.setContentsMargins(10, 10, 10, 10)
        dv.addWidget(self.view)
        self.tabs.addTab(data_page, "📄 Dados")
        self.tabs.addTab(prof, "📊 Perfil das colunas")
        self.tabs.currentChanged.connect(lambda *_: self._load_profile())
        self.body.addWidget(self.tabs, 1)
        self._profile_tid = None

    def _current(self):
        data = self.combo.currentData()
        if not data:
            return None, None, None
        layer, name = data
        return layer, name, self.store.data[layer].get(name)

    def refresh(self):
        current = self.combo.currentData()
        flt = self.layer.currentData()
        self.combo.blockSignals(True)
        self.combo.clear()
        for layer in LAYERS:
            if flt and layer != flt:
                continue
            for name in self.store.data[layer]:
                self.combo.addItem(layer_icon(layer), name, (layer, name))
        idx = next((k for k in range(self.combo.count()) if current and tuple(self.combo.itemData(k)) == tuple(current)), -1)
        if idx < 0 and self.combo.count():
            idx = self.combo.count() - 1          # o mais recente
        self.combo.setCurrentIndex(idx)
        self.combo.blockSignals(False)
        self._show()

    def select(self, layer, name):
        if self.layer.currentData() not in (None, layer):
            self.layer.blockSignals(True)
            self.layer.setCurrentIndex(0)
            self.layer.blockSignals(False)
        self.refresh()
        for k in range(self.combo.count()):
            if tuple(self.combo.itemData(k)) == (layer, name):
                self.combo.setCurrentIndex(k)
                break

    def filter_layer(self, layer):
        self.layer.setCurrentIndex(max(self.layer.findData(layer), 0))

    def _show(self, *_):
        layer, name, tid = self._current()
        self.empty.setVisible(tid is None)
        if tid is None:
            for m in (self.m_rows, self.m_cols, self.m_null, self.m_layer):
                m.value.setText("—")
            self.view.set_table(None, None)
            self.profile.set_df(None)
            self._profile_tid = None
            return
        m = self.engine.meta(tid)
        self.m_rows.value.setText(fmt_int(m["rows"]))
        self.m_cols.value.setText(str(len(m["cols"])))
        self.m_null.value.setText(fmt_int(sum(self.engine.null_counts(tid))))
        self.m_layer.value.setText(LAYER_LABEL[layer])
        self.m_layer.value.setStyleSheet(f"color: {LAYER_COLOR[layer]};")
        self.view.set_table(self.engine, tid)
        if self._profile_tid != tid:
            self.profile.set_df(None)
            self._profile_tid = None
        self._load_profile()

    def _load_profile(self):
        if self.tabs.currentIndex() != 1:
            return
        _, _, tid = self._current()
        if tid is None or self._profile_tid == tid:
            return

        def done(df, tid=tid):
            if self._current()[2] == tid:
                self._profile_tid = tid
                self.profile.set_df(df)

        if tid in self.engine._profile:
            done(self.engine._profile[tid])
        else:
            self.win.run_task(profile_table, (self.engine, tid), done, "Calculando o perfil das colunas…")

    def _export(self):
        layer, name, tid = self._current()
        if tid:
            self.win.goto_export(layer, name)

    def _remove(self):
        layer, name, tid = self._current()
        if tid is None:
            return
        if QMessageBox.question(self, APP_NAME, f"Remover o dataset '{name}' da camada {LAYER_LABEL[layer]}?\n\n"
                                                "Ele some da sessão (a receita continua igual).") == QMessageBox.Yes:
            self.store.remove(layer, name)


def split_with_preview(page, tabs):
    """Formulários à esquerda (com rolagem) e pré-visualização à direita."""
    split = QSplitter(Qt.Horizontal)
    scroll = QScrollArea()
    scroll.setWidgetResizable(True)
    scroll.setFrameShape(QFrame.NoFrame)
    scroll.setWidget(tabs)
    scroll.setMinimumWidth(470)
    split.addWidget(scroll)
    page.preview = PreviewPanel(page)
    split.addWidget(page.preview)
    split.setStretchFactor(0, 0)
    split.setStretchFactor(1, 1)
    split.setSizes([540, 760])
    split.setChildrenCollapsible(False)
    page.body.addWidget(split, 1)
    return split


class SilverPage(Page):
    LAYER_LABEL = "Prata"
    ACCENT = LAYER_COLOR["silver"]

    def __init__(self, win):
        super().__init__(win, "Prata: limpeza e tratamento",
                         "Cada operação gera um novo dataset na Prata, a partir da Bronze ou da própria Prata. "
                         "À direita você vê a origem escolhida e, depois de aplicar, o resultado.")
        self.tabs = QTabWidget()
        split_with_preview(self, self.tabs)
        self._build_clean()
        self._build_text()
        self._build_merge()
        self._build_filter()
        self.tabs.currentChanged.connect(lambda *_: self.preview_source())

    def _cols(self, name):
        return self.cols(SILVER_SOURCES, name)

    def _source_combo(self):
        return [self.cl_src, self.tx_src, self.mg_a, self.fl_src][self.tabs.currentIndex()]

    def preview_source(self, *_):
        name = self._source_combo().currentText()
        self.preview.show_dataset(layer_of(self.store, name, SILVER_SOURCES), name, "Origem: ")

    def on_result(self, step, results):
        if step["layer"] != self.LAYER_LABEL or not results:
            return
        layer, name, tid = results[-1]
        p = step["params"]
        src = p.get("src") or p.get("a")
        note = f"<span style='color:{PRIMARY}'>●</span> <b>Resultado</b> — "
        src_layer = layer_of(self.store, src, LAYERS) if src else None
        if src_layer and src != name:
            before = self.engine.rows(self.store.data[src_layer][src])
            after = self.engine.rows(tid)
            diff = after - before
            delta = f", {'+' if diff > 0 else '−'}{fmt_int(abs(diff))}" if diff else ""
            note = (f"<span style='color:{PRIMARY}'>●</span> <b>Resultado</b> "
                    f"({src}: {fmt_int(before)} → {fmt_int(after)} linhas{delta}) — ")
        self.preview.show_dataset(layer, name, note)

    def refresh(self):
        sources = list(self.store.available(SILVER_SOURCES))
        for combo in (self.cl_src, self.tx_src, self.mg_a, self.mg_b, self.fl_src):
            fill_combo(combo, sources)
        if (not self._mg_b_touched and self.mg_b.count() > 1
                and self.mg_b.currentText() == self.mg_a.currentText()):
            self.mg_b.setCurrentIndex(1)
        self._clean_cols()
        self._text_cols()
        self._merge_cols()
        self._filter_cols()
        self.preview_source()

    def _build_clean(self):
        w = QWidget()
        v = QVBoxLayout(w)
        form = QFormLayout()
        self.cl_src = QComboBox()
        self.cl_src.setMinimumWidth(320)
        self.cl_src.currentTextChanged.connect(self._clean_cols)
        self.cl_src.currentTextChanged.connect(self.preview_source)
        form.addRow("Dataset de origem", self.cl_src)
        v.addLayout(form)

        g = QGroupBox("Deduplicação")
        gv = QVBoxLayout(g)
        gv.addWidget(hint("Colunas chave (nenhuma marcada = considera todas as colunas). Mantém a primeira ocorrência."))
        self.dd_cols = CheckList()
        gv.addWidget(self.dd_cols)
        gv.addWidget(make_button("Remover duplicatas", self._dedup, True), 0, Qt.AlignLeft)
        v.addWidget(g)

        g = QGroupBox("Valores nulos")
        gf = QFormLayout(g)
        self.fill = QLineEdit()
        self.fill.setPlaceholderText("vazio = não preencher")
        self.dropna = QCheckBox("Remover linhas com nulos")
        gf.addRow("Preencher nulos com", self.fill)
        gf.addRow("", self.dropna)
        gf.addRow("", left(make_button("Aplicar nulos", self._nulls, True)))
        v.addWidget(g)

        g = QGroupBox("Tipagem de dados")
        gf = QFormLayout(g)
        self.cast_col = QComboBox()
        self.cast_type = QComboBox()
        self.cast_type.addItems(CAST_TYPES)
        gf.addRow("Coluna", self.cast_col)
        gf.addRow("Converter para", self.cast_type)
        gf.addRow("", hint("Números aceitam '1.234,56'; datas aceitam dd/mm/aaaa e números de série do Excel."))
        gf.addRow("", left(make_button("Converter tipo", self._cast, True)))
        v.addWidget(g)
        v.addStretch()
        self.tabs.addTab(w, "🧹 Limpeza")

    def _build_text(self):
        w = QWidget()
        v = QVBoxLayout(w)
        v.addWidget(hint("Espaços sobrando, o espaço 'invisível' que vem do Excel e diferenças de maiúsculas ou "
                         "acentos fazem o merge não encontrar chaves que parecem iguais. Padronize antes de mesclar."))
        form = QFormLayout()
        self.tx_src = QComboBox()
        self.tx_src.currentTextChanged.connect(self._text_cols)
        self.tx_src.currentTextChanged.connect(self.preview_source)
        form.addRow("Dataset de origem", self.tx_src)
        v.addLayout(form)
        g = QGroupBox("Colunas de texto")
        gv = QVBoxLayout(g)
        self.tx_cols = CheckList(max_height=170)
        gv.addWidget(self.tx_cols)
        r = QHBoxLayout()
        r.addWidget(make_button("Todas", lambda: self._tx_check(True)))
        r.addWidget(make_button("Nenhuma", lambda: self._tx_check(False)))
        r.addStretch()
        gv.addLayout(r)
        v.addWidget(g)
        g = QGroupBox("O que fazer")
        gv = QVBoxLayout(g)
        self.tx_trim = QCheckBox("Tirar espaços do começo e do fim")
        self.tx_spaces = QCheckBox("Trocar espaços repetidos por um só")
        self.tx_empty = QCheckBox("Texto vazio vira nulo")
        self.tx_accents = QCheckBox("Remover acentos (ação → acao)")
        for cb, on in ((self.tx_trim, True), (self.tx_spaces, True), (self.tx_empty, True), (self.tx_accents, False)):
            cb.setChecked(on)
            gv.addWidget(cb)
        r = QHBoxLayout()
        r.addWidget(QLabel("Maiúsculas/minúsculas"))
        self.tx_case = QComboBox()
        self.tx_case.addItems(list(TEXT_CASES))
        r.addWidget(self.tx_case)
        r.addStretch()
        gv.addLayout(r)
        v.addWidget(g)
        v.addWidget(make_button("✨ Padronizar texto", self._text_clean, True), 0, Qt.AlignLeft)
        v.addStretch()
        self.tabs.addTab(w, "✨ Texto")

    def _text_cols(self, *_):
        tid = self.store.available(SILVER_SOURCES).get(self.tx_src.currentText())
        cols = [c for c, t in self.engine.types(tid).items() if t == "VARCHAR"] if tid else []
        self.tx_cols.set_items(cols)

    def _tx_check(self, on):
        for i in range(self.tx_cols.count()):
            self.tx_cols.item(i).setCheckState(Qt.Checked if on else Qt.Unchecked)

    def _text_clean(self):
        src = self.tx_src.currentText()
        cols = self.tx_cols.checked()
        if not src:
            return self.warn("Carregue dados primeiro.")
        if not cols:
            return self.warn("Marque ao menos uma coluna de texto.")
        opts = {"trim": self.tx_trim.isChecked(), "spaces": self.tx_spaces.isChecked(),
                "accents": self.tx_accents.isChecked(), "case": TEXT_CASES[self.tx_case.currentText()],
                "empty_null": self.tx_empty.isChecked()}
        if not any([opts["trim"], opts["spaces"], opts["accents"], opts["case"], opts["empty_null"]]):
            return self.warn("Escolha ao menos uma ação.")
        acoes = [n for n, k in (("espaços", "trim"), ("espaços repetidos", "spaces"), ("acentos", "accents"),
                                ("vazio→nulo", "empty_null")) if opts[k]]
        if opts["case"]:
            acoes.append(self.tx_case.currentText())
        name = f"{src}_padronizado"
        self.win.run_step("Prata", "text_clean", {"src": src, "columns": cols, "name": name, **opts},
                          "Padronização de texto", f"{', '.join(cols)} ({', '.join(acoes)}) → {name}")

    def _clean_cols(self, *_):
        cols = self._cols(self.cl_src.currentText())
        self.dd_cols.set_items(cols)
        fill_combo(self.cast_col, cols)

    def _dedup(self):
        src = self.cl_src.currentText()
        if not src:
            return self.warn("Carregue dados primeiro.")
        cols = self.dd_cols.checked()
        desc = "Removidas duplicatas" + (f" por {', '.join(cols)}" if cols else "")
        self.win.run_step("Prata", "dedup", {"src": src, "subset": cols, "name": f"{src}_sem_duplicatas"},
                          "Deduplicação", f"{desc} → {src}_sem_duplicatas")

    def _nulls(self):
        src = self.cl_src.currentText()
        if not src:
            return self.warn("Carregue dados primeiro.")
        fill, drop = self.fill.text(), self.dropna.isChecked()
        if not fill and not drop:
            return self.warn("Informe um valor para preencher ou marque 'Remover linhas com nulos'.")
        partes = ([f"preenchidos com '{fill}'"] if fill else []) + (["linhas com nulos removidas"] if drop else [])
        self.win.run_step("Prata", "nulls", {"src": src, "fill": fill, "dropna": drop, "name": f"{src}_sem_nulos"},
                          "Tratamento de nulos", f"{', '.join(partes)} → {src}_sem_nulos")

    def _cast(self):
        src, col, typ = self.cl_src.currentText(), self.cast_col.currentText(), self.cast_type.currentText()
        if not src or not col:
            return self.warn("Escolha o dataset e a coluna.")
        self.win.run_step("Prata", "cast", {"src": src, "column": col, "type": typ, "name": f"{src}_tipado"},
                          "Tipagem", f"{col} → {typ} → {src}_tipado")

    def _build_merge(self):
        w = QWidget()
        v = QVBoxLayout(w)
        form = QFormLayout()
        self.mg_a, self.mg_b = QComboBox(), QComboBox()
        self.mg_how = QComboBox()
        self.mg_how.addItems(list(JOIN_TYPES))
        self.mg_left, self.mg_right = QComboBox(), QComboBox()
        self.mg_name = QLineEdit()
        self._mg_auto = ""
        self._mg_b_touched = False
        form.addRow("Tabela A (base)", self.mg_a)
        form.addRow("Tabela B (consulta)", self.mg_b)
        form.addRow("Tipo de join", self.mg_how)
        form.addRow("Chave — Tabela A", self.mg_left)
        form.addRow("Chave — Tabela B", self.mg_right)
        form.addRow("Nome do resultado", self.mg_name)
        v.addLayout(form)
        v.addWidget(hint("left: mantém tudo de A  ·  inner: só o que casa nas duas  ·  "
                         "full: tudo de A e de B  ·  cross: todas as combinações (sem chave). "
                         "Chaves de tipos diferentes (ex.: número × texto) são comparadas automaticamente."))
        r = QHBoxLayout()
        r.addWidget(make_button("🔍 Verificar chaves", self._check_merge))
        r.addWidget(make_button("🔗 Aplicar mesclagem", self._merge, True))
        r.addStretch()
        v.addLayout(r)
        self.mg_report = QLabel()
        self.mg_report.setWordWrap(True)
        self.mg_report.setTextFormat(Qt.RichText)
        self.mg_report.setObjectName("report")
        self.mg_report.setVisible(False)
        v.addWidget(self.mg_report)
        v.addStretch()
        self.mg_a.currentTextChanged.connect(self.preview_source)
        for wdg in (self.mg_a, self.mg_b, self.mg_left, self.mg_right, self.mg_how):
            wdg.currentTextChanged.connect(lambda *_: self.mg_report.setVisible(False))
        self.mg_a.currentTextChanged.connect(self._merge_cols)
        self.mg_b.currentTextChanged.connect(self._merge_cols)
        self.mg_b.activated.connect(lambda *_: setattr(self, "_mg_b_touched", True))
        self.mg_how.currentTextChanged.connect(self._merge_keys_enabled)
        self.tabs.addTab(w, "🔗 Merge")

    def _merge_cols(self, *_):
        a = self.mg_a.currentText()
        fill_combo(self.mg_left, self._cols(a))
        fill_combo(self.mg_right, self._cols(self.mg_b.currentText()))
        if a and self.mg_name.text() in ("", self._mg_auto):
            self._mg_auto = f"{a}_merge"
            self.mg_name.setText(self._mg_auto)
        self._merge_keys_enabled()

    def _merge_keys_enabled(self, *_):
        cross = self.mg_how.currentText() == "cross"
        self.mg_left.setEnabled(not cross)
        self.mg_right.setEnabled(not cross)

    def _check_merge(self):
        a, b, how = self.mg_a.currentText(), self.mg_b.currentText(), self.mg_how.currentText()
        if not a or not b:
            return self.warn("São necessários 2 datasets.")
        if how == "cross":
            return self.warn("O cross join não usa chaves: o resultado tem todas as combinações de A com B.")
        ds = self.store.available(SILVER_SOURCES)
        ta, tb = ds[a], ds[b]
        lk, rk = self.mg_left.currentText(), self.mg_right.currentText()

        def done(r):
            pct = 100 * r["matched"] / max(r["rows_a"], 1)
            cor = "#047857" if pct >= 90 else ("#B45309" if pct >= 50 else "#B91C1C")
            linhas = [f"<b>A</b> ({a}): {fmt_int(r['rows_a'])} linhas"
                      + (f", {fmt_int(r['null_a'])} com chave vazia" if r["null_a"] else ""),
                      f"<b>B</b> ({b}): {fmt_int(r['rows_b'])} linhas, {fmt_int(r['keys_b'])} chaves distintas",
                      f"<span style='color:{cor}'><b>{fmt_int(r['matched'])} linhas de A ({pct:.1f}%".replace(".", ",")
                      + ") encontram correspondência em B</b></span>"]
            if r["dup_b"]:
                um = r["dup_b"] == 1
                linhas.append(f"<b style='color:{WARN}'>Atenção:</b> {fmt_int(r['dup_b'])} {'chave aparece' if um else 'chaves aparecem'} mais de uma "
                              f"vez em B — as linhas de A com {'essa chave' if um else 'essas chaves'} vão se repetir "
                              f"no resultado. Se não for o esperado, deduplique B pela chave antes.")
            if r["matched"] == 0:
                linhas.append(f"<b style='color:{DANGER}'>Nenhuma chave casou.</b> Confira se são as colunas certas; se forem textos, "
                              "use <b>Texto › Padronizar texto</b> nas duas tabelas (espaços e maiúsculas costumam ser o motivo).")
            linhas.append(f"Resultado esperado do <b>{how}</b> join: <b>{fmt_int(r['result'])} linhas</b>.")
            self.mg_report.setText("<br>".join(linhas))
            self.mg_report.setVisible(True)

        self.win.run_task(merge_check, (self.engine, ta, tb, lk, rk, how), done, "Verificando as chaves…")

    def _merge(self):
        a, b, how = self.mg_a.currentText(), self.mg_b.currentText(), self.mg_how.currentText()
        name = self.mg_name.text().strip()
        if not a or not b:
            return self.warn("São necessários 2 datasets.")
        if not name:
            return self.warn("Dê um nome ao resultado.")
        cross = how == "cross"
        params = {"a": a, "b": b, "how": how, "name": name,
                  "left_on": None if cross else self.mg_left.currentText(),
                  "right_on": None if cross else self.mg_right.currentText()}
        chave = "" if cross else f" ({params['left_on']} = {params['right_on']})"
        self.win.run_step("Prata", "merge", params, "Mesclagem", f"{how} join de {a} com {b}{chave} → {name}")

    def _build_filter(self):
        w = QWidget()
        v = QVBoxLayout(w)
        form = QFormLayout()
        self.fl_src, self.fl_col, self.fl_op = QComboBox(), QComboBox(), QComboBox()
        self.fl_op.addItems(FILTER_OPERATORS)
        self.fl_val = QLineEdit()
        self.fl_val.setPlaceholderText("número, texto ou data (dd/mm/aaaa)")
        form.addRow("Dataset para filtrar", self.fl_src)
        form.addRow("Coluna", self.fl_col)
        form.addRow("Operador", self.fl_op)
        form.addRow("Valor", self.fl_val)
        v.addLayout(form)
        v.addWidget(make_button("🎯 Aplicar filtro", self._filter, True), 0, Qt.AlignLeft)
        v.addStretch()
        self.fl_src.currentTextChanged.connect(self._filter_cols)
        self.fl_src.currentTextChanged.connect(self.preview_source)
        self.fl_op.currentTextChanged.connect(self._filter_op_changed)
        self.tabs.addTab(w, "🎯 Filtros")

    def _filter_cols(self, *_):
        fill_combo(self.fl_col, self._cols(self.fl_src.currentText()))

    def _filter_op_changed(self, op):
        self.fl_val.setEnabled(op not in NULL_OPERATORS)
        if op in LIST_OPERATORS:
            self.fl_val.setPlaceholderText("valores separados por ; (ex.: ATIVO; PENDENTE; 1234)")
        else:
            self.fl_val.setPlaceholderText("número, texto ou data (dd/mm/aaaa)")

    def _filter(self):
        src, col, op, val = (self.fl_src.currentText(), self.fl_col.currentText(),
                             self.fl_op.currentText(), self.fl_val.text())
        if not src or not col:
            return self.warn("Escolha o dataset e a coluna.")
        self.win.run_step("Prata", "filter", {"src": src, "column": col, "operator": op, "value": val,
                                              "name": f"{src}_filtrado"},
                          "Filtro", f"{col} {op} '{val}' → {src}_filtrado")


class GoldPage(Page):
    LAYER_LABEL = "Ouro"
    ACCENT = LAYER_COLOR["gold"]

    def __init__(self, win):
        super().__init__(win, "Ouro: consolidação",
                         "Agregações e regras de negócio a partir da Prata, prontas para exportar.")
        self.tabs = QTabWidget()
        self.split = split_with_preview(self, self.tabs)
        self._build_agg()
        self._build_rules()
        self._build_export()
        self.tabs.currentChanged.connect(self._tab_changed)
        self.ag_src.currentTextChanged.connect(self.preview_source)
        self.rc_src.currentTextChanged.connect(self.preview_source)

    on_result = SilverPage.on_result

    def _tab_changed(self, idx):
        export = idx == 2
        self.preview.setVisible(not export)     # a aba Exportar já tem a própria pré-visualização
        self.preview_source()

    def preview_source(self, *_):
        idx = self.tabs.currentIndex()
        if idx == 2:
            return
        name = (self.ag_src if idx == 0 else self.rc_src).currentText()
        self.preview.show_dataset(layer_of(self.store, name, GOLD_SOURCES), name, "Origem: ")

    def refresh(self):
        sources = list(self.store.available(GOLD_SOURCES))
        fill_combo(self.ag_src, sources)
        fill_combo(self.rc_src, sources)
        self._agg_cols()
        self._rule_cols()
        self._refresh_export()
        self.preview_source()

    def _build_agg(self):
        w = QWidget()
        v = QVBoxLayout(w)
        form = QFormLayout()
        self.ag_src = QComboBox()
        self.ag_src.setMinimumWidth(320)
        self.ag_func = QComboBox()
        self.ag_func.addItems(AGG_FUNCTIONS)
        self.ag_group, self.ag_cols = CheckList(), CheckList()
        self.ag_name = QLineEdit()
        self._ag_auto = ""
        form.addRow("Origem (Prata/Ouro)", self.ag_src)
        form.addRow("Agrupar por (opcional)", self.ag_group)
        form.addRow("Função", self.ag_func)
        form.addRow("Colunas para calcular", self.ag_cols)
        form.addRow("Nome final", self.ag_name)
        v.addLayout(form)
        v.addWidget(make_button("📊 Executar agregação", self._aggregate, True), 0, Qt.AlignLeft)
        v.addStretch()
        self.ag_src.currentTextChanged.connect(self._agg_cols)
        self.ag_func.currentTextChanged.connect(self._agg_cols)
        self.tabs.addTab(w, "📊 Agregar")

    def _agg_cols(self, *_):
        src = self.ag_src.currentText()
        self.ag_group.set_items(self.cols(GOLD_SOURCES, src))
        self.ag_cols.set_items(self.cols(GOLD_SOURCES, src, numeric=self.ag_func.currentText() in ("sum", "mean")))
        if src and self.ag_name.text() in ("", self._ag_auto):
            self._ag_auto = f"{src}_agregado"
            self.ag_name.setText(self._ag_auto)

    def _aggregate(self):
        src, func, name = self.ag_src.currentText(), self.ag_func.currentText(), self.ag_name.text().strip()
        group = self.ag_group.checked()
        cols = [c for c in self.ag_cols.checked() if c not in group]
        if not src:
            return self.warn("Gere dados na Prata primeiro.")
        if not cols:
            return self.warn("Selecione ao menos uma coluna para calcular (que não esteja no agrupamento).")
        if not name:
            return self.warn("Dê um nome ao resultado.")
        por = f" agrupado por {', '.join(group)}" if group else ""
        self.win.run_step("Ouro", "aggregate", {"src": src, "group_by": group, "columns": cols, "func": func, "name": name},
                          "Agregação", f"{func} de {', '.join(cols)}{por} → {name}")

    def _build_rules(self):
        w = QWidget()
        v = QVBoxLayout(w)
        form = QFormLayout()
        self.rc_src, self.rc_a, self.rc_op, self.rc_b = QComboBox(), QComboBox(), QComboBox(), QComboBox()
        self.rc_op.addItems(["+", "-", "*", "/"])
        self.rc_new = QLineEdit("calculo_novo")
        form.addRow("Origem", self.rc_src)
        form.addRow("Coluna A", self.rc_a)
        form.addRow("Operador", self.rc_op)
        form.addRow("Coluna B", self.rc_b)
        form.addRow("Nova coluna", self.rc_new)
        v.addLayout(form)
        v.addWidget(hint("Divisão por zero resulta em vazio. O resultado é salvo como <origem>_math."))
        v.addWidget(make_button("🧮 Criar coluna", self._compute, True), 0, Qt.AlignLeft)
        v.addStretch()
        self.rc_src.currentTextChanged.connect(self._rule_cols)
        self.tabs.addTab(w, "🧮 Regras")

    def _rule_cols(self, *_):
        cols = self.cols(GOLD_SOURCES, self.rc_src.currentText())
        fill_combo(self.rc_a, cols)
        fill_combo(self.rc_b, cols)

    def _compute(self):
        src, a, op, b, new = (self.rc_src.currentText(), self.rc_a.currentText(), self.rc_op.currentText(),
                              self.rc_b.currentText(), self.rc_new.text().strip())
        if not src or not a or not b or not new:
            return self.warn("Preencha origem, colunas e o nome da nova coluna.")
        self.win.run_step("Ouro", "compute", {"src": src, "col_a": a, "operation": op, "col_b": b,
                                             "new_col": new, "name": f"{src}_math"},
                          "Matemática", f"{new} = {a} {op} {b} → {src}_math")

    def _build_export(self):
        w = QWidget()
        v = QVBoxLayout(w)
        row = QHBoxLayout()
        self.ex_combo = QComboBox()
        self.ex_combo.setMinimumWidth(360)
        self.ex_combo.currentIndexChanged.connect(self._show_export)
        row.addWidget(QLabel("Dataset"))
        row.addWidget(self.ex_combo, 1)
        row.addWidget(make_button("⬇️ Salvar Parquet…", lambda: self._export("parquet")))
        row.addWidget(make_button("🦆 Salvar no DuckDB…", self._export_duckdb))
        row.addWidget(make_button("⬇️ Salvar Excel (XLSX)…", lambda: self._export("xlsx"), True))
        v.addLayout(row)
        v.addWidget(hint("Escolha as colunas, renomeie (clique duas vezes no nome) e mude a ordem. A escolha vale "
                         "para este botão e também para a execução agendada. Acima de 1.048.575 linhas o Excel é "
                         "dividido em várias abas; para bases muito grandes, prefira Parquet."))

        split = QSplitter(Qt.Horizontal)
        panel = QGroupBox("Colunas do arquivo")
        pv = QVBoxLayout(panel)
        self.col_filter = QLineEdit()
        self.col_filter.setPlaceholderText("🔎 Filtrar colunas")
        self.col_filter.textChanged.connect(self._filter_col_rows)
        pv.addWidget(self.col_filter)
        self.col_table = QTableWidget(0, 3)
        self.col_table.setHorizontalHeaderLabels(["Coluna original", "Nome no arquivo", "Tipo"])
        self.col_table.verticalHeader().setVisible(False)
        self.col_table.verticalHeader().setDefaultSectionSize(26)
        self.col_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.col_table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.col_table.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.EditKeyPressed
                                       | QAbstractItemView.SelectedClicked)
        hh = self.col_table.horizontalHeader()
        hh.setSectionResizeMode(0, QHeaderView.Stretch)
        hh.setSectionResizeMode(1, QHeaderView.Stretch)
        hh.setSectionResizeMode(2, QHeaderView.ResizeToContents)
        self.col_table.itemChanged.connect(self._col_item_changed)
        pv.addWidget(self.col_table, 1)
        r1 = QHBoxLayout()
        up = make_button("▲", lambda: self._move_col(-1))
        up.setToolTip("Mover a coluna selecionada para cima")
        down = make_button("▼", lambda: self._move_col(1))
        down.setToolTip("Mover a coluna selecionada para baixo")
        for b in (up, down):
            b.setFixedWidth(40)
            r1.addWidget(b)
        r1.addWidget(make_button("Todas", lambda: self._check_all(True)))
        r1.addWidget(make_button("Nenhuma", lambda: self._check_all(False)))
        r1.addWidget(make_button("Restaurar", self._reset_cols))
        pv.addLayout(r1)
        self.col_info = hint("")
        pv.addWidget(self.col_info)
        pv.addWidget(make_button("➕ Salvar como novo dataset na Ouro", self._save_as_dataset))
        panel.setMinimumWidth(360)
        split.addWidget(panel)
        self.ex_view = DataView()
        split.addWidget(self.ex_view)
        split.setStretchFactor(0, 0)
        split.setStretchFactor(1, 1)
        split.setSizes([400, 900])
        v.addWidget(split, 1)
        # a aba Exportar é larga; não deixa ela forçar a largura mínima das outras abas
        w.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)
        self.cfg, self.cfg_name = [], None
        self._preview_timer = QTimer(self)
        self._preview_timer.setSingleShot(True)
        self._preview_timer.setInterval(250)
        self._preview_timer.timeout.connect(self._update_preview)
        self.tabs.addTab(w, "📦 Exportar")

    def _refresh_export(self):
        current = self.ex_combo.currentData()
        self.ex_combo.blockSignals(True)
        self.ex_combo.clear()
        for layer in ("gold", "silver", "bronze"):
            for name in self.store.data[layer]:
                self.ex_combo.addItem(layer_icon(layer), name, (layer, name))
        idx = self.find_export(current)
        self.ex_combo.setCurrentIndex(max(idx, 0) if self.ex_combo.count() else -1)
        self.ex_combo.blockSignals(False)
        self._show_export()

    def find_export(self, data) -> int:
        """Posição do dataset (camada, nome) no combo (findData não compara tuplas)."""
        if data:
            for k in range(self.ex_combo.count()):
                if tuple(self.ex_combo.itemData(k)) == tuple(data):
                    return k
        return -1

    def select_export(self, layer, name):
        idx = self.find_export((layer, name))
        if idx >= 0:
            self.ex_combo.setCurrentIndex(idx)

    def _current_export(self):
        data = self.ex_combo.currentData()
        if not data:
            return None, None
        return data[1], self.store.data[data[0]].get(data[1])

    # --- configuração de colunas ------------------------------------------------
    def _show_export(self, *_):
        name, tid = self._current_export()
        self.cfg_name = name
        self.cfg = []
        if tid:
            types = self.engine.types(tid)
            saved = self.store.export_cols.get(name)
            if saved:
                by_lower = {c.lower(): c for c in types}
                used = set()
                for item in saved:
                    col = by_lower.get(str(item["col"]).lower())
                    if col and col not in used:
                        self.cfg.append({"col": col, "as": item.get("as") or col, "on": True, "type": types[col]})
                        used.add(col)
                self.cfg += [{"col": c, "as": c, "on": False, "type": t} for c, t in types.items() if c not in used]
            else:
                self.cfg = [{"col": c, "as": c, "on": True, "type": t} for c, t in types.items()]
        self._fill_col_table()
        self.ex_view.set_table(self.engine if tid else None, tid, self._projection())
        self._update_col_info()

    def _fill_col_table(self, select_row=None):
        t = self.col_table
        t.blockSignals(True)
        t.setRowCount(len(self.cfg))
        for i, c in enumerate(self.cfg):
            orig = QTableWidgetItem(c["col"])
            orig.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable | Qt.ItemIsUserCheckable)
            orig.setCheckState(Qt.Checked if c["on"] else Qt.Unchecked)
            new = QTableWidgetItem(c["as"])
            new.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable | Qt.ItemIsEditable)
            if c["as"] != c["col"]:
                f = new.font()
                f.setBold(True)
                new.setFont(f)
            typ = QTableWidgetItem(c["type"])
            typ.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            typ.setForeground(QColor("#6E6884"))
            t.setItem(i, 0, orig)
            t.setItem(i, 1, new)
            t.setItem(i, 2, typ)
        t.blockSignals(False)
        self._filter_col_rows()
        self._mark_duplicates()
        if select_row is not None:
            t.selectRow(select_row)

    def _filter_col_rows(self, *_):
        txt = self.col_filter.text().strip().lower()
        for i, c in enumerate(self.cfg):
            hide = bool(txt) and txt not in c["col"].lower() and txt not in c["as"].lower()
            self.col_table.setRowHidden(i, hide)

    def _col_item_changed(self, item):
        i = item.row()
        if i >= len(self.cfg):
            return
        if item.column() == 0:
            self.cfg[i]["on"] = item.checkState() == Qt.Checked
        elif item.column() == 1:
            new = " ".join(item.text().split())
            self.cfg[i]["as"] = new or self.cfg[i]["col"]
            self.col_table.blockSignals(True)
            item.setText(self.cfg[i]["as"])
            f = item.font()
            f.setBold(self.cfg[i]["as"] != self.cfg[i]["col"])
            item.setFont(f)
            self.col_table.blockSignals(False)
        self._mark_duplicates()
        self._config_changed()

    def _duplicates(self):
        seen, dup = {}, set()
        for c in self.cfg:
            if c["on"]:
                k = c["as"].lower()
                if k in seen:
                    dup.add(k)
                seen[k] = True
        return dup

    def _mark_duplicates(self):
        dup = self._duplicates()
        self.col_table.blockSignals(True)
        for i, c in enumerate(self.cfg):
            it = self.col_table.item(i, 1)
            if it is not None:
                it.setBackground(QColor("#FEE2E2") if c["on"] and c["as"].lower() in dup else QColor("#FFFFFF"))
        self.col_table.blockSignals(False)

    def _move_col(self, delta):
        rows = self.col_table.selectionModel().selectedRows()
        if not rows:
            return
        i = rows[0].row()
        j = i + delta
        if 0 <= j < len(self.cfg):
            self.cfg[i], self.cfg[j] = self.cfg[j], self.cfg[i]
            self._fill_col_table(select_row=j)
            self._config_changed()

    def _check_all(self, on):
        for i, c in enumerate(self.cfg):
            if not self.col_table.isRowHidden(i):     # respeita o filtro de colunas
                c["on"] = on
        self._fill_col_table()
        self._config_changed()

    def _reset_cols(self):
        _, tid = self._current_export()
        if not tid:
            return
        self.cfg = [{"col": c, "as": c, "on": True, "type": t} for c, t in self.engine.types(tid).items()]
        self._fill_col_table()
        self._config_changed()

    def _config_changed(self):
        name, tid = self._current_export()
        if name is None:
            return
        spec = [{"col": c["col"], "as": c["as"]} for c in self.cfg if c["on"]]
        default = [{"col": c, "as": c} for c in self.engine.columns(tid)]
        if spec == default:
            self.store.export_cols.pop(name, None)
        else:
            self.store.export_cols[name] = spec
        self._update_col_info()
        self._preview_timer.start()

    def _projection(self):
        if not self.cfg or self._duplicates():
            return None
        proj = [(c["col"], c["as"]) for c in self.cfg if c["on"]]
        if not proj or proj == [(c["col"], c["col"]) for c in self.cfg]:
            return None
        return proj

    def _update_preview(self):
        _, tid = self._current_export()
        if tid and not self._duplicates() and any(c["on"] for c in self.cfg):
            self.ex_view.set_projection(self._projection())

    def _update_col_info(self):
        on = sum(1 for c in self.cfg if c["on"])
        ren = sum(1 for c in self.cfg if c["on"] and c["as"] != c["col"])
        txt = f"{on} de {len(self.cfg)} colunas no arquivo"
        if ren:
            txt += f" · {ren} renomeada(s)"
        dup = self._duplicates()
        if dup:
            txt += " · nomes repetidos: " + ", ".join(sorted(dup))
        elif not on and self.cfg:
            txt += " · marque ao menos uma coluna"
        self.col_info.setText(txt)

    def _export_columns(self):
        """Valida e devolve [(coluna, nome)] para exportar, ou None se houver problema."""
        if not any(c["on"] for c in self.cfg):
            self.warn("Marque ao menos uma coluna para exportar.")
            return None
        dup = self._duplicates()
        if dup:
            self.warn("Há nomes repetidos entre as colunas exportadas: " + ", ".join(sorted(dup)))
            return None
        return [(c["col"], c["as"]) for c in self.cfg if c["on"]]

    def _save_as_dataset(self):
        name, tid = self._current_export()
        if tid is None:
            return self.warn("Nenhum dataset selecionado.")
        cols = self._export_columns()
        if cols is None:
            return
        new = f"{name}_colunas"
        desc = f"{len(cols)} de {len(self.cfg)} colunas"
        ren = [f"{c}→{a}" for c, a in cols if c != a]
        if ren:
            desc += f", renomeadas: {', '.join(ren)}"
        if len(desc) > 200:
            desc = desc[:197] + "…"
        self.win.run_step("Ouro", "select_cols", {"src": name, "columns": [{"col": c, "as": a} for c, a in cols],
                                                  "name": new},
                          "Seleção de colunas", f"{desc} → {new}")

    def _export_duckdb(self):
        name, tid = self._current_export()
        if tid is None:
            return self.warn("Nenhum dataset para exportar.")
        cols = self._export_columns()
        if cols is None:
            return
        if cols == [(c, c) for c in self.engine.columns(tid)]:
            cols = None
        start = self.win.settings.value("last_duckdb", "dados.duckdb")
        path, _ = QFileDialog.getSaveFileName(self, "Gravar em qual banco DuckDB? (se já existir, a tabela é "
                                                    "adicionada ou atualizada)", start, "DuckDB (*.duckdb)",
                                              options=QFileDialog.Option.DontConfirmOverwrite)
        if not path:
            return
        if not path.lower().endswith((".duckdb", ".db")):
            path += ".duckdb"
        self.win.settings.setValue("last_duckdb", path)
        self.win.run_task(write_duckdb, (self.engine, [(name, tid, cols)], path),
                          lambda p: self.win.statusBar().showMessage(f"Tabela '{name}' gravada em {p}", 10000),
                          "Gravando no DuckDB…")

    def _export(self, kind):
        name, tid = self._current_export()
        if tid is None:
            return self.warn("Nenhum dataset para exportar.")
        cols = self._export_columns()
        if cols is None:
            return
        if cols == [(c, c) for c in self.engine.columns(tid)]:
            cols = None
        flt = "Parquet (*.parquet)" if kind == "parquet" else "Excel (*.xlsx)"
        path, _ = QFileDialog.getSaveFileName(self, "Salvar como", f"{_safe_filename(name)}.{kind}", flt)
        if not path:
            return
        fn = write_parquet if kind == "parquet" else write_xlsx
        self.win.run_task(fn, (self.engine, tid, path, cols),
                          lambda p: self.win.statusBar().showMessage(f"Arquivo salvo: {p}", 10000),
                          "Exportando…")


class CreateTableDialog(QDialog):
    """Nova tabela no banco: colunas e tipos (estilo BD.py) ou copiados de um dataset."""

    def __init__(self, page):
        super().__init__(page)
        self.page = page
        self.setWindowTitle("Nova tabela")
        self.setWindowIcon(QIcon(svg_pixmap(icon_svg("logo", ""), 32)))
        self.resize(720, 600)
        v = QVBoxLayout(self)
        form = QFormLayout()
        self.name = QLineEdit()
        self.name.setPlaceholderText("Ex.: Vendas_bronze_2024")
        self.src = QComboBox()
        self.src.addItem("Nenhum — definir as colunas abaixo", None)
        for layer in LAYERS:
            for n in page.store.data[layer]:
                self.src.addItem(layer_icon(layer), n, (layer, n))
        form.addRow("Nome da tabela", self.name)
        form.addRow("Copiar colunas de", self.src)
        v.addLayout(form)
        self.grid = QTableWidget(0, 2)
        self.grid.setHorizontalHeaderLabels(["Coluna", "Tipo"])
        self.grid.verticalHeader().setVisible(False)
        self.grid.verticalHeader().setDefaultSectionSize(30)
        self.grid.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.grid.horizontalHeader().setSectionResizeMode(1, QHeaderView.Fixed)
        self.grid.setColumnWidth(1, 170)
        v.addWidget(self.grid, 1)
        row = QHBoxLayout()
        row.addWidget(make_button("➕ Coluna", lambda: self._add_row("", "TEXT")))
        row.addWidget(make_button("Remover coluna", self._remove_row))
        row.addStretch()
        v.addLayout(row)
        self.fill = QCheckBox("Inserir também os dados desse dataset na nova tabela")
        self.fill.setEnabled(False)
        self.record = QCheckBox("Registrar na receita (repetir na execução agendada)")
        self.record.setChecked(True)
        v.addWidget(self.fill)
        v.addWidget(self.record)
        v.addWidget(hint("SQL que será executado:"))
        self.preview = QPlainTextEdit()
        self.preview.setReadOnly(True)
        self.preview.setObjectName("code")
        self.preview.setMaximumHeight(130)
        v.addWidget(self.preview)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("Criar tabela")
        bb.button(QDialogButtonBox.Ok).setProperty("primary", True)
        bb.button(QDialogButtonBox.Cancel).setText("Cancelar")
        bb.accepted.connect(self._accept)
        bb.rejected.connect(self.reject)
        v.addWidget(bb)
        self.src.currentIndexChanged.connect(self._from_dataset)
        self.name.textChanged.connect(self._update_preview)
        self.grid.itemChanged.connect(self._update_preview)
        for _ in range(3):
            self._add_row("", "TEXT")

    def _add_row(self, name, typ):
        r = self.grid.rowCount()
        self.grid.blockSignals(True)
        self.grid.insertRow(r)
        self.grid.setItem(r, 0, QTableWidgetItem(name))
        self.grid.blockSignals(False)
        combo = QComboBox()
        combo.setEditable(True)
        combo.addItems(DB_TYPES)
        combo.setCurrentText(typ)
        combo.currentTextChanged.connect(self._update_preview)
        self.grid.setCellWidget(r, 1, combo)
        self._update_preview()

    def _remove_row(self):
        r = self.grid.currentRow()
        if r < 0:
            r = self.grid.rowCount() - 1
        if r >= 0:
            self.grid.removeRow(r)
            self._update_preview()

    def _from_dataset(self, *_):
        data = self.src.currentData()
        self.fill.setEnabled(bool(data))
        if not data:
            self.fill.setChecked(False)
            return
        tid = self.page.store.data[data[0]].get(data[1])
        self.grid.setRowCount(0)
        for c, t in self.page.engine.types(tid).items():
            self._add_row(c, db_type_for(t))
        if not self.name.text().strip():
            self.name.setText(default_table_name(data[1], data[0]))
        self.fill.setChecked(True)

    def columns(self):
        out = []
        for r in range(self.grid.rowCount()):
            it = self.grid.item(r, 0)
            name = it.text().strip() if it else ""
            combo = self.grid.cellWidget(r, 1)
            if name:
                out.append({"name": name, "type": (combo.currentText() if combo else "TEXT").strip().upper()})
        return out

    def _update_preview(self, *_):
        cols = self.columns()
        body = ",\n".join(f"    {qi(c['name'])} {c['type']}" for c in cols) or "    -- adicione colunas"
        self.preview.setPlainText(f"CREATE TABLE IF NOT EXISTS {qi(self.name.text().strip() or 'nova_tabela')} (\n"
                                  f"{body}\n)")

    def _accept(self):
        if not self.name.text().strip():
            return QMessageBox.warning(self, APP_NAME, "Dê um nome à tabela.")
        cols = self.columns()
        if not cols:
            return QMessageBox.warning(self, APP_NAME, "Adicione pelo menos uma coluna.")
        low = [c["name"].lower() for c in cols]
        dup = sorted({n for n in low if low.count(n) > 1})
        if dup:
            return QMessageBox.warning(self, APP_NAME, "Colunas repetidas: " + ", ".join(dup))
        bad = [c["name"] for c in cols if not _TYPE_RE.match(c["type"])]
        if bad:
            return QMessageBox.warning(self, APP_NAME, "Tipo inválido nas colunas: " + ", ".join(bad))
        self.accept()


class InsertDialog(QDialog):
    """Grava um dataset do app numa tabela do banco."""

    def __init__(self, page, table=None):
        super().__init__(page)
        self.page = page
        self.setWindowTitle("Inserir dados no banco")
        self.setWindowIcon(QIcon(svg_pixmap(icon_svg("logo", ""), 32)))
        self.resize(660, 470)
        v = QVBoxLayout(self)
        form = QFormLayout()
        self.src = QComboBox()
        for layer in ("gold", "silver", "bronze"):
            for n in page.store.data[layer]:
                self.src.addItem(layer_icon(layer), n, (layer, n))
        self.dest = QComboBox()
        self.dest.setEditable(True)
        self.dest.addItems(page.table_names())
        if table:
            self.dest.setCurrentText(table)
        self.dest.lineEdit().setPlaceholderText("tabela existente ou nome de uma nova")
        self.mode = QComboBox()
        for k, label in DB_WRITE_MODES.items():
            self.mode.addItem(label, k)
        form.addRow("Dataset de origem", self.src)
        form.addRow("Tabela de destino", self.dest)
        form.addRow("Como gravar", self.mode)
        v.addLayout(form)
        self.try_cast = QCheckBox("Valores incompatíveis com o tipo da coluna viram vazio (em vez de dar erro)")
        self.record = QCheckBox("Registrar na receita (repetir na execução agendada)")
        self.record.setChecked(True)
        v.addWidget(self.try_cast)
        v.addWidget(self.record)
        self.report = QLabel()
        self.report.setObjectName("report")
        self.report.setWordWrap(True)
        self.report.setTextFormat(Qt.RichText)
        v.addWidget(self.report)
        v.addStretch()
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("Gravar no banco")
        bb.button(QDialogButtonBox.Ok).setProperty("primary", True)
        bb.button(QDialogButtonBox.Cancel).setText("Cancelar")
        bb.accepted.connect(self._accept)
        bb.rejected.connect(self.reject)
        v.addWidget(bb)
        for w in (self.src, self.mode):
            w.currentIndexChanged.connect(self._update)
        self.dest.currentTextChanged.connect(self._update)
        self._update()

    def _update(self, *_):
        data = self.src.currentData()
        table = self.dest.currentText().strip()
        if not data or not table:
            self.report.setText("Escolha o dataset e a tabela de destino.")
            return
        tid = self.page.store.data[data[0]].get(data[1])
        scols = self.page.engine.columns(tid)
        rows = self.page.engine.rows(tid)
        mode = self.mode.currentData()
        tcols = self.page.table_columns(table)
        if tcols is None or mode == "replace":
            acao = "criada" if tcols is None else "recriada"
            self.report.setText(f"A tabela <b>{table}</b> será {acao} com as {len(scols)} colunas do dataset e "
                                f"{fmt_int(rows)} linhas.")
            return
        tlow = {t.lower() for t in tcols}
        slow = {c.lower() for c in scols}
        common = [t for t in tcols if t.lower() in slow]
        ignored = [c for c in scols if c.lower() not in tlow]
        missing = [t for t in tcols if t.lower() not in slow]
        cor = PRIMARY if common else DANGER
        txt = [f"<b style='color:{cor}'>{len(common)} de {len(tcols)} colunas da tabela</b> recebem dados "
               f"({fmt_int(rows)} linhas)."]
        if mode == "truncate":
            txt.append("O conteúdo atual da tabela é apagado antes (se der erro, nada muda).")
        if ignored:
            txt.append("Colunas do dataset que não existem na tabela (ignoradas): " + ", ".join(ignored[:12])
                       + ("…" if len(ignored) > 12 else ""))
        if missing:
            txt.append("Colunas da tabela sem dado no dataset (ficam vazias): " + ", ".join(missing[:12])
                       + ("…" if len(missing) > 12 else ""))
        self.report.setText("<br>".join(txt))

    def _accept(self):
        if not self.src.currentData():
            return QMessageBox.warning(self, APP_NAME, "Não há datasets para gravar. Carregue dados primeiro.")
        if not self.dest.currentText().strip():
            return QMessageBox.warning(self, APP_NAME, "Informe a tabela de destino.")
        self.accept()


class DatabasePage(Page):
    """Conecta um banco DuckDB, mostra as tabelas, cria tabelas, grava datasets e roda SQL."""

    def __init__(self, win):
        super().__init__(win, "Banco de dados",
                         "Conecte um banco DuckDB (.duckdb) para criar tabelas, gravar datasets e rodar SQL. "
                         "O que for registrado na receita roda também na execução agendada.")
        self.path_open = None
        g = QGroupBox("Conexão")
        gv = QVBoxLayout(g)
        row = QHBoxLayout()
        self.db_path = QLineEdit(win.settings.value("db_path", ""))
        self.db_path.setPlaceholderText(r"C:\...\pipeline.duckdb   (se não existir, é criado)")
        self.db_path.returnPressed.connect(self._toggle)
        row.addWidget(self.db_path, 1)
        row.addWidget(make_button("Escolher…", self._pick))
        self.conn_btn = make_button("🔗 Conectar", self._toggle, True)
        row.addWidget(self.conn_btn)
        gv.addLayout(row)
        self.status = hint("Nenhum banco conectado.")
        gv.addWidget(self.status)
        self.body.addWidget(g)

        split = QSplitter(Qt.Horizontal)
        left = QGroupBox("Tabelas")
        lv = QVBoxLayout(left)
        self.tables = QListWidget()
        self.tables.setTextElideMode(Qt.ElideRight)
        self.tables.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.tables.currentItemChanged.connect(lambda *_: self._show_table())
        lv.addWidget(self.tables, 1)
        r1 = QHBoxLayout()
        self.btn_new = make_button("➕ Nova tabela…", self._new_table)
        self.btn_ins = make_button("⬇️ Inserir dados…", self._insert, True)
        r1.addWidget(self.btn_new)
        r1.addWidget(self.btn_ins)
        lv.addLayout(r1)
        r2 = QHBoxLayout()
        self.btn_ref = make_button("🔄 Atualizar", self.refresh_tables)
        self.btn_del = make_button("🗑️ Excluir tabela", self._drop)
        r2.addWidget(self.btn_ref)
        r2.addWidget(self.btn_del)
        lv.addLayout(r2)
        left.setMinimumWidth(320)
        split.addWidget(left)

        self.tabs = QTabWidget()
        data_page = QWidget()
        dv = QVBoxLayout(data_page)
        dv.setContentsMargins(10, 10, 10, 10)
        self.view = DataView(min_height=260)
        dv.addWidget(self.view)
        self.tabs.addTab(data_page, "📄 Dados")
        self.struct = SmallTable(min_height=200)
        sp = QWidget()
        sv = QVBoxLayout(sp)
        sv.setContentsMargins(10, 10, 10, 10)
        sv.addWidget(self.struct)
        self.tabs.addTab(sp, "📊 Estrutura")
        sql_page = QWidget()
        qv = QVBoxLayout(sql_page)
        qv.setContentsMargins(10, 10, 10, 10)
        qv.addWidget(hint("SQL do DuckDB. Nomes de tabela sem prefixo se referem às tabelas deste banco. "
                          "Vários comandos podem ser separados por ponto e vírgula. Ctrl+Enter executa."))
        self.sql = QPlainTextEdit()
        self.sql.setObjectName("code")
        self.sql.setPlaceholderText("-- Ex.:\nCREATE OR REPLACE TABLE Vendas_prata AS\n"
                                    "SELECT * FROM Vendas_2023_bronze UNION ALL BY NAME SELECT * FROM Vendas_2024_bronze;\n\n"
                                    "SELECT * FROM Vendas_prata LIMIT 100")
        self.sql.setMinimumHeight(120)
        qv.addWidget(self.sql, 1)
        qr = QHBoxLayout()
        qr.addWidget(make_button("▶ Executar", self._run_sql, True))
        self.sql_record = QCheckBox("Registrar na receita")
        self.sql_record.setToolTip("O SQL roda de novo a cada execução da receita (cuidado com INSERT, que duplica linhas).")
        qr.addWidget(self.sql_record)
        qr.addStretch()
        self.sql_info = hint("")
        qr.addWidget(self.sql_info)
        qv.addLayout(qr)
        self.sql_result = SmallTable(min_height=140)
        qv.addWidget(self.sql_result, 1)
        run_sc = QShortcut(QKeySequence("Ctrl+Return"), self.sql)
        run_sc.setContext(Qt.WidgetShortcut)
        run_sc.activated.connect(self._run_sql)
        self.tabs.addTab(sql_page, "SQL")
        split.addWidget(self.tabs)
        split.setStretchFactor(0, 0)
        split.setStretchFactor(1, 1)
        split.setSizes([340, 900])
        split.setChildrenCollapsible(False)
        self.body.addWidget(split, 1)
        self._set_connected(False)

    # --- conexão ---------------------------------------------------------------
    def _pick(self):
        path, _ = QFileDialog.getSaveFileName(self, "Banco DuckDB (escolha um existente ou digite um nome novo)",
                                              self.db_path.text() or "pipeline.duckdb", "DuckDB (*.duckdb *.db)",
                                              options=QFileDialog.Option.DontConfirmOverwrite)
        if path:
            self.db_path.setText(str(Path(path)))
            if self.path_open:
                self._disconnect()
            self._connect()

    def _toggle(self):
        if self.path_open:
            self._disconnect()
        else:
            self._connect()

    def _connect(self):
        try:
            p = db_norm(self.db_path.text())
        except ValueError as exc:
            return self.warn(str(exc))
        if not p.lower().endswith((".duckdb", ".db", ".ddb")):
            p += ".duckdb"
        try:
            db_check_file(p)
        except ValueError as exc:
            return self.warn(str(exc))
        if not Path(p).exists():
            if QMessageBox.question(self, APP_NAME, f"O banco não existe. Criar um novo banco em:\n{p}?") != QMessageBox.Yes:
                return
            Path(p).parent.mkdir(parents=True, exist_ok=True)
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            self.engine.ui.execute(f"ATTACH {ql(p)} AS bd")
        except duckdb.Error as exc:
            QApplication.restoreOverrideCursor()
            return self.warn(str(friendly_db_error(exc, p)))
        QApplication.restoreOverrideCursor()
        self.engine.attached = p
        self.path_open = p
        self.db_path.setText(p)
        self.win.settings.setValue("db_path", p)
        self._set_connected(True)
        self.refresh_tables()

    def _disconnect(self):
        self.view.set_table(None, None)
        with contextlib.suppress(Exception):
            self.engine.ui.execute("DETACH bd")
        self.engine.forget("bd.")
        self.engine.attached = None
        self.path_open = None
        self.tables.clear()
        self.struct.set_df(None)
        self._set_connected(False)

    def _set_connected(self, on):
        self.conn_btn.setText("Desconectar" if on else "Conectar")
        self.conn_btn.setIcon(icon("x", "#4F4A62", 16) if on else icon("link", "#FFFFFF", 16))
        self.conn_btn.setProperty("primary", not on)
        self.conn_btn.style().unpolish(self.conn_btn)
        self.conn_btn.style().polish(self.conn_btn)
        for b in (self.btn_new, self.btn_ins, self.btn_ref, self.btn_del):
            b.setEnabled(on)
        self.tabs.setEnabled(on)
        self.db_path.setReadOnly(on)
        if not on:
            self.status.setText("Nenhum banco conectado. O banco fica reservado para o app enquanto estiver "
                                "conectado; desconecte para usá-lo em outro programa.")

    # --- tabelas ----------------------------------------------------------------
    def table_names(self):
        if not self.path_open:
            return []
        return db_tables(self.engine.ui, "bd")

    def table_columns(self, table):
        if not self.path_open or table.lower() not in {t.lower() for t in self.table_names()}:
            return None
        real = next(t for t in self.table_names() if t.lower() == table.lower())
        return [r[0] for r in self.engine.ui.execute(f"DESCRIBE {db_ref('bd', real)}").fetchall()]

    def refresh_tables(self):
        if not self.path_open:
            return
        current = self.tables.currentItem().data(Qt.UserRole) if self.tables.currentItem() else None
        self.engine.forget("bd.")
        rows = []
        for name, ncols in self.engine.ui.execute(
                "SELECT table_name, column_count FROM duckdb_tables() "
                "WHERE database_name = 'bd' AND schema_name = 'main' ORDER BY table_name").fetchall():
            n = self.engine.ui.execute(f"SELECT count(*) FROM {db_ref('bd', name)}").fetchone()[0]
            rows.append((name, n, ncols))
        self.tables.blockSignals(True)
        self.tables.clear()
        for name, size, ncols in rows:
            it = QListWidgetItem(icon("database", MUTED, 16), f"{name}  —  {fmt_int(size)} linhas · {ncols} colunas")
            it.setData(Qt.UserRole, name)
            it.setToolTip(f"{name}: {fmt_int(size)} linhas, {ncols} colunas")
            self.tables.addItem(it)
            if name == current:
                self.tables.setCurrentItem(it)
        self.tables.blockSignals(False)
        size = Path(self.path_open).stat().st_size if Path(self.path_open).exists() else 0
        self.status.setText(f"Conectado · {len(rows)} tabela(s) · {fmt_bytes(size)} em disco. "
                            "Enquanto conectado, o banco fica reservado para o app.")
        if self.tables.currentItem() is None and self.tables.count():
            self.tables.setCurrentRow(0)
        self._show_table()

    def _current_table(self):
        it = self.tables.currentItem()
        return it.data(Qt.UserRole) if it else None

    def _show_table(self):
        name = self._current_table()
        self.view.set_table(None, None)
        if not name:
            self.struct.set_df(None)
            return
        ref = db_ref("bd", name)
        self.engine.forget(ref)
        m = self.engine.meta(ref)
        self.view.set_table(self.engine, ref)
        self.struct.set_df(pd.DataFrame({"Coluna": m["cols"], "Tipo": [db_type_for(t) for t in m["types"]],
                                         "Tipo no DuckDB": m["types"]}))

    def _drop(self):
        name = self._current_table()
        if not name:
            return
        if QMessageBox.question(self, APP_NAME, f"Excluir a tabela '{name}' do banco?\n\nIsso apaga os dados dela "
                                                "no arquivo e não pode ser desfeito.") != QMessageBox.Yes:
            return
        self.view.set_table(None, None)
        try:
            self.engine.ui.execute(f"DROP TABLE {db_ref('bd', name)}")
        except duckdb.Error as exc:
            return self.warn(f"Não consegui excluir: {exc}")
        self.refresh_tables()

    # --- ações (podem ir para a receita) ----------------------------------------
    def _new_table(self):
        dlg = CreateTableDialog(self)
        polish(dlg)
        if dlg.exec() != QDialog.Accepted:
            return
        table, cols = dlg.name.text().strip(), dlg.columns()
        params = {"db": self.path_open, "table": table, "columns": cols}
        then = None
        if dlg.fill.isChecked() and dlg.src.currentData():
            layer, src = dlg.src.currentData()
            wparams = {"db": self.path_open, "table": table, "src": src, "mode": "append", "try_cast": True}
            then = lambda: self.win.run_db("db_write", wparams, "Gravação no banco",  # noqa: E731
                                           f"{src} → {table} (acrescentar)", dlg.record.isChecked())
        self.win.run_db("db_create", params, "Criação de tabela", f"{table} ({len(cols)} colunas)",
                        dlg.record.isChecked(), then=then)

    def _insert(self):
        dlg = InsertDialog(self, self._current_table())
        polish(dlg)
        if dlg.exec() != QDialog.Accepted:
            return
        layer, src = dlg.src.currentData()
        table, mode = dlg.dest.currentText().strip(), dlg.mode.currentData()
        params = {"db": self.path_open, "table": table, "src": src, "mode": mode,
                  "try_cast": dlg.try_cast.isChecked()}
        modo = {"append": "acrescentar", "truncate": "limpar e inserir", "replace": "recriar"}[mode]
        self.win.run_db("db_write", params, "Gravação no banco", f"{src} → {table} ({modo})", dlg.record.isChecked())

    def _run_sql(self):
        sql = self.sql.toPlainText().strip()
        if not sql:
            return self.warn("Escreva o SQL a executar.")
        first = sql.split()[0].upper() if sql.split() else ""
        record = self.sql_record.isChecked() and first not in ("SELECT", "WITH", "DESCRIBE", "SHOW", "SUMMARIZE")
        label = " ".join(sql.split())
        self.win.run_db("db_sql", {"db": self.path_open, "sql": sql}, "SQL no banco",
                        label[:120] + ("…" if len(label) > 120 else ""), record)

    def after_db(self, op):
        """Chamado depois de qualquer operação no banco."""
        if op == "db_sql":
            df = self.engine.last_sql_result
            self.sql_result.set_df(df)
            self.sql_info.setText("Executado." if df is None else f"{fmt_int(len(df))} linha(s) (até 1.000).")
        if op == "db_write" and self.engine.last_report:
            r = self.engine.last_report
            extra = ""
            if not r["created"]:
                extra = f" · {len(r['common'])} colunas"
                if r["ignored"]:
                    extra += f" · ignoradas: {', '.join(r['ignored'][:6])}"
                if r["missing"]:
                    extra += f" · vazias: {', '.join(r['missing'][:6])}"
            self.win.statusBar().showMessage(
                f"Gravado em {r['table']}: {fmt_int(r['rows'])} linhas{extra}", 12000)
        self.refresh_tables()

    def on_result(self, step, results):
        if step.get("layer") == "Banco":
            self.after_db(step["op"])


class AutomationPage(Page):
    def __init__(self, win):
        super().__init__(win, "Automação e agendamento",
                         "Cada passo feito nas outras telas é gravado como uma receita. Salve a receita e "
                         "agende o programa para reexecutá-la sozinho, sem abrir a interface.")
        self.recipe_path = None
        self._saved_steps = None

        g = QGroupBox("1. Receita do pipeline")
        v = QVBoxLayout(g)
        v.addWidget(StepsLog(self.store, min_height=180))
        row = QHBoxLayout()
        row.addWidget(make_button("💾 Salvar receita…", self._save, True))
        row.addWidget(make_button("📂 Abrir receita…", self._open))
        row.addStretch()
        v.addLayout(row)
        self.recipe_lbl = hint("Receita ainda não salva.")
        v.addWidget(self.recipe_lbl)
        self.body.addWidget(g)

        g = QGroupBox("2. Saída da execução automática")
        form = QFormLayout(g)
        row = QHBoxLayout()
        self.out_dir = QLineEdit()
        self.out_dir.setPlaceholderText(r"Ex: \\servidor\relatorios\ouro   (vazio = pasta 'saida' ao lado da receita)")
        row.addWidget(self.out_dir, 1)
        row.addWidget(make_button("Pasta…", self._pick_out))
        form.addRow("Pasta de saída", row)
        fmt_row = QHBoxLayout()
        self.f_xlsx, self.f_parquet = QCheckBox("Excel (.xlsx)"), QCheckBox("Parquet")
        self.f_duckdb = QCheckBox("DuckDB (todas as tabelas em dados.duckdb)")
        self.f_xlsx.setChecked(True)
        fmt_row.addWidget(self.f_xlsx)
        fmt_row.addWidget(self.f_parquet)
        fmt_row.addWidget(self.f_duckdb)
        fmt_row.addStretch()
        form.addRow("Formatos", fmt_row)
        layer_row = QHBoxLayout()
        self.l_checks = {layer: QCheckBox(LAYER_LABEL[layer]) for layer in ("gold", "silver", "bronze")}
        for layer, cb in self.l_checks.items():
            cb.setIcon(layer_icon(layer))
        self.l_checks["gold"].setChecked(True)
        for cb in self.l_checks.values():
            layer_row.addWidget(cb)
        layer_row.addStretch()
        form.addRow("Exportar camadas", layer_row)
        self.body.addWidget(g)

        g = QGroupBox("Código Python (sem o app)")
        v = QVBoxLayout(g)
        v.addWidget(hint("Gera os scripts do pipeline, um por camada: BD.py (tabelas), etl_bronze.py "
                         "(carga), etl_prata.py (tratamentos) e etl_ouro.py (consolidação e Parquet), mais um "
                         "executar_tudo.py. Bom para entregar a quem não usa o app ou versionar no Git."))
        v.addWidget(make_button("📄 Gerar código Python…", self._make_code, True), 0, Qt.AlignLeft)
        self.code_box = g

        g4 = g
        g = QGroupBox("3. Testar e agendar no Windows")
        v = QVBoxLayout(g)
        row = QHBoxLayout()
        row.addWidget(make_button("▶ Executar receita agora", self._run_now, True))
        row.addWidget(make_button("⬇️ Gerar .bat de agendamento…", self._make_bat))
        row.addStretch()
        v.addLayout(row)
        steps = QLabel(
            "<ol style='margin-left:-20px'>"
            "<li>Salve a receita e gere o <b>.bat</b> (fica ao lado da receita).</li>"
            "<li>Pressione <b>Win + S</b>, digite <b>Agendador de Tarefas</b> e abra.</li>"
            "<li>Clique em <b>Criar Tarefa Básica…</b>, dê um nome e escolha a recorrência (ex: diariamente às 06:00).</li>"
            "<li>Em Ação, escolha <b>Iniciar um programa</b> e selecione o .bat gerado.</li>"
            "<li>Conclua. Cada execução grava um <b>.log</b> ao lado da receita com o resultado.</li>"
            "</ol>"
            "Dica: em tarefas agendadas, unidades mapeadas (ex: <b>Z:</b>) podem não existir. "
            "Prefira caminhos de rede completos, como <b>\\\\servidor\\pasta</b>.")
        steps.setWordWrap(True)
        steps.setTextFormat(Qt.RichText)
        v.addWidget(steps)
        self.body.addWidget(g)
        self.body.addWidget(g4)
        self.body.addStretch()

    def refresh(self):
        if self.recipe_path is None:
            self.recipe_lbl.setText("Receita ainda não salva.")
        elif self.store.recipe_steps() != self._saved_steps:
            self.recipe_lbl.setText(f"{self.recipe_path} — há alterações não salvas.")
        else:
            self.recipe_lbl.setText(f"Salva em {self.recipe_path}")

    def _pick_out(self):
        path = QFileDialog.getExistingDirectory(self, "Pasta de saída")
        if path:
            self.out_dir.setText(str(Path(path)))

    def _build_recipe(self):
        steps = self.store.recipe_steps()
        if not steps:
            self.warn("Ainda não há passos. Monte o pipeline nas outras telas primeiro.")
            return None
        formats = [f for f, cb in (("xlsx", self.f_xlsx), ("parquet", self.f_parquet), ("duckdb", self.f_duckdb))
                   if cb.isChecked()]
        layers = [layer for layer, cb in self.l_checks.items() if cb.isChecked()]
        if not formats or not layers:
            self.warn("Escolha ao menos um formato e uma camada para exportar.")
            return None
        return {"app": APP_NAME, "version": 2, "created": dt.datetime.now().isoformat(timespec="seconds"),
                "output_dir": self.out_dir.text().strip(), "formats": formats, "layers": layers,
                "settings": {"mem_mb": self.win.settings.value("mem_mb", 0, int)},
                "export_columns": self.store.recipe_export_cols(),
                "steps": steps}

    def _save(self) -> bool:
        recipe = self._build_recipe()
        if not recipe:
            return False
        start = str(self.recipe_path) if self.recipe_path else "receita.json"
        path, _ = QFileDialog.getSaveFileName(self, "Salvar receita", start, "Receita (*.json)")
        if not path:
            return False
        Path(path).write_text(json.dumps(recipe, ensure_ascii=False, indent=2), encoding="utf-8")
        self.recipe_path = Path(path)
        self._saved_steps = recipe["steps"]
        self.refresh()
        return True

    def _open(self):
        path, _ = QFileDialog.getOpenFileName(self, "Abrir receita", "", "Receita (*.json)")
        if not path:
            return
        try:
            recipe = json.loads(Path(path).read_text(encoding="utf-8"))
            steps = recipe["steps"]
        except Exception as exc:
            return self.warn(f"Arquivo de receita inválido: {exc}")
        self.out_dir.setText(recipe.get("output_dir", ""))
        self.f_xlsx.setChecked("xlsx" in recipe.get("formats", []))
        self.f_parquet.setChecked("parquet" in recipe.get("formats", []))
        self.f_duckdb.setChecked("duckdb" in recipe.get("formats", []))
        for layer, cb in self.l_checks.items():
            cb.setChecked(layer in recipe.get("layers", []))

        def loaded(result):
            self.recipe_path = Path(path)
            self._saved_steps = [clean_step(s) for s in steps]
            self.store.load(*result, recipe.get("export_columns"))
            self.win.statusBar().showMessage(f"Receita carregada: {len(steps)} passos reexecutados", 10000)

        self.win.run_task(replay, (self.engine, steps), loaded, "Reexecutando a receita…")

    def _run_now(self):
        recipe = self._build_recipe()
        if not recipe:
            return
        base = self.recipe_path.parent if self.recipe_path else Path.cwd()

        def done(files):
            pasta = files[0].parent if files else "-"
            QMessageBox.information(self, APP_NAME, f"Receita executada do zero com sucesso.\n\n"
                                                    f"{len(files)} arquivo(s) gerado(s) em:\n{pasta}")

        self.win.run_task(lambda r, b, ctx: run_recipe(r, b, ctx), (recipe, base), done,
                          "Executando a receita do zero…")

    def _make_code(self):
        if not self.store.steps:
            return self.warn("Ainda não há passos. Monte o pipeline nas outras telas primeiro.")
        dlg = CodeDialog(self.win)
        polish(dlg)
        if dlg.exec() != QDialog.Accepted or not dlg.written:
            return
        pasta = dlg.written[0].parent
        box = QMessageBox(QMessageBox.Information, APP_NAME,
                          f"{len(dlg.written)} arquivos gerados em:\n{pasta}\n\n"
                          + "\n".join(f"• {p.name}" for p in dlg.written)
                          + "\n\nPara rodar tudo: python executar_tudo.py", parent=self)
        open_btn = box.addButton("Abrir pasta", QMessageBox.ActionRole)
        box.addButton("Fechar", QMessageBox.RejectRole)
        box.exec()
        if box.clickedButton() is open_btn:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(pasta)))

    def _make_bat(self):
        if self.recipe_path is None or self.store.recipe_steps() != self._saved_steps:
            QMessageBox.information(self, APP_NAME, "Primeiro salve a receita atual; o .bat vai executar esse arquivo.")
            if not self._save():
                return
        recipe = self.recipe_path.resolve()
        bat_path = recipe.with_suffix(".bat")
        if getattr(sys, "frozen", False):
            cmd = f'start "" /wait "{sys.executable}" --run "{recipe}"'
        else:
            cmd = f'"{sys.executable}" "{Path(__file__).resolve()}" --run "{recipe}"'
        bat = ("@echo off\r\nchcp 65001 >nul\r\n"
               f"rem Gerado pelo {APP_NAME} em {dt.datetime.now():%d/%m/%Y %H:%M}\r\n"
               f"{cmd}\r\nexit /b %ERRORLEVEL%\r\n")
        bat_path.write_text(bat, encoding="utf-8")
        QMessageBox.information(self, APP_NAME, f".bat gerado em:\n{bat_path}\n\n"
                                                "Selecione esse arquivo no Agendador de Tarefas.")


# ============================================================================
# JANELA PRINCIPAL
# ============================================================================
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.settings = QSettings("Refina", "Refina")
        splash_step("Limpando sessões antigas…", 65)
        clean_stale_sessions()
        splash_step("Preparando o banco de dados…", 72)
        self.engine = Engine(self.settings.value("mem_mb", 0, int))
        self.store = Store(self.engine)
        self.pool = QThreadPool.globalInstance()
        self._task = None
        self._progress = None
        self._poll = QTimer(self)
        self._poll.setInterval(400)
        self._poll.timeout.connect(self._poll_query)
        self._explicit_pct = False
        self.setWindowTitle(f"{APP_NAME} · Refinaria de dados")
        self.resize(1320, 860)
        self.setAcceptDrops(True)

        central = QWidget()
        h = QHBoxLayout(central)
        h.setContentsMargins(0, 0, 0, 0)
        h.setSpacing(0)
        h.addWidget(self._build_sidebar())

        splash_step("Montando as telas…", 82)
        self.stack = QStackedWidget()
        self.pages = []
        for label, cls in [("upload|Upload", UploadPage), ("search|Explorar", ExplorerPage),
                           ("layers|Prata", SilverPage), ("award|Ouro", GoldPage),
                           ("database|Banco de dados", DatabasePage), ("clock|Automação", AutomationPage)]:
            page = cls(self)
            self.pages.append(page)
            scroll = QScrollArea()
            scroll.setWidgetResizable(True)
            scroll.setFrameShape(QFrame.NoFrame)
            scroll.setWidget(page)
            self.stack.addWidget(scroll)
            ic, text = label.split("|")
            self.nav.addItem(QListWidgetItem(icon(ic, "#958FAB", 18, selected="#FFFFFF"), text))
        self.upload, self.explorer, self.silver, self.gold, self.database, self.automation = self.pages
        h.addWidget(self.stack, 1)
        self.setCentralWidget(central)
        polish(self)
        self.setWindowIcon(QIcon(svg_pixmap(icon_svg("logo", ""), 32)))

        self.nav.currentRowChanged.connect(self.stack.setCurrentIndex)
        self.nav.setCurrentRow(0)
        self.store.changed.connect(self._update_badges)
        self.statusBar().showMessage(f"Pronto. Memória máxima: {fmt_bytes(self.engine.mem_mb * 2**20)}.  "
                                     f"Atalhos: Ctrl+O abrir arquivos · Ctrl+Z desfazer · Ctrl+1…6 telas.")
        self._update_badges()
        for seq, fn in ((QKeySequence.Undo, self._undo), (QKeySequence.Open, self.upload.pick_files)):
            QShortcut(seq, self, fn)
        for i in range(6):
            QShortcut(QKeySequence(f"Ctrl+{i + 1}"), self, lambda i=i: self.nav.setCurrentRow(i))
        geo = self.settings.value("geometry")
        if geo is not None:
            with contextlib.suppress(Exception):
                self.restoreGeometry(geo)

    # --- navegação entre telas ------------------------------------------------------
    def open_in_explorer(self, layer, name):
        self.nav.setCurrentRow(1)
        self.explorer.select(layer, name)

    def goto_export(self, layer, name):
        self.nav.setCurrentRow(3)
        self.gold.tabs.setCurrentIndex(2)
        self.gold.select_export(layer, name)

    def recent_paths(self) -> list:
        v = self.settings.value("recent_paths", [])
        if isinstance(v, str):
            v = [v]
        return [str(p) for p in (v or [])]

    def remove_recent(self, path):
        self.settings.setValue("recent_paths", [p for p in self.recent_paths() if p != path])
        self.upload.refresh_recent()
        self.statusBar().showMessage("Caminho removido do histórico.", 4000)

    def clear_recent(self):
        self.settings.setValue("recent_paths", [])
        self.upload.refresh_recent()

    def add_recent(self, paths):
        cur = [p for p in self.recent_paths() if p not in paths]
        self.settings.setValue("recent_paths", (list(paths) + cur)[:8])
        self.upload.refresh_recent()

    def _build_sidebar(self):
        side = QFrame()
        side.setObjectName("sidebar")
        side.setFixedWidth(240)
        v = QVBoxLayout(side)
        v.setContentsMargins(16, 20, 16, 16)
        brand = QHBoxLayout()
        brand.setSpacing(10)
        logo = QLabel()
        logo.setPixmap(svg_pixmap(icon_svg("logo", ""), 30))
        brand.addWidget(logo)
        names = QVBoxLayout()
        names.setSpacing(0)
        title = QLabel(APP_NAME)
        title.setObjectName("appTitle")
        sub = QLabel("Refinaria de dados")
        sub.setObjectName("appSubtitle")
        names.addWidget(title)
        names.addWidget(sub)
        brand.addLayout(names)
        brand.addStretch()
        v.addLayout(brand)
        v.addSpacing(22)
        self.nav = QListWidget()
        self.nav.setObjectName("nav")
        self.nav.setIconSize(QSize(18, 18))
        self.nav.setFixedHeight(6 * 44)
        self.nav.setCursor(Qt.PointingHandCursor)
        v.addWidget(self.nav)
        v.addSpacing(18)
        sec = QLabel("Camadas")
        sec.setObjectName("sectionLabel")
        v.addWidget(sec)
        v.addSpacing(2)
        self.badges = {}
        for key, label in [("bronze", "Bronze"), ("silver", "Prata"), ("gold", "Ouro"), ("steps", "Passos aplicados")]:
            card = ClickCard(lambda key=key: self._card_clicked(key))
            card.setObjectName("statusCard")
            card.setToolTip("Ver o histórico de passos" if key == "steps" else "Ver os datasets desta camada")
            hl = QHBoxLayout(card)
            hl.setContentsMargins(10, 6, 10, 6)
            hl.setSpacing(8)
            mark = QLabel()
            mark.setPixmap(svg_pixmap(icon_svg("list" if key == "steps" else f"dot:{key}", "#8F89A6"),
                                      16 if key == "steps" else 14))
            mark.setFixedWidth(18)
            hl.addWidget(mark)
            hl.addWidget(QLabel(label))
            hl.addStretch()
            badge = QLabel("0")
            badge.setObjectName("count")
            badge.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            hl.addWidget(badge)
            self.badges[key] = badge
            v.addWidget(card)
        v.addStretch()
        zrow = QHBoxLayout()
        zrow.setContentsMargins(10, 0, 4, 0)
        zlbl = QLabel("Zoom")
        zlbl.setObjectName("sectionLabel")
        zrow.addWidget(zlbl)
        zrow.addStretch()
        self.zoom_lbl = QLabel()
        self.zoom_lbl.setObjectName("count")
        self.zoom_lbl.setAlignment(Qt.AlignCenter)
        self.zoom_lbl.setMinimumWidth(44)
        for text, delta, tip in (("−", -1, "Diminuir (afasta a visualização)"), ("+", 1, "Aumentar")):
            b = QToolButton()
            b.setObjectName("sideZoom")
            b.setText(text)
            b.setToolTip(tip)
            b.setCursor(Qt.PointingHandCursor)
            b.clicked.connect(lambda _=False, d=delta: self._zoom(d))
            if delta < 0:
                zrow.addWidget(b)
                zrow.addWidget(self.zoom_lbl)
            else:
                zrow.addWidget(b)
        v.addLayout(zrow)
        v.addSpacing(8)
        self._update_zoom_label()
        row = QHBoxLayout()
        undo = make_button("↩️ Desfazer", self._undo, icon_color="#BAB4CC")
        undo.setToolTip("Desfaz a última operação (Ctrl+Z)")
        clear = make_button("🗑️ Limpar", self._clear, icon_color="#BAB4CC")
        clear.setToolTip("Apaga todos os dados e o histórico")
        for b in (undo, clear):
            b.setObjectName("sideButton")
        row.addWidget(undo)
        row.addWidget(clear)
        v.addLayout(row)
        return side

    def _update_zoom_label(self):
        saved = saved_ui_scale()
        current = float(os.environ.get("QT_SCALE_FACTOR") or 1)
        self.zoom_lbl.setText(f"{round(saved * 100)}%")
        self.zoom_lbl.setToolTip("Tamanho da interface" + ("" if abs(saved - current) < 1e-3
                                                            else " (vale ao reabrir o app)"))

    def _zoom(self, delta):
        cur = saved_ui_scale()
        i = min(max(UI_SCALES.index(cur) + delta, 0), len(UI_SCALES) - 1)
        new = UI_SCALES[i]
        if new == cur:
            return
        self.settings.setValue("ui_scale", new)
        self.settings.sync()
        self._update_zoom_label()
        box = QMessageBox(QMessageBox.Question, APP_NAME,
                          f"Zoom da interface: {round(new * 100)}%.\n\nO novo tamanho vale ao reabrir o app."
                          + ("\n\nOs dados desta sessão serão fechados (receitas salvas não são afetadas)."
                             if self.store.steps else ""), parent=self)
        now = box.addButton("Reabrir agora", QMessageBox.AcceptRole)
        box.addButton("Depois", QMessageBox.RejectRole)
        box.exec()
        if box.clickedButton() is now:
            self.restart()

    def restart(self):
        if getattr(sys, "frozen", False):
            program, args = sys.executable, sys.argv[1:]
        else:
            program, args = sys.executable, [os.path.abspath(sys.argv[0])] + sys.argv[1:]
        QProcess.startDetached(program, args)
        self.close()
        QApplication.quit()

    def _card_clicked(self, key):
        if key == "steps":
            self.nav.setCurrentRow(5)
        else:
            self.nav.setCurrentRow(1)
            self.explorer.filter_layer(key)

    def _update_badges(self):
        for layer in LAYERS:
            self.badges[layer].setText(str(len(self.store.data[layer])))
        self.badges["steps"].setText(str(len(self.store.steps)))

    def _undo(self):
        if self._task is not None:
            self.statusBar().showMessage("Aguarde a operação atual terminar para desfazer.", 5000)
            return
        last = self.store.steps[-1] if self.store.steps else None
        if self.store.undo():
            if last and last.get("layer") == "Banco":
                QMessageBox.information(self, APP_NAME, "O passo foi tirado da receita, mas o que já foi gravado "
                                                        "no banco continua lá (para remover, use a tela Banco de dados).")
            self.statusBar().showMessage("↩️ Última operação desfeita.", 5000)
        else:
            self.statusBar().showMessage("Nada a desfazer.", 5000)

    def _clear(self):
        if self._task is not None:
            return
        if QMessageBox.question(self, APP_NAME, "Apagar todos os dados e o histórico de passos?") == QMessageBox.Yes:
            self.store.reset()

    def closeEvent(self, event):
        if self._task is not None:
            self._task[0].ctx.cancel()
            self.pool.waitForDone(5000)
        self.settings.setValue("geometry", self.saveGeometry())
        self.engine.close()
        super().closeEvent(event)

    # --- Arrastar e soltar ------------------------------------------------------
    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self, event):
        paths = [u.toLocalFile() for u in event.mimeData().urls() if u.isLocalFile()]
        if paths:
            self.load_paths(paths)

    # --- Carregamento: analisa → usuário escolhe abas → lê em blocos ------------
    def load_paths(self, paths):
        clean = [c for c in (clean_input(p) for p in paths) if c]
        paths = [p if is_url(p) else str(Path(p).resolve()) for p in clean]
        if not paths:
            return
        self.add_recent(paths)
        self.run_task(inspect_paths, (paths,), lambda infos: self._after_inspect(paths, infos),
                      "Analisando arquivos…")

    def _after_inspect(self, paths, infos):
        ok = [i for i in infos if not i["error"]]
        if not ok:
            erros = "\n".join(f"• {i['file']}: {i['error']}" for i in infos)
            return self.show_error("Nenhum arquivo pôde ser lido.", erros)
        n_items = sum(len(i["items"]) for i in ok)
        needs_choice = n_items > 1 or len(infos) > len(ok) or \
            any(i["kind"] in ("excel", "excel_calamine", "csv") for i in ok)
        select, headers = None, {}
        if needs_choice:
            dlg = SelectDialog(infos, self)
            if dlg.exec() != QDialog.Accepted:
                return
            select, headers = dlg.selection(), dlg.headers()
        parts = []
        for i in ok:
            chosen = select.get(i["path"]) if select else [it["name"] for it in i["items"]]
            if not chosen:
                continue
            parts.append(i["file"] if len(i["items"]) == 1 else f"{i['file']} ({', '.join(chosen)})")
        desc = "Lido de: " + "; ".join(parts)
        if len(desc) > 220:
            desc = desc[:217] + "…"
        self.run_step("Bronze", "load_paths", {"paths": paths, "select": select, "headers": headers},
                      "Ingestão", desc)

    # --- Execução em segundo plano ------------------------------------------------
    def run_step(self, layer, op, params, label, desc, then=None):
        step = {"layer": layer, "op": op, "params": params, "label": label, "desc": desc}
        snapshot = {k: dict(v) for k, v in self.store.data.items()}
        self.run_task(run_op, (self.engine, snapshot, op, params),
                      lambda results: self._step_done(step, results, then), f"Processando: {label}…")

    def run_db(self, op, params, label, desc, record, then=None):
        """Operação no banco externo: vira passo da receita (record=True) ou roda só agora."""
        if record:
            return self.run_step("Banco", op, params, label, desc, then=then)
        snapshot = {k: dict(v) for k, v in self.store.data.items()}

        def done(_):
            self.database.after_db(op)
            if then:
                QTimer.singleShot(0, then)

        self.run_task(run_op, (self.engine, snapshot, op, params), done, f"{label}…")

    def _step_done(self, step, results, then=None):
        self.store.apply(step, results)
        for page in self.pages:
            if hasattr(page, "on_result"):
                page.on_result(step, results)
        if then:
            QTimer.singleShot(0, then)
        if not results:
            msg = f"{step['label']}: {step['desc']}"
        elif len(results) == 1:
            _, name, tid = results[0]
            msg = f"{step['label']}: {name} ({fmt_int(self.engine.rows(tid))} linhas)"
        else:
            msg = f"{step['label']}: {len(results)} datasets carregados"
        self.statusBar().showMessage(msg, 10000)

    def run_task(self, fn, args, on_done, message="Processando…"):
        if self._task is not None:
            QMessageBox.information(self, APP_NAME, "Aguarde a operação atual terminar.")
            return
        worker = Worker(fn, args)
        worker.signals.done.connect(self._task_done)
        worker.signals.failed.connect(self._task_failed)
        worker.signals.cancelled.connect(self._task_cancelled)
        worker.signals.progress.connect(self._task_progress)
        self._task = (worker, on_done)
        self._explicit_pct = False
        self._progress = QProgressDialog(message, "Cancelar", 0, 0, self)
        self._progress.setWindowTitle(APP_NAME)
        self._progress.setWindowModality(Qt.WindowModal)
        self._progress.setMinimumWidth(460)
        self._progress.setAutoClose(False)
        self._progress.setAutoReset(False)
        self._progress.setMinimumDuration(10 ** 9)
        self._progress.canceled.connect(self._cancel_task)
        QTimer.singleShot(300, self._maybe_show_progress)
        QApplication.setOverrideCursor(Qt.WaitCursor)
        self.statusBar().showMessage(message)
        self._poll.start()
        self.pool.start(worker)

    def _maybe_show_progress(self):
        if self._task is not None and self._progress is not None:
            self._progress.show()

    def _cancel_task(self):
        if self._task is not None:
            self._progress.setLabelText("Cancelando…")
            self._task[0].ctx.cancel()

    def _task_progress(self, msg, pct):
        if self._progress is None:
            return
        self._progress.setLabelText(msg)
        self.statusBar().showMessage(msg)
        if pct >= 0:
            self._explicit_pct = True
            self._progress.setRange(0, 100)
            self._progress.setValue(pct)
        else:
            self._explicit_pct = False
            self._progress.setRange(0, 0)

    def _poll_query(self):
        """Mostra o progresso das consultas pesadas do DuckDB (merge, agregação, exportação)."""
        if self._task is None or self._progress is None or self._explicit_pct:
            return
        p = self._task[0].ctx.query_progress()
        if p >= 0:
            self._progress.setRange(0, 100)
            self._progress.setValue(int(p))

    def _finish_task(self):
        self._poll.stop()
        QApplication.restoreOverrideCursor()
        if self._progress is not None:
            self._progress.canceled.disconnect()
            self._progress.close()
            self._progress.deleteLater()
            self._progress = None
        _, on_done = self._task
        self._task = None
        self.statusBar().clearMessage()
        return on_done

    def _task_done(self, result):
        on_done = self._finish_task()
        try:
            on_done(result)
        except Exception as exc:
            self.show_error(str(exc), traceback.format_exc())
        self.store.gc()

    def _task_failed(self, message, details):
        self._finish_task()
        self.store.gc()
        self.show_error(message, details)

    def _task_cancelled(self):
        self._finish_task()
        self.store.gc()
        self.statusBar().showMessage("Operação cancelada.", 6000)

    def show_error(self, message, details=""):
        box = QMessageBox(QMessageBox.Critical, APP_NAME, f"{message}", parent=self)
        if details:
            box.setDetailedText(details)
        box.exec()


# ============================================================================
# ENTRADA PRINCIPAL
# ============================================================================
def main():
    if "--run" in sys.argv:
        i = sys.argv.index("--run")
        if i + 1 >= len(sys.argv):
            print("Uso: app.py --run receita.json")
            sys.exit(2)
        sys.exit(cli_run(sys.argv[i + 1]))

    if QApplication.instance() is None:
        apply_ui_scale()
    app = QApplication.instance() or QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setStyle("Fusion")
    try:
        app.styleHints().setColorScheme(Qt.ColorScheme.Light)
    except AttributeError:
        pass
    app.setPalette(light_palette())
    splash_step("Aplicando o visual…", 60)
    app.setStyleSheet(build_style())
    win = MainWindow()
    splash_step("Pronto!", 100)
    win.show()
    if _SPLASH is not None:
        _SPLASH.finish(win)
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
